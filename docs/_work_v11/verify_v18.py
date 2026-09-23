# -*- coding: utf-8 -*-
"""Verify v18: header checks + table 9 arithmetic + citations + terminology."""
import re
import zipfile
import xml.etree.ElementTree as ET

import docx

DOCX = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉精修版-v18.docx"
with zipfile.ZipFile(DOCX) as z:
    xml = z.read('word/document.xml').decode('utf-8')
    print('testzip:', z.testzip() or 'OK')
ET.fromstring(xml)
print('XML: OK')
d = docx.Document(DOCX)
print('段落 %d / 表格 %d / 图片 %d' % (len(d.paragraphs), len(d.tables), len(d.inline_shapes)))

# ------------------------------------------------- 表9 逐行原文 + 算术
t9 = next(t for t in d.tables
          if any('营业收入 基准' in c.text for row in t.rows for c in row.cells))
rows = {}
for row in t9.rows:
    key = row.cells[0].text.strip()
    rows[key] = [c.text.strip() for c in row.cells[1:]]
    print('\n%-22s | %s' % (key, ' | '.join(rows[key])))

print('\n--- 算术校验 ---')
costs = [float(x) for x in re.findall(r'\d+(?:\.\d+)?',
                                      rows['经营成本 保守 / 基准 / 进取'][0])]
print('   成本 保守/基准/进取: %s' % costs)
for k, scen in enumerate(('保守', '基准', '进取')):
    rev = [float(x) for x in re.findall(r'\d+(?:\.\d+)?', rows['营业收入 ' + scen][0])]
    resraw = rows['经营结果 ' + scen][0]
    halves = resraw.split('万元')
    shown = []
    for j, seg in enumerate(halves[:3]):
        v = float(re.search(r'\d+(?:\.\d+)?', seg).group(0))
        shown.append(-v if '负' in seg else v)
    calc = [round(rev[i] - costs[i], 1) for i in range(3)]
    ok = all(abs(calc[i] - shown[i]) < 0.06 for i in range(3))
    print('   %s: 收入%s 成本%s -> 计算%s 表载%s %s  两年累计 %.1f 万元' % (
        scen, rev, costs, calc, shown, '✔' if ok else '✘',
        round(sum(rev[:2]) - sum(costs[:2]), 1)))

# ------------------------------------------------- 正文
def para_with(kw):
    return [p.text for p in d.paragraphs if kw in p.text]

print('\n=== 10.2 口径段 ===')
for t in para_with('年均在付个人数'):
    print('   ' + t)
print('\n=== 10.3 资金段 ===')
for t in para_with('新增资金需求'):
    print('   ' + t)

# ------------------------------------------------- 引用
full = '\n'.join(p.text for p in d.paragraphs)
cut = full.rfind('参考资料')
bodytext = full[:cut]
seq = [int(x) for x in re.findall(r'\[(\d+)\]', bodytext)]
first = []
for n in seq:
    if n not in first:
        first.append(n)
print('\n=== 引用 ===')
print('   首次出现顺序: %s' % first)
print('   严格递增: %s' % ('是 ✔' if first == sorted(first) else '否 ✘'))
refs = re.findall(r'^\[(\d+)\]\s*(.+)$', full[cut:], re.M)
nums = [int(n) for n, _ in refs]
print('   文献 %d 条，编号连续: %s' % (len(refs),
                                   '是 ✔' if nums == list(range(1, len(refs) + 1)) else '否 ✘'))
print('   未被引用的文献: %s' % (sorted(set(nums) - set(seq)) or '无 ✔'))

print('\n=== 术语 / 残留 ===')
print('   输入法 %d 次，输入系统 %d 次' % (bodytext.count('输入法'), bodytext.count('输入系统')))
for m in re.finditer(r'[^。\n]{0,26}输入法[^。\n]{0,26}', bodytext):
    print('   · %s' % m.group(0))
for bad in ('上一版', '47.4', '600 × 240', '66.6', '审批于', '返回原处', '原处', '原地',
            '把模糊需求转化为可检验问题'):
    print('   残留 %-16s %d' % (bad, bodytext.count(bad)))

print('\n=== 4.4 节 ===')
for t in para_with('原位交付、策略控制、结果验证与失败恢复'):
    print('   ' + t)
