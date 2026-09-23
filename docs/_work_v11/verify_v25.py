# -*- coding: utf-8 -*-
"""v25 最终校验：页码规则、目录页码与印刷页一致、中缝、结构。"""
import re
import zipfile
import xml.etree.ElementTree as ET

import docx
import pymupdf

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
DOCX = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉精修版-v25.docx"
PDF = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉精修版-v25.pdf"

doc = pymupdf.open(PDF)
FOOT, FLAT = [], []
for i in range(doc.page_count):
    f = [b[4].replace('\n', '').strip() for b in doc[i].get_text('blocks') if b[1] > 795]
    FOOT.append(f[0] if f else None)
    FLAT.append(re.sub(r'\s+', '', doc[i].get_text()))

print('=== 1. 页码规则 ===')
print('  页脚序列: %s' % ' '.join(x if x else '—' for x in FOOT))
body = FOOT[2:]
print('  封面=%s 目录=%s' % (FOOT[0] or '无页码', FOOT[1] or '无页码'))
print('  正文 %s..%s（%d 页），与 1..%d 一致: %s' % (
    body[0], body[-1], len(body), len(body),
    '✔' if body == [str(i) for i in range(1, len(body) + 1)] else '✘'))

print('\n=== 2. 目录页码 = 实际印刷页码 ===')
with zipfile.ZipFile(DOCX) as z:
    xml = z.read('word/document.xml').decode('utf-8')
toc = ET.fromstring(xml).find(W + 'body').findall(W + 'tbl')[1]
entries = []
for tr in toc.findall(W + 'tr'):
    cells = [''.join(t.text or '' for t in tc.iter(W + 't')) for tc in tr.findall(W + 'tc')]
    for i in (0, 2):
        if i + 1 < len(cells) and cells[i].strip() and cells[i + 1].strip().isdigit():
            entries.append((cells[i].strip(), int(cells[i + 1].strip())))
bad = []
for title, num in entries:
    key = re.sub(r'\s+', '', title)   # 用完整条目标题匹配（编号前缀会撞上正文里的交叉引用）
    found = next((i for i in range(2, doc.page_count) if key in FLAT[i]), None)
    real = FOOT[found] if found is not None else None
    if real is None or int(real) != num:
        bad.append((title, num, real))
print('  共 %d 条；不一致: %s' % (len(entries), bad if bad else '无 ✔'))

print('\n=== 3. 目录中缝 ===')
page = doc[1]
spans = [(s['bbox'], s['text']) for blk in page.get_text('dict')['blocks'] if blk['type'] == 0
         for l in blk['lines'] for s in l['spans'] if s['text'].strip()]
ln = max(b[2] for b, t in spans if t.strip().isdigit() and b[0] < 300)
rl = min(b[0] for b, t in spans if not t.strip().isdigit() and 300 < b[0] < 340)
sep = [d for d in page.get_drawings() if d['rect'].width < 2 and d['rect'].height > 300]
print('  左页码最右 x=%.1f | 竖线 x=%.1f | 右条目最左 x=%.1f' % (
    ln, sep[0]['rect'].x0 if sep else -1, rl))
print('  中缝净空 = %.2f cm；竖线高度 = %.0f pt' % (
    (rl - ln) / 72 * 2.54, sep[0]['rect'].height if sep else 0))
print('  目录行数（应为一页）: 单页条目 %d 条' % len(entries))

print('\n=== 4. 结构 ===')
d = docx.Document(DOCX)
with zipfile.ZipFile(DOCX) as z:
    print('  testzip:', z.testzip() or 'OK')
print('  段落 %d / 表格 %d / 图片 %d' % (len(d.paragraphs), len(d.tables), len(d.inline_shapes)))
print('  分节数 %d；PDF 页数 %d；空白页 %s' % (
    len(re.findall(r'<w:sectPr[ >]', xml)), doc.page_count,
    [i + 1 for i in range(doc.page_count)
     if not doc[i].get_text().strip() and not doc[i].get_images()] or '无'))
full = '\n'.join(p.text for p in d.paragraphs)
print('  用语残留: 返回原处 %d / 原处 %d / 原地 %d / 原位回填 %d' % (
    full.count('返回原处'), full.count('原处'), full.count('原地'), full.count('原位回填')))

