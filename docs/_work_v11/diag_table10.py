# -*- coding: utf-8 -*-
"""诊断表10（项目风险登记表）在 v23 的 docx 与 PDF 中的状态。"""
import re

import docx
import pymupdf

DOCX = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉精修版-v23.docx"
PDF = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉精修版-v23.pdf"

d = docx.Document(DOCX)
print('=== docx 里所有表格 ===')
for i, t in enumerate(d.tables, 1):
    head = ' | '.join(c.text.strip() for c in t.rows[0].cells)
    print('  表%d: %d 行 x %d 列  首行: %s' % (i, len(t.rows), len(t.columns), head[:70]))

print('\n=== 表10（项目风险登记表）docx 全文 ===')
for t in d.tables:
    head = ' '.join(c.text for c in t.rows[0].cells)
    if '风险' in head and '应对' in head:
        for r, row in enumerate(t.rows):
            print('  r%-2d %s' % (r, ' | '.join(c.text.strip()[:34] for c in row.cells)))
        break

print('\n=== PDF 中「项目风险登记表」所在页 ===')
doc = pymupdf.open(PDF)
for i in range(doc.page_count):
    txt = doc[i].get_text()
    if '项目风险登记表' in txt:
        print('  物理 p%d （页脚 %s）' % (i + 1, [
            b[4].strip() for b in doc[i].get_text('blocks') if b[1] > 795]))
        print('  ---- 该页全部文本 ----')
        for line in txt.splitlines():
            if line.strip():
                print('    ', line.strip()[:80])
        break
