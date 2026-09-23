# -*- coding: utf-8 -*-
"""Validate the generated v11 docx."""
import hashlib
import io
import os
import zipfile
import xml.etree.ElementTree as ET

import docx
from PIL import Image

DST = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉精修版-v11.docx"

print('=== 1. zip integrity ===')
with zipfile.ZipFile(DST) as z:
    bad = z.testzip()
    print('  testzip():', bad or 'OK (no corrupt entries)')
    print('  entries:', len(z.namelist()))

print('\n=== 2. XML well-formedness ===')
with zipfile.ZipFile(DST) as z:
    for n in ('word/document.xml', 'word/_rels/document.xml.rels',
              '[Content_Types].xml', 'docProps/app.xml'):
        ET.fromstring(z.read(n))
        print('  OK', n)

print('\n=== 3. python-docx open + paragraphs/tables ===')
d = docx.Document(DST)
print('  paragraphs:', len(d.paragraphs), ' tables:', len(d.tables))
print('  inline_shapes:', len(d.inline_shapes))

print('\n=== 4. embedded images ===')
with zipfile.ZipFile(DST) as z:
    media = sorted(n for n in z.namelist() if n.startswith('word/media/'))
    for n in media:
        raw = z.read(n)
        im = Image.open(io.BytesIO(raw))
        print('  %-22s %5dx%-5d %8.1f KB  md5=%s' % (
            n, im.size[0], im.size[1], len(raw) / 1024.0,
            hashlib.md5(raw).hexdigest()[:10]))

print('\n=== 5. figure captions & key wording ===')
txt = '\n'.join(p.text for p in d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            txt += '\n' + c.text
for key in ('图1  现有 AI 工作流的断点', '图2  Cuttle 从输入到成果原位回填的完整闭环',
            '图3  Cuttle 用户研究与价值验证路径', '图4  Cuttle 一个入口与四个核心系统',
            '图5  三维决策', '图6  意图胶囊', '图7  第一年试点容量',
            '图8  主要产品在任务生命周期', '图9  Cuttle 从知识工作场景样板',
            '图10  Cuttle 二十四个月研发与市场路线', '图11  Cuttle 三年经营测算',
            '3.4 原位回填机制', '4.4 关键工程体系 原位回填与策略验证恢复',
            '原位回填（Return to Origin）', '原位回填完成率'):
    print('  %-5s %s' % ('OK' if key in txt else 'MISS', key))

print('\n=== 6. residual colloquial wording in extracted text ===')
for bad in ('返回原处', '原处', '原地', '送回原处', '回到原'):
    c = txt.count(bad)
    print('  %-8s %d' % (bad, c))

print('\nfile size: %.2f MB' % (os.path.getsize(DST) / 1024.0 / 1024.0))
