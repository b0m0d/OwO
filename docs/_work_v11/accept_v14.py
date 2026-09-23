# -*- coding: utf-8 -*-
"""Final acceptance check on the v14 DOCX + exported PDF."""
import collections
import os
import re
import zipfile
import xml.etree.ElementTree as ET

import docx
import pymupdf

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
DOCX = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉精修版-v14.docx"
PDF = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉精修版-v14.pdf"

ok = []
bad = []

# ------------------------------------------------------------------ docx
with zipfile.ZipFile(DOCX) as z:
    print('DOCX 条目: %d' % len(z.namelist()))
    print('testzip  : %s' % (z.testzip() or 'OK'))
    xml = z.read('word/document.xml').decode('utf-8')
    for n in ('word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml',
              'word/styles.xml', 'word/settings.xml'):
        ET.fromstring(z.read(n))
    print('XML 解析 : OK')

d = docx.Document(DOCX)
print('段落 %d / 表格 %d / 内嵌图 %d' % (len(d.paragraphs), len(d.tables), len(d.inline_shapes)))

text = '\n'.join(p.text for p in d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            text += '\n' + c.text

# ---------------------------------------------------------- 图表号连续性
figs = sorted(int(m.group(1)) for m in re.finditer(r'^图(\d+)\s', text, re.M))
tabs = sorted(int(m.group(1)) for m in re.finditer(r'^表(\d+)\s', text, re.M))
print('\n图号: %s' % figs)
print('表号: %s' % tabs)
(ok if figs == list(range(1, len(figs) + 1)) else bad).append('图号连续性')
(ok if tabs == list(range(1, len(tabs) + 1)) else bad).append('表号连续性')

# ------------------------------------------------------------------ pdf
doc = pymupdf.open(PDF)
flat = [re.sub(r'\s+', '', p.get_text()) for p in doc]
print('\nPDF: %d 页, %.1f x %.1f cm' % (
    doc.page_count, doc[0].rect.width / 72 * 2.54, doc[0].rect.height / 72 * 2.54))
print('PDF 加密: %s' % doc.is_encrypted)
blank = [i + 1 for i, t in enumerate(flat) if not t.strip()]
print('空白页: %s' % (blank or '无'))

# ------------------------------------------- 目录页码 vs 实际页码（全量）
body = ET.fromstring(xml).find(W + 'body')
toc = body.findall(W + 'tbl')[1]
entries = []
for tr in toc.findall(W + 'tr'):
    cells = [''.join(t.text or '' for t in tc.iter(W + 't')) for tc in tr.findall(W + 'tc')]
    for i in (0, 2):
        if i + 1 < len(cells) and cells[i].strip() and cells[i + 1].strip().isdigit():
            entries.append((cells[i].strip(), int(cells[i + 1].strip())))

print('\n目录核对（共 %d 条）:' % len(entries))
mismatch = []
for title, num in entries:
    needle = re.sub(r'\s+', '', title)
    real = None
    for i in range(3, len(flat)):   # skip cover + TOC
        if needle in flat[i]:
            real = i + 1
            break
    if real is None:
        # 项目摘要 / 章节标题可能与正文首行同页，向封面后回退搜索
        for i in range(2, len(flat)):
            if needle in flat[i]:
                real = i + 1
                break
    if real is not None and real != num:
        mismatch.append((title, num, real))
print('  不一致: %s' % (mismatch or '无 ✔'))
(ok if not mismatch else bad).append('目录页码')

# ------------------------------------------------------ 图注与图同页
cap_re = re.compile(r'^(图)(\d+)\s')
print('\n图注与图同页检查:')
for pno in range(doc.page_count):
    page = doc[pno]
    figs_img = [i['bbox'] for i in page.get_image_info() if (i['bbox'][2] - i['bbox'][0]) > 200]
    for line in page.get_text().splitlines():
        line = line.strip()
        if cap_re.match(line):
            rects = page.search_for(line[:18])
            good = bool(rects) and any(f[3] <= rects[0].y0 + 6 for f in figs_img)
            print('  p%-3d %-44s %s' % (pno + 1, line[:44], '✔' if good else '✘'))
            if not good:
                bad.append('图注分离: ' + line[:24])

# ------------------------------------------------------ 用语与字体嵌入
print('\n用语残留:')
for term in ('返回原处', '原处', '原地', '原位回填'):
    c = text.count(term)
    print('  %-8s %d' % (term, c))
    if c:
        bad.append('用语残留 ' + term)

fonts = set()
for pno in range(doc.page_count):
    for f in doc[pno].get_fonts(full=True):
        fonts.add(f[3])
print('\n嵌入字体 (%d): %s' % (len(fonts), sorted(fonts)))
allsub = all('+' in f for f in fonts)
print('全部为子集嵌入: %s' % allsub)
(ok if allsub else bad).append('字体嵌入')

print()
print('=' * 70)
print('通过: %s' % (', '.join(ok) or '—'))
print('问题: %s' % (', '.join(bad) or '无 ✔'))
print('=' * 70)
