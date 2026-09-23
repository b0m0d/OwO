# -*- coding: utf-8 -*-
"""Validate the v12 docx (figures + terminology) and re-render key PDF pages."""
import hashlib
import io
import os
import zipfile
import xml.etree.ElementTree as ET

import docx
import pymupdf
from PIL import Image

DOCX = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉精修版-v12.docx"
PDF = r"T:\创新创业\OwO-master\docs\_work_v11\pdf\Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉精修版-v12.pdf"
OUT = r"T:\创新创业\OwO-master\docs\_work_v11\pdf\pages"
os.makedirs(OUT, exist_ok=True)

print('=== 1. zip / XML ===')
with zipfile.ZipFile(DOCX) as z:
    print('  testzip:', z.testzip() or 'OK')
    print('  entries:', len(z.namelist()))
    for n in ('word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'):
        ET.fromstring(z.read(n))
    print('  XML well-formed: OK')

print('\n=== 2. embedded images ===')
with zipfile.ZipFile(DOCX) as z:
    for n in sorted(x for x in z.namelist() if x.startswith('word/media/')):
        raw = z.read(n)
        im = Image.open(io.BytesIO(raw))
        print('  %-22s %5dx%-5d %8.1f KB  md5=%s' % (
            n, im.size[0], im.size[1], len(raw) / 1024.0, hashlib.md5(raw).hexdigest()[:10]))

print('\n=== 3. document text ===')
d = docx.Document(DOCX)
print('  paragraphs:', len(d.paragraphs), 'tables:', len(d.tables), 'inline_shapes:', len(d.inline_shapes))
txt = '\n'.join(p.text for p in d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            txt += '\n' + c.text

for key in ('图2  Cuttle 从输入到成果原位交付的完整闭环',
            '3.4 原位交付机制', '4.4 关键工程体系 原位交付与策略验证恢复',
            '原位交付（Return to Origin）', '原位交付完成率',
            '并以受控执行完成成果原位交付'):
    print('  %-5s %s' % ('OK' if key in txt else 'MISS', key))

print('\n=== 4. forbidden wording ===')
for bad in ('返回原处', '原处', '原地', '原位回填'):
    print('  %-8s %d' % (bad, txt.count(bad)))

print('\n=== 5. PDF page check ===')
pdoc = pymupdf.open(PDF)
print('  pdf pages:', pdoc.page_count)
for needle, expected in (('3.4 原位交付机制', 11), ('4.4 关键工程体系 原位交付与策略验证恢复', 14)):
    pages = [i + 1 for i, p in enumerate(pdoc) if needle in p.get_text()]
    print('  %-42s expected p%-3d found %-12s %s' % (
        needle, expected, pages, 'OK' if expected in pages else 'MISMATCH'))

for pno in (2, 5, 11, 14):
    pix = pdoc[pno - 1].get_pixmap(dpi=110)
    pix.save(os.path.join(OUT, 'v12_p%03d.png' % pno))
print('  rendered pages 2, 5, 11, 14')
