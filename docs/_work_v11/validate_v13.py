# -*- coding: utf-8 -*-
"""Verify v13: structure, figures, captions, TOC page numbers; render key pages."""
import io
import os
import zipfile
import xml.etree.ElementTree as ET

import docx
import pymupdf
from PIL import Image

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
DOCX = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉精修版-v13.docx"
PDF = r"T:\创新创业\OwO-master\docs\_work_v11\pdf\Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉精修版-v13.pdf"
OUT = r"T:\创新创业\OwO-master\docs\_work_v11\pdf\pages"
os.makedirs(OUT, exist_ok=True)

print('=== 1. zip / XML ===')
with zipfile.ZipFile(DOCX) as z:
    print('  testzip:', z.testzip() or 'OK')
    print('  entries:', len(z.namelist()))
    for n in ('word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'):
        ET.fromstring(z.read(n))
    print('  XML well-formed: OK')

print('\n=== 2. media ===')
with zipfile.ZipFile(DOCX) as z:
    for n in sorted(x for x in z.namelist() if x.startswith('word/media/')):
        raw = z.read(n)
        im = Image.open(io.BytesIO(raw))
        print('  %-22s %5dx%-5d %8.1f KB' % (n, im.size[0], im.size[1], len(raw) / 1024.0))

print('\n=== 3. document ===')
d = docx.Document(DOCX)
print('  paragraphs:', len(d.paragraphs), 'tables:', len(d.tables), 'inline_shapes:', len(d.inline_shapes))
txt = '\n'.join(p.text for p in d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            txt += '\n' + c.text

import re
caps = [m.group(0) for m in re.finditer(r'图\d+[^\n]{0,50}', txt) if m.group(0).startswith('图')]
print('\n=== 4. figure captions ===')
for c in caps:
    print('  ' + c[:70])

print('\n=== 5. former 图8 gone? ===')
print('  竞争格局散点图 caption present:', '主要产品在任务生命周期' in txt)
for bad in ('返回原处', '原处', '原地', '原位回填'):
    print('  %-8s %d' % (bad, txt.count(bad)))

print('\n=== 6. PDF ===')
pdoc = pymupdf.open(PDF)
print('  pages:', pdoc.page_count)
for needle, expected in (('3.4 原位交付机制', 11),
                         ('4.4 关键工程体系 原位交付与策略验证恢复', 14),
                         ('6.2 市场进入规模', 17)):
    pages = [i + 1 for i, p in enumerate(pdoc) if needle in p.get_text()]
    print('  %-42s expected p%-3d found %-12s %s' % (
        needle, expected, pages, 'OK' if expected in pages else 'MISMATCH'))

for pno in (5, 6, 7, 12, 18, 19):
    pdoc[pno - 1].get_pixmap(dpi=110).save(os.path.join(OUT, 'v13_p%03d.png' % pno))
print('  rendered 5,6,7,12,18,19')
