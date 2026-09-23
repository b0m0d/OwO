# -*- coding: utf-8 -*-
"""Diagnose real alignment state of every table cell, from raw XML."""
import glob, os, re
from docx import Document
from docx.oxml.ns import qn

DOCS = os.path.join(r"T:\创新创业\OwO-master\docs")
SRC = [x for x in glob.glob(os.path.join(DOCS, "Cuttle*v8.docx")) if "预览" not in x][0]
print("FILE:", os.path.basename(SRC), os.path.getmtime(SRC))
doc = Document(SRC)

# map table -> caption: walk body, remember last caption paragraph
caps = {}
last_cap = None
for child in doc.element.body.iterchildren():
    if child.tag == qn('w:p'):
        txt = "".join(t.text or "" for t in child.iter(qn('w:t'))).strip()
        if txt.startswith("表"):
            last_cap = txt[:40]
    elif child.tag == qn('w:tbl'):
        caps[id(child)] = last_cap
        last_cap = None

for ti, t in enumerate(doc.tables):
    cap = caps.get(id(t._tbl), "?")
    state = {}
    nojc = 0
    sample = None
    for ri, row in enumerate(t.rows):
        for ci, c in enumerate(row.cells):
            for p in c.paragraphs:
                pPr = p._p.find(qn('w:pPr'))
                jc = None
                if pPr is not None:
                    j = pPr.find(qn('w:jc'))
                    if j is not None:
                        jc = j.get(qn('w:val'))
                if jc is None:
                    nojc += 1
                    if sample is None:
                        sample = (ri, ci, "".join(x.text or "" for x in p._p.iter(qn('w:t')))[:24])
                else:
                    state[jc] = state.get(jc, 0) + 1
    # style-level check
    styles_used = set()
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                styles_used.add(p.style.name)
    print("T%02d cap=%s  jc=%s  NO_JC=%d %s  styles=%s" % (ti, cap, state, nojc, sample or "", sorted(styles_used)))
