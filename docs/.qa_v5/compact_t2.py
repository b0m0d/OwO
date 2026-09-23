# -*- coding: utf-8 -*-
"""Compact 表 2 so it fits on one page; tighten the 2.3 heading spacing."""
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def first_cells(t):
    return [c.text.strip() for c in t.rows[0].cells]

t = None
for cand in doc.tables:
    if first_cells(cand)[:2] == ["验证方向", "验证方法"]:
        t = cand; break
assert t is not None

# 收紧单元格内边距
for r in t.rows:
    for c in r.cells:
        tcPr = c._tc.get_or_add_tcPr()
        mar = tcPr.find(qn('w:tcMar'))
        if mar is not None:
            tcPr.remove(mar)
        mar = OxmlElement('w:tcMar')
        for side, val in (('top', 40), ('start', 70), ('bottom', 40), ('end', 70)):
            e = OxmlElement('w:' + side)
            e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa')
            mar.append(e)
        tcPr.append(mar)
        # 行内段落更紧
        for p in c.paragraphs:
            p.paragraph_format.space_after = 0
            p.paragraph_format.space_before = 0
            p.paragraph_format.line_spacing = 0.95

# 表格上方的表题注与表格之间的间距收紧
for p in doc.paragraphs:
    if p.text.strip() == "表 2 核心验证指标与决策门":
        p.paragraph_format.space_after = 0
        p.paragraph_format.space_before = 0
        p.paragraph_format.keep_with_next = True
    if p.text.strip().startswith("2.3 核心验证指标"):
        p.paragraph_format.space_before = 0
        p.paragraph_format.space_after = 0

doc.save(SRC)
print("compacted")