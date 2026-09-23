# -*- coding: utf-8 -*-
"""Add 年均在付 row and recompute revenue from average paying users, matching the stated method."""
import copy, glob, os
from docx import Document
from docx.oxml.ns import qn
from docx.table import _Row

SRC = [x for x in glob.glob(os.path.join(r"T:\创新创业\OwO-master\docs", "Cuttle*v8.docx"))
       if "预览" not in x][0]
doc = Document(SRC)

def first_cells(t):
    return [c.text.strip() for c in t.rows[0].cells]

def set_cell(cell, text):
    ps = cell.paragraphs
    runs = ps[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
            r._element.getparent().remove(r._element)
    else:
        ps[0].add_run(text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

fin = None
for t in doc.tables:
    fc = first_cells(t)
    if fc[:2] == ["指标", "第一年"] and len(fc) == 5:
        fin = t
        break
assert fin is not None

AVG  = {"保守": [40, 395, 860],   "基准": [60, 855, 1850],   "进取": [90, 1435, 3300]}
TEAM = {"保守": [0, 3, 8],        "基准": [0, 5, 24],        "进取": [0, 10, 40]}
DEP  = {"保守": [0, 0, 0],        "基准": [0, 0, 2],         "进取": [0, 1, 4]}
COST = {"保守": [40, 55, 80],     "基准": [50, 105, 190],    "进取": [65, 190, 260]}
K = ("保守", "基准", "进取")
def wan(x): return round(x / 10000.0, 1)
REV = {k: [wan(AVG[k][y] * 240 + TEAM[k][y] * 40000 + DEP[k][y] * 180000) for y in range(3)] for k in K}
RES = {k: [round(REV[k][y] - COST[k][y], 1) for y in range(3)] for k in K}

# 在“付费个人”行后插入“年均在付”行
pay_tr = None
for tr in fin._element.findall(qn('w:tr')):
    txt = "".join(t.text or "" for t in tr.iter(qn('w:t')))
    if txt.strip().startswith("个人专业版 付费个人"):
        pay_tr = tr
        break
assert pay_tr is not None
new_tr = copy.deepcopy(pay_tr)
pay_tr.addnext(new_tr)
row_avg = _Row(new_tr, fin)
for ci, v in enumerate(["个人专业版 年均在付（测算基数）",
                        " / ".join(str(AVG[k][0]) for k in K),
                        " / ".join(str(AVG[k][1]) for k in K),
                        " / ".join(str(AVG[k][2]) for k in K),
                        "按上年末规模 60% 续订，当年新增按半年加权估算"]):
    set_cell(row_avg.cells[ci], v)

def sgn(v):
    return f"负 {abs(v)} 万元" if v < 0 else f"正 {v} 万元"

# 按标签逐行重写
for r in fin.rows:
    lab = r.cells[0].text.strip()
    if lab.startswith("个人专业版 付费个人"):
        set_cell(r.cells[4], "保守 / 基准 / 进取；年费 240 元，本行用于描述用户规模")
    elif lab.startswith("营业收入 保守"):
        for y in range(3): set_cell(r.cells[1 + y], f"{REV['保守'][y]} 万元")
        set_cell(r.cells[4], "年均在付 %s 人 × 240 元；团队 0 / 12 / 32 万元；无私有部署"
                 % " / ".join(str(AVG['保守'][y]) for y in range(3)))
    elif lab.startswith("营业收入 基准"):
        for y in range(3): set_cell(r.cells[1 + y], f"{REV['基准'][y]} 万元")
        set_cell(r.cells[4], "年均在付 %s 人 × 240 元；团队 0 / 20 / 96 万元；私有部署 0 / 0 / 36 万元"
                 % " / ".join(str(AVG['基准'][y]) for y in range(3)))
    elif lab.startswith("营业收入 进取"):
        for y in range(3): set_cell(r.cells[1 + y], f"{REV['进取'][y]} 万元")
        set_cell(r.cells[4], "年均在付 %s 人 × 240 元；团队 0 / 40 / 160 万元；私有部署 0 / 18 / 72 万元"
                 % " / ".join(str(AVG['进取'][y]) for y in range(3)))
    elif lab.startswith("经营成本"):
        for y in range(3):
            set_cell(r.cells[1 + y], " / ".join(str(COST[k][y]) for k in K) + " 万元")
        set_cell(r.cells[4], "研发、模型、市场与交付")
    elif lab.startswith("经营结果 保守"):
        for y in range(3): set_cell(r.cells[1 + y], sgn(RES['保守'][y]))
        set_cell(r.cells[4], "三年均未转正，用于观察下限")
    elif lab.startswith("经营结果 基准"):
        for y in range(3): set_cell(r.cells[1 + y], sgn(RES['基准'][y]))
        set_cell(r.cells[4], "第三年接近平衡，仍为负值")
    elif lab.startswith("经营结果 进取"):
        for y in range(3): set_cell(r.cells[1 + y], sgn(RES['进取'][y]))
        set_cell(r.cells[4], "第三年转正")
    elif lab.startswith("私有部署"):
        set_cell(r.cells[4], "18 万元每项目，按验收确认")

# 正文口径与资金段同步
for p in doc.paragraphs:
    t = p.text.strip()
    if t.startswith("财务测算采用个人订阅"):
        runs = p.runs
        runs[0].text = ("财务测算采用个人订阅、团队许可和私有部署三类收入，不代表项目已经取得收入、融资或用户规模。"
                        "个人订阅收入按年度平均在付用户数乘年费 240 元测算：年均在付由上年在付规模的 60% 续订、"
                        "当年新增按半年加权估算；团队许可按 4 万元每团队每年计价，自第二年起计入；"
                        "私有部署按 18 万元每项目、以验收确认。本轮预测仅计入基础订阅、团队许可与私有部署收入，"
                        "超额智能体调用收入及其对应的推理成本暂不计入。三种情景分别设定用户与客户数量，"
                        "按同一单价假设计算收入；第一年的规模由第六章的验证容量约束，个人付费用户不超过 200 人，"
                        "三种情景分别按 80、120 与 180 人测算。全部单价与转化率均属待验证假设，验证方式与门槛见表 7。")
        for r in runs[1:]:
            r.text = ""
            r._element.getparent().remove(r._element)
    elif t.startswith("项目已有资源包括学校实验室"):
        runs = p.runs
        runs[0].text = ("项目已有资源包括学校实验室与实验工位、团队成员个人设备、指导教师的方法与安全指导、"
                        "开源模型与本地推理能力，这些不计入现金需求。项目未来两年预计新增资金需求 160 万元，"
                        "其中第一年 50 万元、第二年 110 万元。资金来源以学校及竞赛创新基金、团队自筹与"
                        "外部产业合作为主，具体比例随阶段融资落实情况调整。第一年资金用于把原型推进到"
                        "可验证的试点闭环，并完成对照、消融与红队测试；第二年资金用于场景复制、团队版交付、"
                        "兼容矩阵扩展与支持运维。资金按决策门分两批释放，任一批未达标即暂停后续投入。"
                        "基准情景下前两年累计亏损约 113 万元，该额度可覆盖其所需并保留安全边际。")
        for r in runs[1:]:
            r.text = ""
            r._element.getparent().remove(r._element)

doc.save(SRC)
print("revenue conservative:", REV["保守"], "base:", REV["基准"], "aggressive:", REV["进取"])
print("results:", {k: RES[k] for k in K})
print("saved")
