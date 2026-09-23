# -*- coding: utf-8 -*-
import re, zipfile
from docx import Document
from docx.oxml.ns import qn
SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
z = zipfile.ZipFile(SRC)
print("zip ok:", z.testzip() is None, "| parts:", len(z.namelist()))
d = Document(SRC)
p = d.paragraphs
full = "\n".join(x.text for x in p) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
issues = []
for x in p:
    if x.style.name.startswith("Heading") and not x.text.strip(): issues.append("EMPTY HEADING")
tn = [int(re.match(r"^表 (\d+)", x.text.strip()).group(1)) for x in p if re.match(r"^表 \d+ ", x.text.strip())]
fn = [int(re.match(r"^图 (\d+)", x.text.strip()).group(1)) for x in p if re.match(r"^图 \d+ ", x.text.strip())]
if tn != list(range(1, len(tn)+1)): issues.append(f"TABLE SEQ {tn}")
if fn != list(range(1, len(fn)+1)): issues.append(f"FIGURE SEQ {fn}")
must = ["高教主赛道  本科生创意组","西南大学","杨俊熙","吴栩彪","谈世钊","赵恒军","张子豪",
        "个人付费用户上限 200 人","4 万元每团队每年","18 万元每项目","受控自治升级","适当信任",
        "禁止字段访问率","参考市场边界","当前主要交互形态与价值重心"]
for mk in must:
    if mk not in full: issues.append("MISSING: " + mk)
# table geometry
W = qn('w:w')
bad = 0
for ti, t in enumerate(d.tables):
    grid = t._element.find(qn('w:tblGrid'))
    tot = sum(int(g.get(W)) for g in grid.findall(qn('w:gridCol')))
    if tot != 9638: bad += 1
print("paragraphs:", len(p), "| tables:", len(d.tables), "| table caps:", len(tn), "| figure caps:", len(fn))
print("tables with non-normalized grid:", bad)
print("ISSUES:", len(issues))
for i in issues: print("  !", i)