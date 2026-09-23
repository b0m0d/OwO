# -*- coding: utf-8 -*-
"""Replace the summary evidence-chain table with a closing prose paragraph and renumber."""
import copy
from docx import Document
from docx.text.paragraph import Paragraph

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""; r._element.getparent().remove(r._element)

def first_cells(t):
    return [c.text.strip() for c in t.rows[0].cells]

# 1) 删除摘要证据链表与其题注
for t in list(doc.tables):
    if first_cells(t)[:2] == ["项目要素", "核心内容"]:
        t._element.getparent().remove(t._element)
        print("removed 摘要证据链表")
        break
for p in list(doc.paragraphs):
    if p.text.strip() == "表 1 项目评审证据链":
        p._element.getparent().remove(p._element)
        print("removed its caption")

# 2) 在摘要末尾补一段收束正文
anchor = None
for p in doc.paragraphs:
    if p.text.strip().startswith("项目从问题调研、机制验证"):
        anchor = p
if anchor is not None:
    el = copy.deepcopy(anchor._element)
    anchor._element.addnext(el)
    np = Paragraph(el, anchor._parent)
    set_text(np, "项目围绕一条主线展开：真实问题是 AI 的入口、情境、行动与结果彼此割裂，"
                 "验证方式为访谈、任务日记与跨应用流程观察；项目方案是在输入时刻理解情境并升级为受控行动，"
                 "验证方式为交互原型、演示闭环与任务日志；技术创新是意图胶囊 IC 与受控自治升级闭环，"
                 "验证方式为三组入口对照与消融实验；产业路径从知识工作场景切入，延伸到团队许可与私有部署，"
                 "验证方式为试点留存、付费验证与合作材料。")

# 3) 全部表题注与引用前移一位
RENUM = [
 ("表 2 核心用户、高频任务与验证场景", "表 1 核心用户、高频任务与验证场景"),
 ("表 3 核心验证指标与决策门", "表 2 核心验证指标与决策门"),
 ("表 4 意图胶囊的字段定义", "表 3 意图胶囊的字段定义"),
 ("表 5 三组对照实验与消融实验设计", "表 4 三组对照实验与消融实验设计"),
 ("表 6 安全控制与发布门槛", "表 5 安全控制与发布门槛"),
 ("表 7 首年验证容量与市场空间口径", "表 6 首年验证容量与市场空间口径"),
 ("表 8 产品竞争定位与任务生命周期覆盖", "表 7 产品竞争定位与任务生命周期覆盖"),
 ("表 9 产品版本、定价与单位经济", "表 8 产品版本、定价与单位经济"),
 ("表 10 团队分工与工程基础", "表 9 团队分工与工程基础"),
 ("表 11 三年经营预测与资源需求", "表 10 三年经营预测与资源需求"),
 ("表 12 项目风险登记表", "表 11 项目风险登记表"),
]
for old, new in RENUM:
    for p in doc.paragraphs:
        if p.text.strip() == old:
            set_text(p, new)
            break
# 正文引用
for p in doc.paragraphs:
    t = p.text
    if "项目的价值主张与验证场景见表 2" in t:
        set_text(p, t.replace("见表 2", "见表 1"))
    if "验证方式与门槛见表 9" in t:
        set_text(p, t.replace("见表 9", "见表 8"))
    if "核心指标与阶段目标见表 3" in t:
        set_text(p, t.replace("见表 3", "见表 2"))
    if "三年经营口径见表 11" in t:
        set_text(p, t.replace("见表 11", "见表 10"))

doc.save(SRC)
print("done. tables:", len(doc.tables))