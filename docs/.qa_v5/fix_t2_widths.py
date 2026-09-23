# -*- coding: utf-8 -*-
"""Widen the 分组 column so labels stay on one line, and simplify the header label."""
import glob, os
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCS = os.path.join(r"T:\创新创业\OwO-master\docs")
SRC = [x for x in glob.glob(os.path.join(DOCS, "Cuttle*v8.docx")) if "预览" not in x][0]
doc = Document(SRC)

WIDTHS = [1150, 2200, 1500, 3188, 1600]   # 合计 9638
assert sum(WIDTHS) == 9638, sum(WIDTHS)

def set_cell_text(cell, text):
    ps = cell.paragraphs
    runs = ps[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
            r._element.getparent().remove(r._element)
    else:
        ps[0].add_run(text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

for t in doc.tables:
    fc = [c.text.strip() for c in t.rows[0].cells]
    if len(fc) == 5 and fc[0] == "分组":
        grid = t._element.find(qn('w:tblGrid'))
        for gc, w in zip(grid.findall(qn('w:gridCol')), WIDTHS):
            gc.set(qn('w:w'), str(w))
        for row in t.rows:
            for i, tc in enumerate(row._element.findall(qn('w:tc'))):
                pr = tc.find(qn('w:tcPr'))
                if pr is None:
                    pr = OxmlElement('w:tcPr'); tc.insert(0, pr)
                tcW = pr.find(qn('w:tcW'))
                if tcW is None:
                    tcW = OxmlElement('w:tcW'); pr.insert(0, tcW)
                tcW.set(qn('w:type'), 'dxa'); tcW.set(qn('w:w'), str(WIDTHS[i]))
        set_cell_text(t.rows[0].cells[1], "验证方向")
        print("  表 2 列宽与表头已调整:", WIDTHS)
        break

doc.save(SRC)
print("saved")
