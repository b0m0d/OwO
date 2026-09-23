# -*- coding: utf-8 -*-
from docx import Document
from docx.oxml.ns import qn
SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)
n = 0
for p in doc.paragraphs:
    if not p.style.name.startswith("Heading"):
        continue
    for r in p.runs:
        rPr = r._element.find(qn('w:rPr'))
        if rPr is not None:
            r._element.remove(rPr)
            n += 1
            print("  cleared override in:", p.text.strip()[:44])
doc.save(SRC)
print("cleared:", n)
d = Document(SRC)
left = sum(1 for p in d.paragraphs if p.style.name.startswith("Heading")
           for r in p.runs if r._element.find(qn('w:rPr')) is not None)
print("headings with overrides left:", left)