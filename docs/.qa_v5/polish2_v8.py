# -*- coding: utf-8 -*-
import copy
from docx import Document
from docx.text.paragraph import Paragraph

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def ptxt(p):
    return p.text.strip()

anchor = [p for p in doc.paragraphs if ptxt(p).startswith("六个字段分别是：意图 Intent")][0]
NEW = ("字段信度（Confidence）是每个字段的元数据，用于标注该字段是否足以支撑自动执行，它本身不触发任何动作。"
       "动作风险取决于系统准备做什么，只能在生成执行计划时依据具体动作与目标对象判定，因此完全归策略引擎处理，不进入 IC。"
       "同一条意图在“只读取文件”与“覆盖原文件”两种计划下风险完全不同，把风险写入 IC 会掩盖这一区别。"
       "全文出现的 IC 均指上述六个字段的集合，不再使用其他写法。")
el = copy.deepcopy(anchor._element)
anchor._element.addnext(el)
np = Paragraph(el, anchor._parent)
runs = np.runs
runs[0].text = NEW
for r in runs[1:]:
    r.text = ""
    r._element.getparent().remove(r._element)

# relocate chapter-4.3 closing paragraph to just before the 表 7 caption
closing = [p for p in doc.paragraphs if ptxt(p).startswith("本节的机制共同回答一个问题")][0]
caption = [p for p in doc.paragraphs if ptxt(p).startswith("表 7 三维决策示例")][0]
closing._element.getparent().remove(closing._element)
caption._element.addprevious(closing._element)

# caption alignment
fixes = {
    "图 3 ": "图 3 主要产品在任务生命周期覆盖程度与执行治理深度上的定位",
    "图 7 ": "图 7 第一年试点容量与验证容量测算",
}
for cap in list(doc.paragraphs):
    t = ptxt(cap)
    for pref, new in fixes.items():
        if t.startswith(pref):
            runs = cap.runs
            runs[0].text = new
            for r in runs[1:]:
                r.text = ""
doc.save(SRC)

d = Document(SRC)
for i, p in enumerate(d.paragraphs):
    t = ptxt(p)
    if t.startswith(("四个" ,"六个字段", "字段信度", "定义之后", "IC 自身", "本节的机制", "表 7", "4.1 ", "4.2 ", "4.3 ", "图 3", "图 7")):
        print(f"P{i:03d} [{p.style.name}] {t[:100]}")
print("paras", len(d.paragraphs))