# -*- coding: utf-8 -*-
"""Compare alignment state across every candidate plan file the user might open."""
import glob, os
from docx import Document
from docx.oxml.ns import qn

CAND = [
    r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx",
    r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v6.docx",
    r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v5.docx",
]
CAND += glob.glob(r"T:\创新创业\OwO-master\docs\.qa_v5\v8_before_*.docx")

for src in CAND:
    if not os.path.exists(src):
        continue
    doc = Document(src)
    total_par = 0
    centered = 0
    other = 0
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    total_par += 1
                    pPr = p._p.find(qn('w:pPr'))
                    jc = None
                    if pPr is not None:
                        j = pPr.find(qn('w:jc'))
                        if j is not None:
                            jc = j.get(qn('w:val'))
                    if jc == 'center':
                        centered += 1
                    else:
                        other += 1
    has_t3 = any("意图胶囊的字段定义" in p.text for p in doc.paragraphs)
    mt = os.path.getmtime(src)
    print("%-30s tables=%2d cellparas=%3d CENTER=%3d other=%3d  has表3=%s  mtime=%s" % (
        os.path.basename(src)[-32:], len(doc.tables), total_par, centered, other, has_t3,
        __import__("datetime").datetime.fromtimestamp(mt).strftime("%m-%d %H:%M")))
