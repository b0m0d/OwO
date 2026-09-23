# -*- coding: utf-8 -*-
"""第六轮：术语收尾与交叉引用校验。"""
import docx
from docx.oxml.ns import qn

DOC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v6.docx"
doc = docx.Document(DOC)
body = doc.element.body


def txt_of(el):
    return ''.join(n.text or '' for n in el.iter(qn('w:t'))).strip()


def set_p(el, text):
    rs = el.findall(qn('w:r'))
    t0 = rs[0].findall(qn('w:t'))
    t0[0].text = text
    for t in t0[1:]:
        t.text = ""
    for r in rs[1:]:
        for t in r.findall(qn('w:t')):
            t.text = ""


# 段落级替换
para_fix = [
    ("Cuttle 的差异在于把输入焦点、持续情境、分级执行、成果验证与返回原处组成一条系统链路。",
     "Cuttle 的差异在于把输入焦点、持续情境、分级执行、成果验证与返回原处组成一条系统链路。"),
]
for el in body.iterchildren():
    if not el.tag.endswith('}p'):
        continue
    t = txt_of(el)
    if t.startswith("产品保持四级渐进交互。"):
        set_p(el, "产品保持四级渐进交互。L0 只生成可编辑文本；L1 调用范围明确、结果可预览的工具；"
                  "L2 进入单 Agent 工作流并显示计划、权限和验收条件；L3 仅在并行收益经判据确认后启用"
                  "多智能体协作，或在数据位置受限时进入跨设备扩展。层级越高，系统展示越多的计划、成本、风险和审批信息，"
                  "而层级本身由 4.2 节的联合决策给出，不由用户手工选择。")
    if t.startswith("其可验证指标不是"):
        set_p(el, "其可验证指标不是“是否感知到了更多信息”，而是路由 Macro-F1、用户纠正率、"
                  "本可避免的重复输入次数，以及情境过度读取率（Context Overreach Rate）。"
                  "前两项说明胶囊是否被正确理解，后两项说明胶囊是否读得过多。")

# 表格单元格替换
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                if "用情境胶囊减少重复说明" in p.text:
                    for r in p.runs:
                        if "情境胶囊" in r.text:
                            r.text = r.text.replace("情境胶囊", "意图胶囊")
                if p.text.strip() == "过度升级率低于 10%":
                    for r in p.runs:
                        r.text = r.text.replace("过度升级率低于 10%", "自治过度升级率低于 10%")

doc.save(DOC)

# ---- 交叉引用校验
issues = []
refs = []
for el in body.iterchildren():
    if el.tag.endswith('}p'):
        t = txt_of(el)
        for m in __import__('re').finditer(r'([0-9]\.[0-9])\s?节', t):
            refs.append((m.group(1), t[:40]))
        if '式(' in t:
            for m in __import__('re').finditer(r'式\((\d-\d)\)', t):
                refs.append(("式" + m.group(1), t[:40]))
print("交叉引用:")
for r in refs:
    print("  ", r)
doc2 = docx.Document(DOC)
alltext = "\n".join(p.text for p in doc2.paragraphs)
for t in doc2.tables:
    for row in t.rows:
        for c in row.cells:
            alltext += "\n" + c.text
for bad in ["情境胶囊", "技能市场", "交易分成 15%", "1740", "47.4", "8% 至 12%", "创新五", "设备网格"]:
    if bad in alltext:
        print("残留:", bad)
print("done")
