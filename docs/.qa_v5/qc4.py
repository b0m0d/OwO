# -*- coding: utf-8 -*-
"""Step 4: final layout balance + TOC."""
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

# 1) 表 6 另起一页，避免与图 7 同页挤压
for p in doc.paragraphs:
    if p.text.strip() == "表 6 首年验证容量与市场空间口径":
        p.paragraph_format.page_break_before = True
        p.paragraph_format.keep_with_next = True
        print("  表 6 page break set")

# 2) 图 7 放大到页宽约 80%
NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
target = Cm(15.0)
for p in doc.paragraphs:
    if p.text.strip().startswith("图 7 "):
        prev = p._element.getprevious()
        while prev is not None:
            ext = prev.findall('.//{%s}extent' % NS_WP)
            if ext:
                cx = int(ext[0].get("cx")); cy = int(ext[0].get("cy"))
                ratio = cy / float(cx)
                ext[0].set("cx", str(int(target))); ext[0].set("cy", str(int(int(target) * ratio)))
                for el in prev.findall('.//{%s}ext' % NS_A):
                    el.set("cx", str(int(target))); el.set("cy", str(int(int(target) * ratio)))
                print("  图 7 enlarged")
                break
            prev = prev.getprevious()

# 3) 更新目录
PAGES = {"项目摘要":"3","第一章":"4","第二章":"7","第三章":"10","第四章":"13","第五章":"18",
         "第六章":"20","第七章":"23","第八章":"25","第九章":"28","第十章":"30",
         "第十一章":"33","第十二章":"34","参考资料":"35"}
def set_cell(cell, text):
    ps = cell.paragraphs
    runs = ps[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""; r._element.getparent().remove(r._element)
    else:
        ps[0].add_run(text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)
for row in doc.tables[1].rows:
    left = row.cells[0].text.strip()
    for k, v in PAGES.items():
        if left.startswith(k):
            set_cell(row.cells[1], v); break

doc.save(SRC)
print("step 4 done")