# -*- coding: utf-8 -*-
"""Recompute and write the year-one-capped three-year table (80/120/180 in year one)."""
from docx import Document

P, T, D = 240.0, 40000.0, 180000.0
IND  = {"保守": [80, 600, 1300],  "基准": [120, 1200, 2800], "进取": [180, 2000, 5000]}
TEAM = {"保守": [0, 3, 8],        "基准": [0, 5, 24],        "进取": [0, 10, 40]}
DEP  = {"保守": [0, 0, 0],        "基准": [0, 0, 2],         "进取": [0, 1, 4]}
COST = {"保守": [40, 55, 80],     "基准": [50, 105, 190],    "进取": [65, 190, 260]}
KEYS = ("保守", "基准", "进取")
def w(x): return round(x/1e4, 1)
REV = {k: [w(IND[k][y]*P + TEAM[k][y]*T + DEP[k][y]*D) for y in range(3)] for k in KEYS}
RES = {k: [round(REV[k][y]-COST[k][y], 1) for y in range(3)] for k in KEYS}
for k in KEYS:
    print(k, "ind", IND[k], "team", TEAM[k], "dep", DEP[k], "rev", REV[k], "cost", COST[k],
          "res", RES[k], "cum2", round(RES[k][0]+RES[k][1], 1))

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)
def set_cell(cell, text):
    ps = cell.paragraphs
    runs = ps[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""; r._element.getparent().remove(r._element)
    else:
        ps[0].add_run(text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

t = doc.tables[20]
BREAK = {
 "营业收入 保守": "个人 %s 万元；团队 %s 万元；部署 0 / 0 / 0" % (
     " / ".join(str(w(IND['保守'][y]*P)) for y in range(3)),
     " / ".join(str(w(TEAM['保守'][y]*T)) for y in range(3))),
 "营业收入 基准": "个人 %s 万元；团队 %s 万元；部署 %s 万元" % (
     " / ".join(str(w(IND['基准'][y]*P)) for y in range(3)),
     " / ".join(str(w(TEAM['基准'][y]*T)) for y in range(3)),
     " / ".join(str(w(DEP['基准'][y]*D)) for y in range(3))),
 "营业收入 进取": "个人 %s 万元；团队 %s 万元；部署 %s 万元" % (
     " / ".join(str(w(IND['进取'][y]*P)) for y in range(3)),
     " / ".join(str(w(TEAM['进取'][y]*T)) for y in range(3)),
     " / ".join(str(w(DEP['进取'][y]*D)) for y in range(3))),
}
NOTE = {"保守": "三年均未转正；该情景用于观察下限", "基准": "第三年进入盈亏平衡上方", "进取": "第三年转正"}

for r in t.rows:
    lab = r.cells[0].text.strip()
    if lab.startswith("个人专业版"):
        for y in range(3):
            set_cell(r.cells[1+y], " / ".join(str(IND[k][y]) for k in KEYS))
    elif lab.startswith("团队许可"):
        for y in range(3):
            set_cell(r.cells[1+y], " / ".join(str(TEAM[k][y]) for k in KEYS))
    elif lab.startswith("私有部署"):
        for y in range(3):
            set_cell(r.cells[1+y], " / ".join(str(DEP[k][y]) for k in KEYS))
    elif lab.startswith("第一年规模约束"):
        set_cell(r.cells[1], "个人付费用户上限 200 人；本表按 80 / 120 / 180 人测算")
        set_cell(r.cells[4], "来自表 11 的验证容量；三种情景均在 200 人上限之内")
    elif lab.startswith("营业收入 "):
        k = lab.split(" ")[1]
        for y in range(3):
            set_cell(r.cells[1+y], f"{REV[k][y]} 万元")
        set_cell(r.cells[4], BREAK[lab])
    elif lab.startswith("经营成本"):
        for y in range(3):
            set_cell(r.cells[1+y], " / ".join(str(COST[k][y]) for k in KEYS) + " 万元")
    elif lab.startswith("经营结果 "):
        k = lab.split(" ")[1]
        for y in range(3):
            v = RES[k][y]
            set_cell(r.cells[1+y], (f"负 {abs(v)} 万元" if v < 0 else f"正 {v} 万元"))
        set_cell(r.cells[4], NOTE[k])
doc.save(SRC)

# 资金段同步
for x in doc.paragraphs:
    if x.text.strip().startswith("项目资源需求分三层表述"):
        cum_cons = round(RES["保守"][0] + RES["保守"][1], 1)
        cum_base = round(RES["基准"][0] + RES["基准"][1], 1)
        runs = x.runs
        runs[0].text = ("项目资源需求分三层表述。已有资源：学校实验室与实验工位、团队成员个人设备、"
                        "指导教师的方法与安全指导、开源模型与本地推理能力，这些不计入现金需求，"
                        "因此资金来源中团队自筹与学校支持的占比高于对外融资。新增现金需求分两批："
                        "第一年 50 万元，用于把原型推进到可验证的试点闭环并完成对照、消融与红队测试；"
                        f"第二年追加 110 万元，用于场景复制、团队版交付与兼容矩阵扩展。首轮对外表述的资源配置总额为 160 万元，"
                        f"按决策门分两批释放，任一批未达标即暂停后续投入；保守情景下前两年累计亏损约 {abs(cum_cons)} 万元，"
                        f"基准情景约 {abs(cum_base)} 万元，该额度可覆盖基准情景所需并保留安全边际。")
        for r in runs[1:]:
            r.text = ""; r._element.getparent().remove(r._element)
doc.save(SRC)

d = Document(SRC)
for ri, r in enumerate(d.tables[20].rows):
    print(f"R{ri}: " + " | ".join(c.text.strip()[:64] for c in r.cells))