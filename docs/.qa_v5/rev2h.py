# -*- coding: utf-8 -*-
from docx import Document
SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)
def set_cell(cell, text):
    ps = cell.paragraphs
    runs = ps[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""; r._element.getparent().remove(r._element)
    else:
        ps[0].add_run(text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)
def set_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""; r._element.getparent().remove(r._element)

t = doc.tables[20]
for r in t.rows:
    if r.cells[0].text.strip().startswith("第一年规模约束"):
        set_cell(r.cells[1], "深度服务与完整观察约 200 人")
        set_cell(r.cells[4], "第一年在付用户是服务范围上限，深度服务与完整观察按约 200 人配置，两个口径用途不同")
# 10.1 段落补一句说明两个口径
for x in doc.paragraphs:
    if x.text.strip().startswith("本章用于检验商业模型能否形成可持续经营"):
        set_text(x, "本章用于检验商业模型能否形成可持续经营，不代表项目已经取得收入、融资或用户规模。"
                    "测算采用统一口径，三种情景分别设定用户与客户数量，再按同一单价假设计算收入："
                    "个人专业版按年末在付用户数乘年费 240 元确认收入；团队许可按 4 万元每年计价，自第二年起计入；"
                    "私有部署按 18 万元每项目、以验收确认收入；每年在付用户由上年在付规模的 60% 续订加当年新增构成。"
                    "需要区分两个口径：第一年的在付用户是服务范围上限，而深度服务与完整观察的规模按第六章的验证容量"
                    "约 200 人配置，两者的用途不同，前者用于测算收入，后者用于配置试点人力与实验样本。"
                    "全部单价与转化率均属待验证假设，验证方式与门槛见表 14。")
doc.save(SRC)
d = Document(SRC)
for r in d.tables[20].rows:
    if r.cells[0].text.strip().startswith(("第一年规模约束", "个人专业版")):
        print(" | ".join(c.text.strip()[:70] for c in r.cells))