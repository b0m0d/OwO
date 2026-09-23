# -*- coding: utf-8 -*-
"""Dump the exact text of every sentence the reviewer flagged, plus frequency counts of
the AI-ish rhetorical patterns and the list of per-chapter conclusion boxes."""
import re
from docx import Document
from docx.oxml.ns import qn

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def ptxt(p):
    return "".join(t.text or "" for t in p._element.iter(qn('w:t')))

KEYS = [
 "把用户在任意输入框中的自然语言", "项目的核心命题是", "这一思路把产品竞争",
 "判断项目是否成立的最终标准", "上述四条说明了", "项目要验证的正是",
 "因此项目把这条命题做成可验证的实验", "系统为每一次输入回答三个相互独立的问题",
 "任何一方都无法单独完成一次越权动作", "跨设备能力是意图协议", "市场问题正在从是否使用",
 "模型能力越强", "这些产品的能力边界正在快速移动", "生态与信任",
 "首用价值", "留存诊断", "收入增长来自三条", "第二种增长", "本章用于检验商业模型",
 "使用户把时间用于理解", "12.3", "项目的衡量标准不是调用了多少",
 "IC 的评价对象是情境质量", "两者作用于不同层次", "指导教师不列入项目团队成员名单",
 "第二年的增长", "这类语言本质", "收入来自个人订阅",
]
print("=== 目标段落 ===")
for i, p in enumerate(doc.paragraphs):
    t = ptxt(p).strip()
    if any(k in t for k in KEYS):
        print(f"P{i:03d} [{p.style.name}] {t}")
        print()

print("=== 结论框清单 ===")
for i, p in enumerate(doc.paragraphs):
    t = ptxt(p).strip()
    if t.startswith("核心结论") or t.startswith("核心创新："):
        print(f"P{i:03d}: {t[:70]}")

print()
print("=== 句式频次 ===")
full = "\n".join(ptxt(p) for p in doc.paragraphs)
for pat, name in [
    (r"不是[^，。；]{0,24}，而是", "不是A而是B"),
    (r"而非", "而非"),
    (r"从[^，。；]{0,14}转向", "从A转向B"),
    (r"越[^，。；]{0,8}越", "越A越B"),
    (r"上述", "上述"),
    (r"这一[^，。；]{0,6}(思路|路径|原则|设计|判断|差别|命题|方式)", "这一X"),
    (r"项目要验证|项目把这条命题|本章用于|该原则用于|这也是.{0,12}的原因", "作者旁白"),
    (r"其中[一二三四五六]，", "其一其二"),
]:
    n = len(re.findall(pat, full))
    print(f"  {name}: {n}")
