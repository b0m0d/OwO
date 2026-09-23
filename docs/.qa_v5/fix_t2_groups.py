# -*- coding: utf-8 -*-
"""表 2: move the group labels out of the cell text into their own narrow column,
so every cell reads cleanly and headers stay visually centered."""
import glob, os, copy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCS = os.path.join(r"T:\创新创业\OwO-master\docs")
SRC = [x for x in glob.glob(os.path.join(DOCS, "Cuttle*v8.docx")) if "预览" not in x][0]
doc = Document(SRC)

def first_cells(t):
    return [c.text.strip() for c in t.rows[0].cells]

def set_cell(cell, text):
    ps = cell.paragraphs
    runs = ps[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
            r._element.getparent().remove(r._element)
    else:
        ps[0].add_run(text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

t = None
for cand in doc.tables:
    if first_cells(cand)[:2] == ["验证方向（分组）", "验证方法"]:
        t = cand
        break
assert t is not None, "表 2 not found"

GROUPS = {
 "入口机制": "用户价值", "意图胶囊": "核心机制", "三维任务决策": "核心机制",
 "可见审批": "安全可靠", "端到端任务": "安全可靠",
 "团队与机构": "商业验证", "纵向留存": "商业验证",
}
# 清掉单元格里的【…】前缀，记录分组
n = 0
for row in t.rows[1:]:
    txt = row.cells[0].text.strip()
    grp = ""
    for k, v in GROUPS.items():
        if f"【{v}】" in txt:
            grp = v
            txt = txt.replace(f"【{v}】", "").strip()
            break
    set_cell(row.cells[0], txt)
    row._element.set(qn('w:rsidTr'), grp)   # 暂存分组，稍后建列时读取
    n += 1
print("prefixes removed from", n, "rows")

# 为每行在最左插入一列“分组”
NCOL = 5
# 重建 tblGrid：原 4 列 -> 5 列
grid = t._element.find(qn('w:tblGrid'))
cols = [int(g.get(qn('w:w'))) for g in grid.findall(qn('w:gridCol'))]
new_widths = [900, cols[0] - 900 + 300, cols[1] - 300, cols[2], cols[3]]
total = sum(new_widths)
# 保证合计 9638
new_widths[-1] += 9638 - total
newgrid = OxmlElement('w:tblGrid')
for w in new_widths:
    gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(w)); newgrid.append(gc)
t._element.replace(grid, newgrid)

for row in t.rows:
    tr = row._element
    tcs = tr.findall(qn('w:tc'))
    src = tcs[0]
    newtc = copy.deepcopy(src)
    # 清空内容
    for child in list(newtc):
        if child.tag != qn('w:tcPr'):
            newtc.remove(child)
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); pPr.append(jc)
    ind = OxmlElement('w:ind'); ind.set(qn('w:firstLine'), '0'); pPr.append(ind)
    p.append(pPr)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rf.set(qn(a), 'Microsoft YaHei')
    rPr.append(rf)
    if row is t.rows[0]:
        rPr.append(OxmlElement('w:b'))
        rPr.append(OxmlElement('w:color')) if False else None
        c = OxmlElement('w:color'); c.set(qn('w:val'), 'FFFFFF'); rPr.append(c)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '18'); rPr.append(sz)
    r.append(rPr)
    tt = OxmlElement('w:t')
    if row is t.rows[0]:
        tt.text = "分组"
    else:
        tt.text = tr.get(qn('w:rsidTr')) or ""
    r.append(tt)
    p.append(r)
    newtc.append(p)
    # 宽度
    pr = newtc.find(qn('w:tcPr'))
    if pr is not None:
        tcW = pr.find(qn('w:tcW'))
        if tcW is None:
            tcW = OxmlElement('w:tcW'); pr.insert(0, tcW)
        tcW.set(qn('w:type'), 'dxa'); tcW.set(qn('w:w'), str(new_widths[0]))
    tr.insert(list(tr).index(src), newtc)
    tr.attrib.pop(qn('w:rsidTr'), None)

# 重新分配其余列宽
for row in t.rows:
    tcs = [tc for tc in row._element.findall(qn('w:tc'))]
    if len(tcs) != 5:
        continue
    for i, tc in enumerate(tcs):
        pr = tc.find(qn('w:tcPr'))
        if pr is None:
            pr = OxmlElement('w:tcPr'); tc.insert(0, pr)
        tcW = pr.find(qn('w:tcW'))
        if tcW is None:
            tcW = OxmlElement('w:tcW'); pr.insert(0, tcW)
        tcW.set(qn('w:type'), 'dxa'); tcW.set(qn('w:w'), str(new_widths[i]))

doc.save(SRC)
print("表 2 restructured to 5 columns; widths:", new_widths)
d2 = Document(SRC)
for cand in d2.tables:
    if [c.text.strip() for c in cand.rows[0].cells][:2] == ["分组", "验证方向"]:
        for r in cand.rows:
            print("   ", " | ".join(c.text.strip()[:26] for c in r.cells))
        break
