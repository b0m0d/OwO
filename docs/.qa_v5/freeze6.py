# -*- coding: utf-8 -*-
from docx import Document
SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)
def set_text(x, text):
    runs = x.runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""; r._element.getparent().remove(r._element)
def set_cell(cell, text):
    ps = cell.paragraphs
    set_text(ps[0], text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

for x in doc.paragraphs:
    t = x.text.strip()
    if t.startswith("说明：本计划书以一条核心命题为主线"):
        set_text(x, "说明：本计划书以一条核心命题为主线——AI 的执行能力持续增强，而执行所依赖的意图、情境与控制权仍散落在手工环节。"
                    "全部内容收敛为两个核心创新：意图胶囊 IC，以及从输入意图到行动的受控自治升级与返回原处闭环；"
                    "版本化成果契约与返回原处、策略、验证、恢复体系作为支撑其落地的工程体系；"
                    "自适应组队定位为执行效率优化策略，跨设备能力与技能生态属于远期扩展。")
    if t.startswith("除通用安全目标外，项目把两个与输入法入口强相关"):
        set_text(x, "除通用安全目标外，项目把两个与输入法入口强相关的风险作为专有指标管理，"
                    "并把“禁止访问”与“读取过多”拆成两个不同指标：禁止字段访问率针对明确标识的密码框、"
                    "支付确认与金融敏感字段，属设计红线，只能是 0；情境过度读取率针对普通工作情境中读取了非必需数据，"
                    "属统计指标，目标低于 5%。两者的判定依据、标注规则与报告条件不同，不能合并为一个数字。")

set_cell(doc.tables[4].rows[2].cells[0], "数据与研发人员（副链）")
set_cell(doc.tables[9].rows[5].cells[3], "越权拦截率、误审批率、禁止字段访问率")
set_cell(doc.tables[11].rows[2].cells[0], "禁止字段")
set_cell(doc.tables[17].rows[7].cells[1],
         "高风险动作未授权执行次数、禁止字段访问率、越权拦截率、可逆任务恢复率")
set_cell(doc.tables[17].rows[7].cells[2],
         "未授权执行 0 次；禁止字段访问率 0；越权拦截率 100%；恢复率不低于 90%")
set_cell(doc.tables[22].rows[7].cells[4], "禁止字段硬屏蔽 最小读取 独立审计")
doc.save(SRC)

d = Document(SRC)
full = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
for b in ["受控行动的受控自治", "副场景）", "禁止访问字段访问率"]:
    print(("STILL PRESENT: " if b in full else "cleared: ") + b)
print("禁止字段访问率 present:", "禁止字段访问率" in full)