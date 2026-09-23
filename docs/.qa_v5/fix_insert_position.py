# -*- coding: utf-8 -*-
"""修正：把误插到表格内部的段落移到表格之后（找到 <w:tbl> 并 addnext）。"""
import copy
import re

import docx
from docx.oxml.ns import qn

DOC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v7.docx"
doc = docx.Document(DOC)
body = doc.element.body

MARKERS = ["ZZMARKER", "对照实验的公平性前提必须写清", "全部指标共用一套统计口径", "两项专有指标的判定规则必须可复现"]


def el_text(el):
    return ''.join(n.text or '' for n in el.iter(qn('w:t')))


# 1) 删除误插的段落
removed = 0
for el in list(body.iterdescendants()):
    if el.tag.endswith('}p'):
        t = el_text(el)
        if any(m in t for m in MARKERS):
            parent = el.getparent()
            if parent is not None and parent.tag.endswith('}tbl'):
                parent.remove(el)
                removed += 1
print("删除误插在表格内部的段落:", removed)

# 2) 定位需要挂载的表格元素
def outer_table(prefix_cells):
    for el in body.iterchildren():
        if el.tag.endswith('}tbl'):
            txt = el_text(el)
            if all(p in txt for p in prefix_cells):
                return el
    raise KeyError(prefix_cells)


targets = [
    (["实验", "对照组", "处理组", "主指标"], "对照实验的公平性前提必须写清",
     "对照实验的公平性前提必须写清：处理组与对照组使用同一模型与同一版本、相同的工具与权限范围、"
     "相同的任务材料与结果要求，唯一变量是“任务如何进入系统、上下文如何获得”。"
     "若对照组的工具权限被削减或模型版本更低，效率提升就不能归因于入口设计。"
     "实验采用 within-subject crossover：同一用户在两轮中分别先用传统工作台与先用 Cuttle，"
     "以消除学习顺序与熟练度带来的偏差。"),
    (["维度", "指标", "十二个月目标", "数据来源"], "全部指标共用一套统计口径",
     "全部指标共用一套统计口径：以配对样本比较为主，报告效应量与置信区间，而不只报告均值差异；"
     "样本量在正式实验前依据预实验效应量做统计功效分析后确定；所有指标注明数据来源、标注规则与样本量。"
     "指标之间的解释关系是：原地闭环完成率是首要指标，其余指标用于解释它为什么高或低。"),
    (["安全目标", "指标", "阶段门槛", "不达标处理"], "两项专有指标的判定规则必须可复现",
     "两项专有指标的判定规则必须可复现：依据预先文档化的“完成任务最小充分信息集”与“最小必要执行档位”，"
     "由两名不参与研发的独立标注者独立标注，计算标注者一致率，不一致样本由第三方裁决并记录理由；"
     "一致率达到预设水平后该指标才对外报告。"),
]

TEMPLATE_P = doc.paragraphs[10]
for cells, marker, text in targets:
    tbl = outer_table(cells)
    new_el = copy.deepcopy(TEMPLATE_P._element)
    tbl.addnext(new_el)
    p = docx.text.paragraph.Paragraph(new_el, doc)
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""
    assert marker in p.text
    print("   + 已挂到表后:", text[:32])

doc.save(DOC)

d2 = docx.Document(DOC)
txt = "\n".join(x.text for x in d2.paragraphs)
print("复核:")
for m in ["对照实验的公平性前提", "全部指标共用一套统计口径", "两项专有指标的判定规则必须可复现"]:
    print("  ", "OK " if m in txt else "MISS", m)
print("ZZMARKER 残留:", "ZZMARKER" in txt)
print("段落:", len(d2.paragraphs), "表格:", len(d2.tables))
