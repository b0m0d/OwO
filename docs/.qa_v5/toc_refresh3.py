# -*- coding: utf-8 -*-
"""Refresh TOC cached page numbers from the latest render."""
import glob, os, re, copy
import pymupdf
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCS = os.path.join(r"T:\创新创业\OwO-master\docs")
SRC = [x for x in glob.glob(os.path.join(DOCS, "Cuttle*v8.docx")) if "预览" not in x][0]
PDF = sorted(glob.glob(os.path.join(DOCS, ".qa_v5", "pdfHH", "*.pdf")))[0]

doc = Document(SRC)
def visible(p):
    return "".join(t.text or "" for t in p._element.iter(qn('w:t')))

pdf = pymupdf.open(PDF)
pages = [pdf[i].get_text().replace("\n", "").replace(" ", "") for i in range(pdf.page_count)]

# 收集标题
heads = []
for p in doc.paragraphs:
    if p.style.name in ("Heading 1", "Heading 2"):
        t = visible(p).strip()
        if t and t != "目录":
            heads.append(t)

# 目录占 2 页
resolved = {}
for t in heads:
    key = t.replace(" ", "")
    for i in range(3, len(pages)):
        if key in pages[i]:
            resolved[t] = str(i + 1)
            break
missing = [t for t in heads if t not in resolved]
print("resolved:", len(resolved), "/", len(heads), "| missing:", missing[:5])

updated = 0
for p in doc.paragraphs:
    el = p._element
    if not el.findall('.//' + qn('w:tab')):
        continue
    lbl = visible(p).strip()
    if not lbl:
        continue
    title = None
    for t in resolved:
        if lbl.startswith(t):
            title = t
            break
    if title is None:
        continue
    runs = p.runs
    if not runs:
        continue
    runs[0].text = title
    pg = resolved[title]
    pr = None
    for r in runs[1:]:
        if (r.text or "").strip().isdigit():
            pr = r
            break
    if pr is None:
        pr = p.add_run("")
    pr.text = pg
    for r in runs[1:]:
        if r is not pr:
            r.text = ""
    updated += 1
print("TOC entries updated:", updated)
doc.save(SRC)
