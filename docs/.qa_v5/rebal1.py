# -*- coding: utf-8 -*-
"""Rebalance batch 1: expand chapter 3 (product experience), compress chapter 4 (tech)."""
import copy
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""; r._element.getparent().remove(r._element)

def find(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)

def E(prefix, text):
    set_text(find(prefix), text)

def drop(prefix):
    p = find(prefix); p._element.getparent().remove(p._element)

def model_p():
    for p in doc.paragraphs:
        if p.style.name == "Normal" and len(p.text) > 60:
            return p
    return doc.paragraphs[10]

def insert_after(anchor, texts):
    prev = anchor
    for t in texts:
        el = copy.deepcopy(model_p()._element)
        prev._element.addnext(el)
        np = Paragraph(el, prev._parent)
        set_text(np, t)
        prev = np
    return prev

# ================= 第三章：补用户体验流程（+450 字） =================
p = find("案例 B：IDE 错误到可验证差异")
insert_after(p, [
 "一次完整任务包含六个步骤，用户在每个步骤都能看到系统正在做什么。"
 "以案例 A 为例：第一步，用户在 Word 的目标位置输入任务；"
 "第二步，系统识别任务边界并预览准备读取的资料范围，逐项列出网页、PDF、表格与文档来源；"
 "第三步，用户确认读取范围，或按应用与字段关闭其中任意一项；"
 "第四步，系统建立意图胶囊并执行取数与核验，过程中以状态条显示当前步骤、已用时间与预计剩余；"
 "第五步，遇到来源口径冲突或高风险动作时，系统暂停执行并请求确认，同时给出冲突双方的来源与差异；"
 "第六步，成果在需求产生的原位置生成，附带引用、修改痕迹与证据卡，用户可在原处直接审阅并继续编辑。",
 "流程中的失败处理同样可见。执行中断时系统给出明确原因，并支持从最近检查点重试；"
 "无法自动完成的部分保留已完成结果与未解决问题的清单，由用户接管剩余步骤，"
 "不会因为中途失败而丢弃已经产生的成果。每份成果都保留从原始对象到最终结果的完整链路，"
 "用户可以逐条查看某个数字来自哪个文件、哪一页，以及是否经过人工修改。",
])

# ================= 4.2 压缩：只留一个例子 =================
E("决策的输出是三维取值与策略引擎裁定的审批方式",
  "决策的输出是三维取值与策略引擎裁定的审批方式：无需确认、先预览、必须确认或直接拒绝。"
  "维度拆分的必要性可用一个例子说明：让手机读取一张照片，属于低自治、单执行者、跨设备，"
  "按单一层级会被高估为最高档；而多执行者阅读资料并只输出摘要的任务风险很低，"
  "其等级不应高于自动修改版本库的单执行者任务。")

# ================= 4.3 压缩组队论述 =================
E("自适应组队用于执行效率优化",
  "自适应组队用于执行效率优化。多执行者协作在多数任务上协调开销大于收益，因此系统默认单执行者，"
  "仅在预测收益为正时才考虑组队；判断依据是任务可并行性、结果可合并性与预计协调成本，"
  "并使用历史实测结果持续校准。若实测收益与事前判断长期不相关，项目将改为按任务模板静态指派。")
E("执行后的实测收益是真实可观测的量",
  "执行后的实测收益包括相对单执行者的质量变化、耗时变化、额外推理成本、结果冲突与合并失败次数，"
  "用于判定本次组队是否值得，并作为标签校准事前判断。")

# ================= 4.6 压缩公平性说明 =================
E("为保证三组对照的可比性",
  "为保证三组对照的可比性，各组使用相同的模型版本、工具权限、任务材料与结果要求，"
  "仅改变任务入口与情境获取机制；实验采用同组内交叉设计并对顺序做平衡处理，以消除学习效应。")

doc.save(SRC)
print("batch 1 done")