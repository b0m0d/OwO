# -*- coding: utf-8 -*-
"""Repair references [1]-[12]: their paragraphs live inside hyperlink wrappers, so the earlier
text replacement appended instead of replacing. Rebuild each paragraph cleanly."""
import copy, re, zipfile, os, shutil
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
shutil.copy2(SRC, r"T:\创新创业\OwO-master\docs\.qa_v5\v8_before_refrepair.docx")
doc = Document(SRC)

def ptxt(p):
    return "".join(t.text or "" for t in p._element.iter(qn('w:t')))

def clean_set(p, text):
    """Replace all content of a paragraph, keeping pPr and the first run's rPr."""
    el = p._element
    rpr = None
    for r in el.iter(qn('w:r')):
        rp = r.find(qn('w:rPr'))
        if rp is not None:
            rpr = copy.deepcopy(rp)
            break
    for child in list(el):
        if child.tag != qn('w:pPr'):
            el.remove(child)
    r = OxmlElement('w:r')
    if rpr is not None:
        r.append(rpr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    el.append(r)

fixed = 0
for p in doc.paragraphs:
    raw = ptxt(p).strip()
    if not raw.startswith("["):
        continue
    # 目标是：段落文本只出现一次编号
    nums = re.findall(r"^\[(\d+)\]", raw)
    markers = re.findall(r"\[\d+\]", raw)
    if len(markers) == 1:
        continue
    # 取最后一次出现的条目（新版格式）
    parts = re.split(r"(?=\[\d+\]\s)", raw)
    parts = [x.strip() for x in parts if x.strip()]
    if len(parts) >= 2:
        new_text = parts[-1]
        clean_set(p, new_text)
        fixed += 1
print("repaired reference paragraphs:", fixed)

doc.save(SRC)

# verify
d = Document(SRC)
bad = []
for p in d.paragraphs:
    raw = ptxt(p).strip()
    if raw.startswith("["):
        if len(re.findall(r"\[\d+\]", raw)) > 1:
            bad.append(raw[:80])
print("remaining duplicated refs:", len(bad))
for b in bad[:5]:
    print("   ", b)
z = zipfile.ZipFile(SRC)
print("zip ok:", z.testzip() is None)
