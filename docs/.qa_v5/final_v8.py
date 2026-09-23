# -*- coding: utf-8 -*-
import io, re, zipfile
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)
def ptxt(p):
    return "".join(t.text or "" for t in p._element.iter(qn('w:t')))
body = doc.element.body
out = []
pi = ti = 0
for ch in body.iterchildren():
    tag = ch.tag.split('}')[-1]
    if tag == 'p':
        p = Paragraph(ch, doc)
        out.append(f"[P{pi:04d}|{p.style.name}] {p.text.strip()}")
        pi += 1
    elif tag == 'tbl':
        t = Table(ch, doc)
        out.append(f"===== TABLE {ti} ({len(t.rows)}x{len(t.columns)}) =====")
        for ri, row in enumerate(t.rows):
            out.append(f"  T{ti}R{ri}: " + " | ".join(c.text.strip().replace('\n',' / ') for c in row.cells))
        out.append(f"===== END TABLE {ti} =====")
        ti += 1
io.open(r"T:\创新创业\OwO-master\docs\.qa_v5\v8_dump.txt","w",encoding="utf-8").write("\n".join(out))

issues = []
paras = doc.paragraphs
for p in paras:
    if p.style.name.startswith("Heading") and not p.text.strip():
        issues.append("EMPTY HEADING")
    t = p.text.strip()
    if t.endswith(("，", "、", "：", "；")) and len(t) > 25:
        issues.append("TRUNCATION?: " + t[:70])
caps_t = [p.text.strip() for p in paras if re.match(r"^表 \d+ ", p.text.strip())]
caps_f = [p.text.strip() for p in paras if re.match(r"^图 \d+ ", p.text.strip())]
tn = [int(re.match(r"^表 (\d+)", c).group(1)) for c in caps_t]
fn = [int(re.match(r"^图 (\d+)", c).group(1)) for c in caps_f]
if tn != list(range(1, len(tn)+1)): issues.append(f"TABLE SEQ {tn}")
if fn != list(range(1, len(fn)+1)): issues.append(f"FIGURE SEQ {fn}")
refs = [ptxt(p).strip() for p in paras if re.match(r"^\[\d+\]", ptxt(p).strip())]
if len(refs) != 30: issues.append(f"REF COUNT {len(refs)}")
# markers that must be present
must = ["两个核心创新", "意图胶囊（Intent Capsule, IC）", "高频", "全局快捷键", "三组对照实验",
        "统字符", "原地闭环完成率", "禁止访问字段", "情境过度读取", "定价访谈", "谈世钊", "吴栩彪",
        "同组内交叉设计", "统计功效分析", "本轮不下结论"]
full = "\n".join(ptxt(p) for p in paras) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
for mk in must:
    if mk not in full:
        issues.append("MISSING MARKER: " + mk)
banned = ["Swarm", "三个核心创新", "有统计意义", "L0 ", "L1 ", "L2 ", "L3 ", "比赛", "参赛主体",
          "六元组", "R Risk", "设备网格"]
for b in banned:
    if b in full:
        issues.append("BANNED PRESENT: " + b)
print("paras", pi, "tables", ti, "refs", len(refs))
print("ISSUES:", len(issues))
for i in issues: print("  !", i)