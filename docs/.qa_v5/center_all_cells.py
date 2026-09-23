# -*- coding: utf-8 -*-
"""Center every body cell in every table (header already centered), with uniform spacing."""
import glob, os
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCS = os.path.join(r"T:\创新创业\OwO-master\docs")
SRC = [x for x in glob.glob(os.path.join(DOCS, "Cuttle*v8.docx")) if "预览" not in x][0]
doc = Document(SRC)

n_cells = 0
n_paras = 0
for ti, t in enumerate(doc.tables):
    if ti <= 1:          # 封面表与目录不在处理范围
        continue
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.first_line_indent = 0
                pf.left_indent = 0
                pf.space_before = 0
                pf.space_after = 0
                n_paras += 1
            tcPr = c._tc.get_or_add_tcPr()
            va = tcPr.find(qn('w:vAlign'))
            if va is None:
                va = OxmlElement('w:vAlign'); tcPr.append(va)
            va.set(qn('w:val'), 'center')
            n_cells += 1

doc.save(SRC)
print("cells centered:", n_cells, "| paragraphs:", n_paras)
