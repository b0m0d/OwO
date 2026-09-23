# -*- coding: utf-8 -*-
"""Final audit after the de-AI pass: pattern frequency, conclusion boxes, typo check, integrity."""
import re, zipfile
from docx import Document
from docx.oxml.ns import qn

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def ptxt(p):
    return "".join(t.text or "" for t in p._element.iter(qn('w:t')))

paras = doc.paragraphs
full = "\n".join(ptxt(p) for p in paras)

print("=== 句式频次 ===")
for pat, name in [
    (r"不是[^，。；]{0,24}，而是", "不是A而是B"),
    (r"而非", "而非"),
    (r"从[^，。；]{0,14}转向", "从A转向B"),
    (r"上述", "上述"),
    (r"这一(思路|路径|原则|设计|判断|差别|命题|方式)", "这一X"),
    (r"项目要验证|项目把这条命题|本章用于|该原则用于|这也是.{0,14}的原因|项目的衡量标准", "作者旁白"),
]:
    print("  %-12s %d" % (name, len(re.findall(pat, full))))

print()
print("=== 结论框 ===")
boxes = [ptxt(p).strip() for p in paras if ptxt(p).strip().startswith(("核心结论", "核心创新："))]
for b in boxes:
    print("  ", b[:80])
print("  共", len(boxes), "个")

print()
print("=== 关键修正点核对 ===")
checks = {
 "封面新表述": "将用户在任意输入场景中的自然语言请求",
 "摘要新开头": "Cuttle 所解决的问题在于",
 "1.2 新表述": "Cuttle 不以基础模型能力作为主要竞争点",
 "1.3 直接比较": "与全局快捷键加悬浮窗口相比",
 "4.2 非拟人": "任务路由由自治程度、协作形态和执行位置三个相互独立的维度构成",
 "4.2 去口号": "职责分离降低了单一模块自行扩大权限",
 "4.5 去自然扩展": "可进一步扩展至跨设备执行",
 "6.1 去宏大论断": "跨应用调用、权限控制和结果追溯等问题的重要性随之提升",
 "6.3 去咨询腔": "现有产品分别强化输入表达或任务执行",
 "6.4 去套话": "兼容性测试、安全规则与实施经验可持续积累",
 "10.1 去不是A而是B": "收入增长主要由个人用户新增、团队转化和开发者渠道三部分构成",
 "10.1 直接进入": "财务测算采用个人订阅、团队许可和私有部署三类收入",
 "typo 修正": "第二年增长主要来自前两项",
 "教师免责句已删": None,
 "删除模型能力越强越需要": None,
}
for name, key in checks.items():
    if key is None:
        continue
    print("  %-22s %s" % (name, "OK" if key in full else "MISSING"))

print("  教师免责句: ", "仍存在" if "不列入项目团队成员名单" in full else "OK 已删")
print("  '模型能力越强，用户越需要': ", "仍存在" if "模型能力越强，用户越需要" in full else "OK 已删")
print("  '第二种增长': ", "仍存在" if "第二种增长" in full else "OK 已删")

print()
z = zipfile.ZipFile(SRC)
print("zip ok:", z.testzip() is None, "| paragraphs:", len(paras), "| tables:", len(doc.tables))
dx = z.read("word/document.xml").decode("utf-8")
print("TOC field:", "TOC " in dx, "| fldChar:", dx.count("fldChar"))
