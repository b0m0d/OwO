# -*- coding: utf-8 -*-
import re, zipfile
from docx import Document
SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
d = Document(SRC)
p = d.paragraphs
full = "\n".join(x.text for x in p) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
issues = []
for x in p:
    if x.style.name.startswith("Heading") and not x.text.strip(): issues.append("EMPTY HEADING")
    t = x.text.strip()
    if t.endswith(("，", "、", "：", "；")) and len(t) > 25: issues.append("TRUNCATION: " + t[:50])
tn = [int(re.match(r"^表 (\d+)", x.text.strip()).group(1)) for x in p if re.match(r"^表 \d+ ", x.text.strip())]
fn = [int(re.match(r"^图 (\d+)", x.text.strip()).group(1)) for x in p if re.match(r"^图 \d+ ", x.text.strip())]
if tn != list(range(1, len(tn)+1)): issues.append(f"TABLE SEQ {tn}")
if fn != list(range(1, len(fn)+1)): issues.append(f"FIGURE SEQ {fn}")
banned = ["Swarm","三个核心创新","有统计意义","R Risk","设备网格","连续自治","交易额分成 15","命中",
          "参赛主体","决定任务层级","作为所有 Agent 前的意图入口","20 万元每项目","9 万元实施与维护",
          "按席位","4 万元每年每客户","核心设计目标的不同","禁止读取率","受控行动的受控自治",
          "拿不到","副场景","禁止访问字段访问率","在付用户不超过 200 人"]
for b in banned:
    if b in full: issues.append("BANNED: " + b)
must = ["受控自治升级","适当信任","Task Contract","Artifact Contract","执行权限的上界","禁止字段访问率",
        "越权拦截率","学生主导、教师指导","知识工作的直接价值","可积累优势与拟形成壁垒","参考市场边界",
        "首批可触达验证样本池","三组入口对照","20 份定价访谈","个人付费用户上限 200 人","4 万元每团队每年",
        "18 万元每项目","稳定锚点","当前主要交互形态与价值重心","约束：Origin 绑定准确率不下降",
        "主链子任务","在权限、能力与数据位置的约束下组合","入口与情境获取的组合设计"]
for mk in must:
    if mk not in full: issues.append("MISSING: " + mk)
P,T,D = 240.0,40000.0,180000.0
IND=[[80,600,1300],[120,1200,2800],[180,2000,5000]]; TEAM=[[0,3,8],[0,5,24],[0,10,40]]; DEP=[[0,0,0],[0,0,2],[0,1,4]]
PUB={"保守":[1.9,26.4,63.2],"基准":[2.9,48.8,199.2],"进取":[4.3,106.0,352.0]}
for i,k in enumerate(("保守","基准","进取")):
    for y in range(3):
        rev = round((IND[i][y]*P+TEAM[i][y]*T+DEP[i][y]*D)/1e4,1)
        if abs(rev-PUB[k][y])>0.05: issues.append(f"ARITH {k} y{y} {rev} vs {PUB[k][y]}")
z = zipfile.ZipFile(SRC)
print("zip ok:", z.testzip() is None, "| paragraphs", len(p), "| tables", len(d.tables),
      "| table caps", len(tn), "| figure caps", len(fn), "| refs", sum(1 for x in p if re.match(r"^\[\d+\]", x.text.strip())))
print("ISSUES:", len(issues))
for i in issues: print("  !", i)