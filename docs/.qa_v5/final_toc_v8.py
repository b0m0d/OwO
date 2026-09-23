# -*- coding: utf-8 -*-
"""Rebuild TOC page numbers and align remaining figure captions."""
from docx import Document
from docx.oxml.ns import qn

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

PAGES = [
    ("项目摘要", "3"),
    ("第一章", "5"), ("第二章", "9"), ("第三章", "12"), ("第四章", "14"),
    ("第五章", "20"), ("第六章", "22"), ("第七章", "25"), ("第八章", "27"),
    ("第九章", "30"), ("第十章", "31"), ("第十一章", "34"), ("第十二章", "36"),
    ("参考资料", "37"),
]

toc = doc.tables[1]
print("toc rows:", len(toc.rows))
for row in toc.rows:
    left = row.cells[0].text.strip()
    for key, pg in PAGES:
        if left.startswith(key):
            cell = row.cells[1]
            ps = cell.paragraphs
            runs = ps[0].runs
            if runs:
                runs[0].text = pg
                for r in runs[1:]:
                    r.text = ""
                    r._element.getparent().remove(r._element)
            else:
                ps[0].add_run(pg)
            for extra in ps[1:]:
                extra._element.getparent().remove(extra._element)
            break

CAPS = {
    "图 3 ": "图 3 主要产品在任务生命周期覆盖程度与执行治理深度上的定位",
    "图 5 ": "图 5 三维决策：自治程度 · 协作形态 · 执行位置",
    "图 9 ": "图 9 Cuttle 从知识工作场景样板到团队许可与私有部署的市场进入路径",
    "图 11 ": "图 11 Cuttle 三年经营测算（三情景，统一口径）",
}
for p in list(doc.paragraphs):
    t = p.text.strip()
    for pref, new in CAPS.items():
        if t.startswith(pref):
            runs = p.runs
            runs[0].text = new
            for r in runs[1:]:
                r.text = ""
                r._element.getparent().remove(r._element)
doc.save(SRC)

d = Document(SRC)
print("--- TOC ---")
for row in d.tables[1].rows:
    print("  ", row.cells[0].text.strip(), "->", row.cells[1].text.strip())
print("--- captions ---")
for p in d.paragraphs:
    t = p.text.strip()
    if t.startswith(("图 3 ", "图 5 ", "图 9 ", "图 11 ")):
        print("  ", t)