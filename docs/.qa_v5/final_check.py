import docx, re
p=r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v6.docx"
d=docx.Document(p)
alltext="\n".join(x.text for x in d.paragraphs)
for t in d.tables:
    for r in t.rows:
        for c in r.cells: alltext+="\n"+c.text
bad = ["情境胶囊","技能市场","交易分成 15%","1740","47.4","8% 至 12%","创新五","创新四","设备网格","必须为 0（泛化）"]
print("残留检查:")
for b in bad:
    n=alltext.count(b)
    if n: print("  ",b,n)
comma = [x.text for x in d.paragraphs if "，" in x.text and "。" not in x.text and len(x.text)>30]
print("无句号长段落:", len(comma))
print("段落数", len(d.paragraphs), "表格数", len(d.tables))
words=sum(len(x.text) for x in d.paragraphs)
twords=sum(len(c.text) for t in d.tables for r in t.rows for c in r.cells)
print("正文字数(含表格重复)", words, twords)
