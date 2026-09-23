# -*- coding: utf-8 -*-
"""Final wording sweep: remove remaining student-first framing."""
from docx import Document

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def set_cell(cell, text):
    ps = cell.paragraphs
    runs = ps[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
            r._element.getparent().remove(r._element)
    else:
        ps[0].add_run(text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

def set_row(table, ri, values):
    for ci, v in enumerate(values):
        set_cell(table.rows[ri].cells[ci], v)

# 表 4 假设：校园团队 -> 团队与机构客户
set_row(doc.tables[5], 5, ["团队与机构客户愿意为治理能力付费", "试点与报价测试",
                           "至少 1 个付费意向或采购流程", "至少 3 个付费意向或采购流程",
                           "转向个人订阅或开发者生态"])

# 表 5 典型场景：课程报告 -> 跨来源资料汇总
set_row(doc.tables[6], 1, ["项目汇报 | 在 Word 输入 按这些资料更新汇报口径",
                           "识别文档结构与要求 调用资料检索 核验与写作流程",
                           "带来源标注与修改痕迹的 Word 内容"])

# 表 15 单位经济：校园团队版 -> 团队版
set_row(doc.tables[15], 2, ["团队版", "4 万元每团队", "1.2 万元交付支持", "70%", "使用深度 续约"])

# 表 16 路线：校园试点 -> 场景试点；校园团队版 -> 团队版
set_row(doc.tables[16], 3, ["7 至 12 个月",
                            "输入法兼容矩阵与 Origin/Return 适配器标准化",
                            "完成三组对照实验与消融实验 标定路由阈值",
                            "三个场景试点（办公文档 报告交付 代码修复）",
                            "连续四周留存与安全门槛同时达标"])
set_row(doc.tables[16], 4, ["13 至 18 个月", "团队版与成果契约接口",
                            "验证协同净收益与团队交付成本", "首批付费团队",
                            "续用意愿与可控交付成本"])

# 表 19 人才培养：校园试点 -> 场景试点
set_row(doc.tables[19], 4, ["市场验证", "商业模式 财务与合规", "场景试点 报价 合作复盘", "意向 反馈与经营口径"])

# 表 21 资源：校园复制 -> 场景复制
set_row(doc.tables[21], 3, ["规模化资源（第二年）",
                            "场景复制 团队版交付 兼容矩阵扩展 支持与运维",
                            "80 万元", "第二批试点 团队许可交付能力"])
doc.save(SRC)
print("ok")
d = Document(SRC)
for i, p in enumerate(d.paragraphs):
    if "校园" in p.text or "课程报告" in p.text:
        print(f"P{i:03d} {p.text[:110]}")