# -*- coding: utf-8 -*-
"""v7 题注终稿：按“表格表头”识别每张表，重建题注文本并放到表前，统一编号。"""
import re

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

DOC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v7.docx"
doc = docx.Document(DOC)
body = doc.element.body


def el_text(el):
    return ''.join(n.text or '' for n in el.iter(qn('w:t'))).strip()


def set_text(p_obj, text):
    runs = p_obj.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        p_obj.add_run(text)


# (表头前缀, 题注描述)；封面表与目录表无题注
SPEC = [
    ("项目要素", "项目评审证据链"),
    ("对象当前成本", "输入时刻入口价值与首批验证口径"),
    ("用户高频任务", "目标用户与基准任务链"),
    ("假设验证方法", "预设验证标准与证伪条件"),
    ("场景用户输入", "典型场景的输入到交付闭环"),
    ("六元组分量", "意图胶囊的字段定义"),
    ("任务特征", "三维决策示例 自治程度 协作形态 执行位置"),
    ("实验对照组", "核心机制的对照与消融设计"),
    ("风险层级", "风险分级与控制策略"),
    ("安全目标指标", "安全门槛 设计红线 统计指标与生产事故指标"),
    ("层级测算口径", "第一年试点容量测算"),
    ("类别代表产品", "任务生命周期六维度竞品定位"),
    ("产品目标用户", "产品版本与收入结构"),
    ("收入单元", "单位经济模型的规划假设"),
    ("阶段产品目标", "分阶段研发与市场决策门"),
    ("维度指标", "项目十二个月核心指标"),
    ("成员角色", "团队分工与贡献证据"),
    ("阶段学习任务", "项目驱动的人才培养路径"),
    ("指标第一年", "三年经营测算（统一口径）"),
    ("资源层级", "分层资源需求与用途"),
    ("风险概率", "项目风险登记表"),
]

# 1) 表格元素按文档顺序
tbls = [el for el in body.iterchildren() if el.tag.endswith('}tbl')]
body_tbls = []
for el in tbls:
    hdr = "".join(c.text for c in Table(el, doc).rows[0].cells)
    if hdr.startswith(("参赛组别", "项目摘要")):
        continue
    body_tbls.append((hdr, el))
print("正文表格:", len(body_tbls), "期望:", len(SPEC))
assert len(body_tbls) == len(SPEC), len(body_tbls)

# 2) 取出所有现有“表 ”题注段落（按文档顺序）
caps = [el for el in body.iterchildren()
        if el.tag.endswith('}p') and el_text(el).startswith('表 ')]
print("现有表题注:", len(caps))
extra = caps[len(SPEC):]
for el in extra:
    el.getparent().remove(el)
    print("  删除多余题注:", el_text(el)[:34])
caps = caps[:len(SPEC)]

# 3) 逐表：题注放到表前并重写文本
for i, ((hdr, tbl), cap) in enumerate(zip(body_tbls, caps)):
    if not hdr.startswith(SPEC[i][0]):
        print("  警告：第 %d 张表表头 %r 与期望前缀 %r 不符" % (i + 1, hdr[:20], SPEC[i][0]))
    prev = tbl.getprevious()
    if prev is not cap:
        cap.getparent().remove(cap)
        tbl.addprevious(cap)
    set_text(Paragraph(cap, doc), f"表 {i+1} {SPEC[i][1]}")

# 4) 图题注编号顺排
figs = [el for el in body.iterchildren()
        if el.tag.endswith('}p') and el_text(el).startswith('图 ')]
for i, el in enumerate(figs, start=1):
    set_text(Paragraph(el, doc), re.sub(r'^图\s*\d+', f'图 {i}', el_text(el)))

doc.save(DOC)
print("saved | paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))
print()
for i, ((hdr, tbl), cap) in enumerate(zip(body_tbls, caps)):
    print(f"  表{i+1}: {el_text(cap)[:34]:36s} | {hdr[:26]}")
