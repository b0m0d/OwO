# -*- coding: utf-8 -*-
"""Comprehensive final audit: duplication, captions, refs, citations, tone, fonts, TOC field."""
import re, zipfile
from docx import Document
from docx.oxml.ns import qn

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def ptxt(p):
    return "".join(t.text or "" for t in p._element.iter(qn('w:t')))

issues = []
paras = doc.paragraphs
cells = [c for t in doc.tables for r in t.rows for c in r.cells]
full = "\n".join(ptxt(p) for p in paras) + "\n" + "\n".join(c.text for c in cells)

# 1) 段落内重复文本（本轮踩过的坑）
for i, p in enumerate(paras):
    raw = ptxt(p)
    for m in re.finditer(r"(.{28,}?)\1", raw):
        issues.append(f"DUP P{i}: {m.group(1)[:50]}")
        break
    # 参考条目重复编号
    if raw.strip().startswith("[") and len(re.findall(r"\[\d+\]", raw)) > 1:
        issues.append(f"DUP REF P{i}: {raw[:60]}")

# 2) 题注连续性
tc = [int(re.match(r"^表 (\d+)", p.text.strip()).group(1)) for p in paras if re.match(r"^表 \d+ ", p.text.strip())]
fc = [int(re.match(r"^图 (\d+)", p.text.strip()).group(1)) for p in paras if re.match(r"^图 \d+ ", p.text.strip())]
if tc != list(range(1, len(tc) + 1)): issues.append(f"TABLE CAPTIONS {tc}")
if fc != list(range(1, len(fc) + 1)): issues.append(f"FIGURE CAPTIONS {fc}")

# 3) 悬空引用
refs = set(int(m.group(1)) for m in re.finditer(r"表 (\d+)", full))
dangling = sorted(r for r in refs if r > len(tc))
if dangling: issues.append(f"DANGLING TABLE REFS {dangling}")

# 4) 参考文献编号与引用
entries = {}
for p in paras:
    m = re.match(r"^\[(\d+)\]\s*(.+)$", ptxt(p).strip())
    if m: entries[int(m.group(1))] = m.group(2)
if sorted(entries) != list(range(1, len(entries) + 1)):
    issues.append(f"REF NUMBERING {sorted(entries)}")
body = "\n".join(ptxt(p) for p in paras if not re.match(r"^\[\d+\]", ptxt(p).strip()))
body += "\n" + "\n".join(c.text for c in cells)
cited = set(int(m.group(1)) for m in re.finditer(r"\[(\d+)\]", body))
uncited = sorted(set(entries) - cited)
if uncited != [1, 2]: issues.append(f"UNCITED REFS {uncited}")

# 5) 禁用/元叙事表述
banned = ["需要说明的是","本章不采用","这里需要区分","不作为独立创新点主张","因此本节标题","本计划书以一条核心命题",
          "必须写清","额外说明一点","全文出现的 IC","首轮对外表述","原因很简单","赛事整合","席位计价",
          "高等教育在学总规模","三维自治","连续自治","Swarm","有统计意义","命中","参赛主体","统一口径"]
for b in banned:
    if b in full: issues.append("BANNED: " + b)

# 5b) TOC 条目行不算中英文紧贴
toc_count = sum(1 for p in paras if re.match(r"^[^	]+	\d+$", ptxt(p)))

# 6) 字体
fonts = set()
for p in paras:
    for r in p.runs:
        rp = r._element.find(qn('w:rPr'))
        if rp is not None:
            f = rp.find(qn('w:rFonts'))
            if f is not None:
                fonts.add(f.get(qn('w:eastAsia')) or f.get(qn('w:ascii')))
if fonts - {"Microsoft YaHei"}: issues.append(f"UNEXPECTED FONTS {fonts}")

# 7) 中英文紧贴
tight = sum(len(re.findall(r"[\u4e00-\u9fa5][A-Za-z0-9]|[A-Za-z0-9][\u4e00-\u9fa5]", ptxt(p))) for p in paras)
if tight - toc_count > 0: issues.append(f"TIGHT CJK-LATIN SPACING {tight - toc_count}")

# 8) TOC 域
dx = zipfile.ZipFile(SRC).read("word/document.xml").decode("utf-8")
if "TOC " not in dx or dx.count("fldChar") < 4: issues.append("TOC FIELD MISSING")
if "updateFields" in zipfile.ZipFile(SRC).read("word/settings.xml").decode("utf-8"):
    issues.append("updateFields present (Word will prompt)")

# 9) 教师不得作为成员
team = None
for t in doc.tables:
    fc0 = [c.text.strip() for c in t.rows[0].cells]
    if fc0 and fc0[0] in ("学生团队成员", "成员"):
        team = t
if team is None:
    issues.append("TEAM TABLE MISSING")
else:
    names = [r.cells[0].text.strip() for r in team.rows[1:]]
    if any("赵" in n or "教师" in n for n in names):
        issues.append(f"ADVISOR IN MEMBER TABLE {names}")

print("paragraphs:", len(paras), "| tables:", len(doc.tables))
print("table captions:", len(tc), "| figure captions:", len(fc), "| refs:", len(entries), "| cited:", len(cited))
print("ISSUES:", len(issues))
for i in issues:
    print("   !", i)
