# -*- coding: utf-8 -*-
"""Editorial pass B5: unify the two user tables, renumber captions, fix cross-references."""
import copy
from docx import Document
from docx.text.paragraph import Paragraph

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""; r._element.getparent().remove(r._element)

def set_cell(cell, text):
    ps = cell.paragraphs
    set_text(ps[0], text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

def first_cells(t):
    return [c.text.strip() for c in t.rows[0].cells]

def get_table(sig, ncols=None):
    for t in doc.tables:
        fc = first_cells(t)
        if fc[:len(sig)] == sig and (ncols is None or len(fc) == ncols):
            return t
    raise KeyError(str(sig))

def drop_table(t):
    t._element.getparent().remove(t._element)

def drop_para(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            p._element.getparent().remove(p._element); return True
    return False

old = get_table(["主要任务", "现状与缺口"], 4)
drop_table(old)
t = get_table(["用户与岗位", "高频任务"], 4)
NEW = [
 ["用户与岗位", "高频任务", "现有障碍", "首个验证场景"],
 ["运营、项目、咨询、研究、行政、产品（主场景）", "跨来源取数、口径核验、报告与材料交付",
  "数据分散在网页、PDF、表格与聊天记录中，来源与口径难以复核",
  "跨来源资料到分析与 Word 报告（主链）"],
 ["数据与研发人员（副链）", "代码修复、测试与文档",
  "任务在 IDE、浏览器与终端之间切换，失败恢复成本高",
  "错误到修复、测试与可审阅差异（副链）"],
 ["高校学生与科研人员（早期样本）", "资料检索、综述与项目材料",
  "重复格式工作，审核链条长",
  "作为独立验证组的样本来源，不单独设定场景"],
 ["实验室与机构（后续扩展）", "数据整理、材料归档、多人协作",
  "模型与工具分散，缺少权限治理与审计",
  "本轮不纳入，作为团队许可与私有部署的需求来源"],
]
for ri, vals in enumerate(NEW):
    for ci, v in enumerate(vals):
        set_cell(t.rows[ri].cells[ci], v)

for c in ["表 1 项目评审证据链", "表 2 核心用户、任务与价值", "表 3 核心验证指标与决策门",
          "表 4 意图胶囊的字段定义", "表 5 三组对照实验与消融实验设计", "表 6 安全控制与发布门槛",
          "表 7 首年验证容量与市场空间口径", "表 8 产品竞争定位与任务生命周期覆盖",
          "表 9 产品版本、定价与单位经济", "表 10 团队分工与工程基础",
          "表 19 三年经营测算（三情景，各自设定用户与客户数量）", "表 11 项目风险登记表"]:
    drop_para(c)

def caption_before(sig, ncols, text):
    t = get_table(sig, ncols)
    model = None
    for p in doc.paragraphs:
        if p.style.name == "Normal" and len(p.text) > 60:
            model = p; break
    el = copy.deepcopy(model._element)
    t._element.addprevious(el)
    np = Paragraph(el, t._parent)
    set_text(np, text)

caption_before(["项目要素", "核心内容"], 3, "表 1 项目评审证据链")
caption_before(["用户与岗位", "高频任务"], 4, "表 2 核心用户、高频任务与验证场景")
caption_before(["验证方向", "验证方式"], 5, "表 3 核心验证指标与决策门")
caption_before(["IC 字段", "含义"], 5, "表 4 意图胶囊的字段定义")
caption_before(["实验", "对照组", "处理组"], 4, "表 5 三组对照实验与消融实验设计")
caption_before(["安全目标", "指标", "阶段门槛"], 4, "表 6 安全控制与发布门槛")
caption_before(["口径", "测算对象", "测算依据"], 4, "表 7 首年验证容量与市场空间口径")
caption_before(["类别", "代表产品", "强项"], 5, "表 8 产品竞争定位与任务生命周期覆盖")
caption_before(["产品版本", "目标用户", "规划定价"], 5, "表 9 产品版本、定价与单位经济")
caption_before(["指标", "第一年", "第二年", "第三年"], 5, "表 10 三年经营预测与资源需求")
caption_before(["成员", "角色与专业方向"], 4, "表 11 团队分工与工程基础")
caption_before(["风险", "概率", "影响", "预警信号"], 5, "表 12 项目风险登记表")

def fix_ref(old_s, new_s):
    for p in doc.paragraphs:
        if old_s in p.text:
            set_text(p, p.text.replace(old_s, new_s))
            print("  ref:", old_s[:22], "->", new_s[:22])

fix_ref("项目的价值主张与首批验证口径见表 2", "项目的价值主张与验证场景见表 2")
fix_ref("验证方式与门槛见表 14", "验证方式与门槛见表 9")
fix_ref("表 14 与本节后续说明", "表 9 与本节后续说明")
fix_ref("数量与单价口径见表 19", "数量与单价口径见表 10")
fix_ref("具体数值见表 19", "具体数值见表 10")

doc.save(SRC)
print("B5 done. tables:", len(doc.tables), "paras:", len(doc.paragraphs))