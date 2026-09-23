# -*- coding: utf-8 -*-
import re, zipfile
from docx import Document
SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
z = zipfile.ZipFile(SRC)
d = Document(SRC)
p = d.paragraphs
full = "\n".join(x.text for x in p) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
issues = []
tc = [int(re.match(r"^表 (\d+)", x.text.strip()).group(1)) for x in p if re.match(r"^表 \d+ ", x.text.strip())]
fc = [int(re.match(r"^图 (\d+)", x.text.strip()).group(1)) for x in p if re.match(r"^图 \d+ ", x.text.strip())]
if tc != list(range(1, len(tc)+1)): issues.append(f"TABLE CAPTIONS {tc}")
if fc != list(range(1, len(fc)+1)): issues.append(f"FIGURE CAPTIONS {fc}")
# dangling references
refs = set(int(m.group(1)) for m in re.finditer(r"表 (\d+)", full))
dangling = sorted(r for r in refs if r > len(tc))
if dangling: issues.append(f"DANGLING TABLE REFS {dangling}")
for x in p:
    if x.style.name.startswith("Heading") and not x.text.strip(): issues.append("EMPTY HEADING")
    t = x.text.strip()
    if t.endswith(("，", "、", "：", "；")) and len(t) > 25: issues.append("TRUNCATION: " + t[:60])
banned = ["需要说明的是","本章不采用","这里需要区分","不作为独立创新点主张","因此本节标题","本计划书以一条核心命题",
          "与三组对照实验一致","边界示例","仅作参照","计划形成","证伪","停止与转向条件","统一口径"]
for b in banned:
    if b in full: issues.append("BANNED: " + b)
print("zip ok:", z.testzip() is None)
print("paragraphs:", len(p), "| tables:", len(d.tables), "| table captions:", len(tc), "| figure captions:", len(fc))
print("ISSUES:", len(issues))
for i in issues: print("  !", i)