# -*- coding: utf-8 -*-
"""v8 revision round 2, part C: terminology, safety metrics, chapter 9 and 12."""
from docx import Document
from docx.oxml.ns import qn
import copy
from docx.text.paragraph import Paragraph

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""; r._element.getparent().remove(r._element)

def set_cell(cell, text):
    ps = cell.paragraphs
    set_text(ps[0], text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

def set_row(table, ri, values):
    for ci, v in enumerate(values):
        set_cell(table.rows[ri].cells[ci], v)

def find_p(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)

def E(prefix, text):
    p = find_p(prefix); set_text(p, text); return p

# ---------- 术语：连续自治 -> 受控自治升级 ----------
E("围绕这一命题，项目收敛出两个核心创新",
  "围绕这一命题，项目收敛出两个核心创新。核心创新一是意图胶囊（IC）：面向输入时刻、可校验、可失效、可追溯的最小情境结构，"
  "它把一次输入从一段自然语言变成有作用域、有生命周期、可供路由与执行共用的中间表示。"
  "核心创新二是受控自治升级闭环：由任务特征判定自治程度、协作形态与执行位置，"
  "把一次输入从文本表达升级为受控行动，并以 Return to Origin 把成果返回需求产生的原处。"
  "版本化成果契约、返回原处机制与策略、验证、恢复体系作为支撑上述机制落地的工程体系；"
  "自适应多智能体组队属于执行效率优化策略，不作为独立创新点主张；跨设备能力与技能生态为远期扩展，"
  "不是项目成立的必要条件。")
E("4.2 核心创新二 从输入意图到受控行动的连续自治闭环",
  "4.2 核心创新二 从输入意图到受控行动的受控自治升级闭环")
E("上述实现是团队工程能力的证据",
  "上述实现是团队工程能力的证据，不是本项目已经完成的产品，也不作为本项目的创新结论。"
  "项目主张的创新是意图胶囊与受控自治升级闭环两项机制；而运行时、权限、审计、回滚与多 Worker 等基础层"
  "恰好是这两项机制能够被真实执行而非仅停留在概念的前提。"
  "后续研发重点集中在输入入口、意图胶囊、返回原处、自治升级阈值与真实用户验证五件事上。")
E("知识产权围绕真正形成差异的机制布局",
  "知识产权围绕真正形成差异的机制布局，而不是围绕功能数量布局。近期完成客户端与运行时软件著作权登记，"
  "沉淀意图胶囊、受控自治升级判据与版本化成果契约的发明交底，建立商标与开源依赖清单；"
  "中期依据新颖性检索与实验结果决定专利申请，形成用户研究报告、评测数据集、技术白皮书与论文。"
  "所有成果明确学生贡献、数据授权与第三方许可，不以申请数量替代创新质量。")
E("说明：本计划书以一条核心命题为主线",
  "说明：本计划书以一条核心命题为主线——AI 的执行能力持续增强，而执行所依赖的意图、情境与控制权仍散落在手工环节。"
  "全部内容收敛为两个核心创新：意图胶囊 IC，以及从输入意图到受控行动的受控自治升级与返回原处闭环；"
  "版本化成果契约与返回原处、策略、验证、恢复体系作为支撑其落地的工程体系；"
  "自适应组队定位为执行效率优化策略，跨设备能力与技能生态属于远期扩展。")
t = doc.tables[2]
set_cell(t.rows[3].cells[1], "意图胶囊 IC；受控自治升级与返回原处闭环")

# ---------- 4.2 补一句“离散三档”的说明 ----------
anchor = find_p("当特征不确定性较高或动作不可逆时")
elx = copy.deepcopy(anchor._element)
anchor._element.addnext(elx)
npx = Paragraph(elx, anchor._parent)
set_text(npx, "需要说明“升级”的含义：自治程度本身是文本表达、工具调用、智能体执行三个离散档位，"
              "不是一个连续变化的数值；所谓升级是指任务可以从文本表达逐级上升为受控行动，"
              "而每一次上升都在同一任务生命周期内完成、可被审批、可被撤销。")

# ---------- 表 16 安全指标术语 ----------
t = doc.tables[17]
set_row(t, 7, ["安全",
               "高风险动作未授权执行次数、禁止访问字段访问率、越权拦截率、可逆任务恢复率",
               "未授权执行 0 次；禁止访问字段访问率 0；越权拦截率 100%；恢复率不低于 90%",
               "红队与故障注入"])

# ---------- 第九章：恢复“学生主导、教师指导”，并区分两个概念 ----------
E("项目采用成员主导、教师指导、用户共同验证的协作方式",
  "项目采用学生主导、教师指导、用户共同验证的协作方式。这里需要区分两个不同概念：产品面向的是知识工作者，"
  "不以学生为定位依据；而这个项目本身是学生团队在教师指导下完成的，团队身份与产品用户定位互不影响。"
  "团队分工覆盖产品与架构、智能体工程与评测、用户研究与商业验证，每个里程碑都保留任务单、版本记录、"
  "测试报告与用户证据。指导教师负责方法、伦理、安全与阶段评审，不代替学生完成核心研发与答辩。")

# ---------- 第十二章标题 ----------
h = find_p("12.1 学习与知识工作的直接价值")
set_text(h, "12.1 知识工作的直接价值")

doc.save(SRC)
print("part C done; paras", len(doc.paragraphs))