# -*- coding: utf-8 -*-
"""Recalculate heading page numbers (skipping the 2-page TOC) and refresh the TOC cache."""
import copy, glob, os, io, re
import pymupdf
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCS = os.path.join(r"T:\创新创业\OwO-master\docs")
SRC = [x for x in glob.glob(os.path.join(DOCS, "Cuttle*v8.docx")) if "预览" not in x][0]
PDF = sorted(glob.glob(os.path.join(DOCS, ".qa_v5", "pdfAA", "*.pdf")))[0]

doc = Document(SRC)
def ptxt(p):
    return "".join(t.text or "" for t in p._element.iter(qn('w:t')))

pdf = pymupdf.open(PDF)
# 目录占 2 页 -> 正文自 PDF 第 4 页（含）之后开始
TOC_END = None
for i in range(pdf.page_count):
    if "目录" in pdf[i].get_text() and i < 4:
        TOC_END = i
BODY_START = 3
pages = [pdf[i].get_text().replace("\n", "").replace(" ", "") for i in range(pdf.page_count)]

heads = []
for p in doc.paragraphs:
    if p.style.name in ("Heading 1", "Heading 2"):
        t = ptxt(p).strip()
        if t and t != "目录":
            heads.append((p.style.name, t))

rows = []
for style, text in heads:
    key = text.replace(" ", "")
    for i in range(BODY_START, len(pages)):
        if key in pages[i]:
            rows.append((style, text, str(i + 1)))
            break
print("resolved:", len(rows), "of", len(heads))
missing = [t for s, t in heads if not any(r[1] == t for r in rows)]
if missing:
    print("missing:", missing[:6])

# 更新目录缓存页码
def visible(p):
    return "".join(t.text or "" for t in p._element.iter(qn('w:t')))

by_label = {t: pg for _, t, pg in rows}
updated = 0
for p in doc.paragraphs:
    el = p._element
    if not el.findall('.//' + qn('w:tab')):
        continue
    lbl = visible(p).strip()
    if not lbl:
        continue
    # 段落文本已是“标题+页码”拼接，取出标题部分
    match = None
    for t, pg in by_label.items():
        if lbl.startswith(t):
            match = (t, pg); break
    if not match:
        continue
    t, pg = match
    runs = p.runs
    if not runs:
        continue
    runs[0].text = t
    # 找到页码 run
    page_run = None
    for r in runs[1:]:
        if (r.text or "").strip().isdigit():
            page_run = r; break
    if page_run is None:
        page_run = p.add_run("")
    page_run.text = pg
    for r in runs[1:]:
        if r is not page_run:
            r.text = ""
    updated += 1
print("TOC cache entries updated:", updated)
doc.save(SRC)

# 校验
d2 = Document(SRC)
bad = []
for p in d2.paragraphs:
    if p._element.findall('.//' + qn('w:tab')):
        lbl = "".join(t.text or "" for t in p._element.iter(qn('w:t'))).strip()
        for t, pg in by_label.items():
            if lbl.startswith(t):
                tail = lbl[len(t):].strip()
                if tail != pg:
                    bad.append((t, tail, pg))
                break
print("mismatched entries:", bad[:6])
print("sample:", list(by_label.items())[:6])
