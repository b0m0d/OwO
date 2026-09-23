# -*- coding: utf-8 -*-
"""Tier-1 fixes 1-2: correct the swapped 图3/图8 captions, rewrite 表2 (4 columns), fix cross-references."""
import copy
from docx import Document
from docx.oxml.ns import qn

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""; r._element.getparent().remove(r._element)

def set_cell(cell, text):
    ps = cell.paragraphs
    set_text(ps[0], text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

def find(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)

def E(prefix, text):
    set_text(find(prefix), text)

def first_cells(t):
    return [c.text.strip() for c in t.rows[0].cells]

def get_table(sig, ncols=None):
    for t in doc.tables:
        fc = first_cells(t)
        if fc[:len(sig)] == sig and (ncols is None or len(fc) == ncols):
            return t
    raise KeyError(str(sig))

# ---- 1) 图 3（用户研究路径）题注 ----
E("图 3 主要产品在任务生命周期覆盖程度与执行治理深度上的定位",
  "图 3 Cuttle 用户研究与价值验证路径")
# ---- 2) 图 8（竞争定位）题注 ----
E("图 8 主要产品在任务生命周期六个环节上的覆盖定位",
  "图 8 主要产品在任务生命周期与执行治理深度上的定位")

# ---- 3) 表 2 由五列压成四列，并修正错误列 ----
t = get_table(["验证方向", "验证方式"], 5)
# 先删掉第五列
for r in t.rows:
    cells = r.cells
    if len(cells) > 4:
        cells[-1]._element.getparent().remove(cells[-1]._element)
grid = t._element.find(qn('w:tblGrid'))
cols = grid.findall(qn('w:gridCol'))
if len(cols) > 4:
    grid.remove(cols[-1])
ROWS2 = [
 ["验证方向", "验证方法", "判定标准", "未达标调整"],
 ["入口机制优于独立工作台与全局快捷 Agent", "三组入口同任务对照，顺序平衡",
  "上下文重复说明次数降低 15% 以上；Origin 绑定准确率不下降；发起步骤数与结果返回成本至少一项显著改善",
  "缩小入口范围，保留快捷指令"],
 ["意图胶囊减少重复说明并抑制多余读取", "对照任务计时与记录，情境标注集",
  "Context Recall 不低于 90%；情境过度读取率低于 5%；重复输入减少 30%",
  "减少自动感知，增强手动选区"],
 ["三维自治决策优于固定档位", "同任务计时与接管计数，标注集评测",
  "三维决策 Macro-F1 不低于 0.85；过度升级率低于 10%；高风险低估率低于 3%",
  "只保留高价值任务升级"],
 ["可见审批提高影响范围理解", "敏感任务可用性测试",
  "影响范围判断正确率不低于 95%；误批准率低于 3%", "增加解释与二次确认"],
 ["端到端任务可稳定完成", "固定任务套件与真实任务",
  "任务成功率不低于 85%；可逆任务恢复率不低于 90%；完成时间降低 25%",
  "收敛任务范围并加固恢复路径"],
 ["团队与机构客户愿意为治理能力付费", "试点与报价测试",
  "取得至少 3 份付费意向或采购流程", "转向个人订阅与开发者生态"],
 ["纵向留存与首用接受度", "产品分析与合作记录",
  "纵向组四周留存率 40%；独立组首用意愿达标", "修复新用户上手路径与首用价值"],
]
for ri, vals in enumerate(ROWS2):
    for ci, v in enumerate(vals):
        set_cell(t.rows[ri].cells[ci], v)

# 列宽按四列重设
CONTENT = 9638
for r in t.rows:
    widths = [3000, 1900, 3438, 1300]
    ci = 0
    for c in r.cells:
        pr = c._tc.find(qn('w:tcPr'))
        if pr is None:
            pr = c._tc.makeelement(qn('w:tcPr'), {}); c._tc.insert(0, pr)
        tcW = pr.find(qn('w:tcW'))
        if tcW is None:
            tcW = c._tc.makeelement(qn('w:tcW'), {}); pr.insert(0, tcW)
        tcW.set(qn('w:type'), 'dxa'); tcW.set(qn('w:w'), str(widths[ci]))
        ci += 1
grid = t._element.find(qn('w:tblGrid'))
for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
    gc.set(qn('w:w'), str(w))

# 表注说明数据来源
import copy as _c
from docx.text.paragraph import Paragraph
cap = find("表 2 核心验证指标与决策门")
set_text(cap, "表 2 核心验证指标与决策门")
model = None
for p in doc.paragraphs:
    if p.style.name == "Normal" and len(p.text) > 60:
        model = p; break
el = _c.deepcopy(model._element)
# 注放在表之后
tbl_el = t._element
nxt = tbl_el.getnext()
note_el = _c.deepcopy(model._element)
tbl_el.addnext(note_el)
note = Paragraph(note_el, t._parent)
set_text(note, "注：判定数据来源包括任务日志、情境标注集、产品分析记录与试点合作材料。")

# ---- 4) 交叉引用修正 ----
for p in doc.paragraphs:
    tt = p.text
    if "核心指标与阶段目标见表 3" in tt:
        set_text(p, tt.replace("核心指标与阶段目标见表 3", "核心验证指标与阶段目标见表 2"))
    if "席位计价 4 万元" in tt:
        set_text(p, tt.replace("席位计价 4 万元", "团队许可 4 万元每团队每年"))
    if "席位计价 4 万元" in p.text:
        pass

# ---- 5) 表 10 标题去掉资源需求 ----
for p in doc.paragraphs:
    if p.text.strip() == "表 10 三年经营预测与资源需求":
        set_text(p, "表 10 三年经营预测")

doc.save(SRC)
print("tier-1 batch 1 done")