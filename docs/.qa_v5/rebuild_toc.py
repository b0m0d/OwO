# -*- coding: utf-8 -*-
"""Rebuild the TOC field cleanly (it was damaged by a run-level refresh) with correct
dot-leader tab stops and the final cached page numbers."""
import copy, re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)
BODY = doc.element.body

PAGES = [("项目摘要", "3"), ("第一章  真实问题与项目机会", "4"), ("第二章  用户需求与验证设计", "7"),
         ("第三章  产品方案与体验闭环", "10"), ("第四章  核心技术与创新机制", "13"),
         ("第五章  安全伦理与数据治理", "18"), ("第六章  市场空间与竞争定位", "20"),
         ("第七章  商业模式与市场进入", "22"), ("第八章  研发验证与实施路径", "25"),
         ("第九章  团队协作与人才培养", "27"), ("第十章  财务规划与资源配置", "29"),
         ("第十一章  风险管理", "32"), ("第十二章  社会价值与发展愿景", "33"),
         ("参考资料", "34")]

def ptxt(p):
    return "".join(t.text or "" for t in p._element.iter(qn('w:t')))

# 1) 定位并删除现有 TOC 段落（含残留域字符）
toc_paras = []
for p in list(doc.paragraphs):
    el = p._element
    has_tab = bool(el.findall('.//' + qn('w:tab')))
    has_fld = bool(el.findall('.//' + qn('w:fldChar'))) or bool(el.findall('.//' + qn('w:instrText')))
    label = ptxt(p).strip()
    if (has_tab or has_fld) and (label == "" or any(label.startswith(k) for k, _ in PAGES)):
        toc_paras.append(p)
print("found TOC paragraphs:", len(toc_paras))
anchor = toc_paras[0]._element.getprevious()
for p in toc_paras:
    p._element.getparent().remove(p._element)

# 2) 取一个正文段落作为格式模板
model = None
for p in doc.paragraphs:
    if p.style.name == "Normal" and len(p.text) > 60:
        model = p; break

def make_entry(text, page, bold):
    el = copy.deepcopy(model._element)
    # 清空内容，保留 pPr
    for child in list(el):
        if child.tag != qn('w:pPr'):
            el.remove(child)
    pPr = el.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr'); el.insert(0, pPr)
    for tag in ('w:ind', 'w:tabs', 'w:spacing'):
        old = pPr.find(qn(tag))
        if old is not None:
            pPr.remove(old)
    ind = OxmlElement('w:ind'); ind.set(qn('w:firstLine'), '0'); pPr.append(ind)
    sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), '0'); sp.set(qn('w:line'), '264')
    sp.set(qn('w:lineRule'), 'auto'); pPr.append(sp)
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right'); tab.set(qn('w:leader'), 'dot'); tab.set(qn('w:pos'), '9638')
    tabs.append(tab); pPr.append(tabs)
    # 文本 run
    r1 = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rf.set(qn(a), 'Microsoft YaHei')
    rPr.append(rf)
    if bold:
        rPr.append(OxmlElement('w:b'))
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '21'); rPr.append(sz)
    r1.append(rPr)
    t1 = OxmlElement('w:t'); t1.set(qn('xml:space'), 'preserve'); t1.text = text
    r1.append(t1)
    # 制表符 run
    r2 = OxmlElement('w:r')
    r2.append(copy.deepcopy(rPr))
    r2.append(OxmlElement('w:tab'))
    # 页码 run
    r3 = OxmlElement('w:r')
    r3.append(copy.deepcopy(rPr))
    t3 = OxmlElement('w:t'); t3.text = page
    r3.append(t3)
    for r in (r1, r2, r3):
        el.append(r)
    return el

def field_run(kind=None, instr=None):
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rf.set(qn(a), 'Microsoft YaHei')
    rPr.append(rf)
    r.append(rPr)
    if instr is not None:
        t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = instr
        r.append(t)
    else:
        fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); r.append(fc)
    return r

# 3) begin / instrText / separate 段落
prev = anchor
def add_para_with_runs(runs):
    global prev
    el = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    ind = OxmlElement('w:ind'); ind.set(qn('w:firstLine'), '0'); pPr.append(ind)
    sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), '0'); pPr.append(sp)
    el.append(pPr)
    for r in runs:
        el.append(r)
    prev.addnext(el)
    prev = el
    return el

add_para_with_runs([field_run('begin')])
add_para_with_runs([field_run(instr=' TOC \\o "1-2" \\h \\z \\u '), field_run('separate')])

# 4) 缓存条目
for text, page in PAGES:
    el = make_entry(text, page, bold=(text.startswith("第") or text in ("项目摘要", "参考资料")))
    prev.addnext(el)
    prev = el

# 5) end
add_para_with_runs([field_run('end')])

doc.save(SRC)

# 6) 校验
d2 = Document(SRC)
import zipfile
dx = zipfile.ZipFile(SRC).read("word/document.xml").decode("utf-8")
print("TOC 指令存在:", "TOC " in dx)
print("fldChar 数量:", dx.count("fldChar"), "(应为 6：begin/separate/end × ...)")
print("instrText 数量:", dx.count("instrText"))
n = 0
for p in d2.paragraphs:
    t = p.text.strip()
    if t and not t.startswith("目录") and any(t.startswith(k) for k, _ in PAGES):
        n += 1
print("目录条目数:", n)
for p in d2.paragraphs[:24]:
    t = p.text.strip()
    if t and (t.startswith("第") or t.startswith("项目摘要") or t.startswith("参考资料")):
        print("   ", repr(t[:52]))
