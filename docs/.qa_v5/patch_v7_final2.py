# -*- coding: utf-8 -*-
"""v7 补丁 2：删除残留公式行，补实验设计的两段缺失说明。"""
import copy

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

DOC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v7.docx"
doc = docx.Document(DOC)
body = doc.element.body
TEMPLATE_P = doc.paragraphs[10]


def el_text(el):
    return ''.join(n.text or '' for n in el.iter(qn('w:t'))).strip()


def find_p(sub):
    hits = [p for p in doc.paragraphs if sub in p.text]
    return hits[0] if hits else None


def set_text(p_obj, text):
    runs = p_obj.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        p_obj.add_run(text)


def para_after(ref, text):
    el = copy.deepcopy(ref._element)
    ref._element.addnext(el)
    p = Paragraph(el, doc)
    set_text(p, text)
    return p


def para_after_el(el, text):
    new_el = copy.deepcopy(TEMPLATE_P._element)
    el.addnext(new_el)
    p = Paragraph(new_el, doc)
    set_text(p, text)
    return p


def tbl_el(prefix):
    for el in body.iterchildren():
        if el.tag.endswith('}tbl'):
            hdr = "".join(c.text for c in Table(el, doc).rows[0].cells)
            if hdr.startswith(prefix):
                return el
    raise KeyError(prefix)


# 1) 删除残留的 IC 公式行
p = find_p("IC = [I, C, O, P, F, R]")
if p:
    p._element.getparent().remove(p._element)
    print("已删除 IC 公式行")

# 2) 表 4 之后补：阈值口径 + 标注规则
if not find_p("表中阈值是预实验前预设的判断标准"):
    t4 = tbl_el("假设验证方法")
    anchor = para_after_el(list(t4.iterchildren())[-1],
        "表中阈值是预实验前预设的判断标准，不是研究结论，也不是既有门槛。"
        "正式实验前依据预实验效应量做统计功效分析，据此确定各指标所需样本量，再把定性阈值收敛为具体数值。")
    para_after(anchor,
        "过度感知与过度自治的判定不能由开发者自行认定。判定依据是预先文档化的“完成任务最小充分信息集”与"
        "“最小必要执行档位”：由两名不参与研发的独立标注者对同一批任务样本独立标注，计算标注者一致率，"
        "不一致样本由第三方裁决并记录理由；标注一致率达到预设水平后，该指标才对外报告。")
    print("已补表 4 后的阈值与标注规则说明")

# 3) 表 8 之后补：对照实验公平性
if not find_p("对照实验的公平性前提"):
    t8 = tbl_el("实验对照组")
    para_after_el(list(t8.iterchildren())[-1],
        "对照实验的公平性前提必须写清：处理组与对照组使用同一模型与同一版本、相同的工具与权限范围、"
        "相同的任务材料与结果要求，唯一变量是“任务如何进入系统、上下文如何获得”。"
        "若对照组的工具权限被削减或模型版本更低，效率提升就不能归因于入口设计。"
        "实验采用 within-subject crossover：同一用户在两轮中分别先用传统工作台与先用 Cuttle，"
        "以消除学习顺序与熟练度带来的偏差。")
    print("已补表 8 后的公平性说明")

doc.save(DOC)
print("saved | paragraphs:", len(doc.paragraphs))
