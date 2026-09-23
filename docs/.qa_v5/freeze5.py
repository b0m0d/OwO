# -*- coding: utf-8 -*-
import re
from docx import Document
SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
d = Document(SRC)
full = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
issues = []
banned = ["20 万元每项目", "9 万元实施与维护", "按席位", "4 万元每年每客户", "核心设计目标的不同",
          "禁止读取率", "受控行动的受控自治", "拿不到", "哪份文档的哪个位置", "副场景）", "连续自治"]
for b in banned:
    if b in full: issues.append("BANNED: " + b)
must = ["当前主要交互形态与价值重心", "4 万元每团队每年", "18 万元每项目", "8 万元实施与维护",
        "个人付费用户上限 200 人", "资料核验（主链子任务）", "代码修改" if False else "代码修复（副链）",
        "约束：Origin 绑定准确率不下降", "稳定锚点", "禁止字段访问率", "从输入意图到行动的受控自治升级闭环"]
for mk in must:
    if mk not in full: issues.append("MISSING: " + mk)
# arithmetic
P, T, D = 240.0, 40000.0, 180000.0
IND = [ [80,600,1300], [120,1200,2800], [180,2000,5000] ]
TEAM = [ [0,3,8], [0,5,24], [0,10,40] ]
DEP = [ [0,0,0], [0,0,2], [0,1,4] ]
PUB = { "保守": [1.9,26.4,63.2], "基准": [2.9,48.8,199.2], "进取": [4.3,106.0,352.0] }
for i, k in enumerate(("保守","基准","进取")):
    for y in range(3):
        rev = round((IND[i][y]*P + TEAM[i][y]*T + DEP[i][y]*D)/1e4, 1)
        if abs(rev - PUB[k][y]) > 0.05:
            issues.append(f"ARITH {k} y{y}: {rev} vs {PUB[k][y]}")
print("ISSUES:", len(issues))
for i in issues: print("  !", i)
print("paras", len(d.paragraphs), "tables", len(d.tables), "refs", sum(1 for p in d.paragraphs if re.match(r"^\[\d+\]", p.text.strip())))