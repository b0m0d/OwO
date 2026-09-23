# -*- coding: utf-8 -*-
"""Step 3: compress 表 10 so 11.2 fits its page, closing the near-empty page."""
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def first_cells(t):
    return [c.text.strip() for c in t.rows[0].cells]

def set_cell(cell, text):
    ps = cell.paragraphs
    runs = ps[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
            r._element.getparent().remove(r._element)
    else:
        ps[0].add_run(text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

def set_row(t, ri, vals):
    for ci, v in enumerate(vals):
        set_cell(t.rows[ri].cells[ci], v)

DATA = [
 ["风险", "概率", "影响", "预警信号", "应对措施"],
 ["输入法信任风险", "高", "极高", "用户担心输入内容被读取", "默认最小读取，可逐应用关闭，独立审计"],
 ["输入延迟风险", "高", "高", "正常打字出现可感延迟", "快慢路径分离，打字链路不调用模型，设延迟上限"],
 ["提示注入与情境污染", "中", "极高", "文档或网页内嵌指令影响授权范围", "情境与指令隔离，来源标记，高风险二次确认"],
 ["情境误判与隐私", "中", "极高", "敏感字段被读取，用户纠正上升", "禁止字段硬屏蔽，最小读取，独立审计"],
 ["执行误操作", "中", "极高", "越权、外发或不可逆动作", "最小权限，独立审批，预览，沙箱与回滚"],
 ["平台厂商集成同类能力", "中", "高", "系统或头部输入法内置相似入口", "聚焦跨应用协议与可靠性，入口与模型可替换"],
]

for t in doc.tables:
    fc = first_cells(t)
    if fc[:2] == ["风险", "概率"]:
        for ri, vals in enumerate(DATA):
            set_row(t, ri, vals)
        for r in t.rows:
            for c in r.cells:
                tcPr = c._tc.get_or_add_tcPr()
                old = tcPr.find(qn('w:tcMar'))
                if old is not None:
                    tcPr.remove(old)
                mar = OxmlElement('w:tcMar')
                for sn, val in (('top', 30), ('start', 60), ('bottom', 30), ('end', 60)):
                    e = OxmlElement('w:' + sn)
                    e.set(qn('w:w'), str(val))
                    e.set(qn('w:type'), 'dxa')
                    mar.append(e)
                tcPr.append(mar)
                for p in c.paragraphs:
                    p.paragraph_format.space_before = 0
                    p.paragraph_format.space_after = 0
                    p.paragraph_format.line_spacing = 0.92
        print("  compacted 表 10 and trimmed wording")
        break

doc.save(SRC)
print("step 3 done")
