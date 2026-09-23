# -*- coding: utf-8 -*-
"""v7 补丁 4：写入最后三段实验/指标/安全说明。"""
import copy

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

DOC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v7.docx"
doc = docx.Document(DOC)
body = doc.element.body
TEMPLATE_P = doc.paragraphs[10]


def find_p(sub):
    return next((p for p in doc.paragraphs if sub in p.text), None)


def set_text(p_obj, text):
    runs = p_obj.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        p_obj.add_run(text)


def tbl_el(prefix):
    for el in body.iterchildren():
        if el.tag.endswith('}tbl'):
            hdr = "".join(c.text for c in Table(el, doc).rows[0].cells)
            if hdr.startswith(prefix):
                return el
    raise KeyError(prefix)


def para_after_el(el, text):
    new_el = copy.deepcopy(TEMPLATE_P._element)
    el.addnext(new_el)
    p = Paragraph(new_el, doc)
    set_text(p, text)
    print("   +", text[:32])
    return p


WRITES = [
    ("实验对照组", "对照实验的公平性前提", 
     "对照实验的公平性前提必须写清：处理组与对照组使用同一模型与同一版本、相同的工具与权限范围、"
     "相同的任务材料与结果要求，唯一变量是“任务如何进入系统、上下文如何获得”。"
     "若对照组的工具权限被削减或模型版本更低，效率提升就不能归因于入口设计。"
     "实验采用 within-subject crossover：同一用户在两轮中分别先用传统工作台与先用 Cuttle，"
     "以消除学习顺序与熟练度带来的偏差。"),
    ("维度指标", "全部指标共用一套统计口径",
     "全部指标共用一套统计口径：以配对样本比较为主，报告效应量与置信区间，而不只报告均值差异；"
     "样本量在正式实验前依据预实验效应量做统计功效分析后确定；所有指标注明数据来源、标注规则与样本量。"
     "指标之间的解释关系是：原地闭环完成率是首要指标，其余指标用于解释它为什么高或低。"),
    ("安全目标指标", "两项专有指标的判定规则必须可复现",
     "两项专有指标的判定规则必须可复现：依据预先文档化的“完成任务最小充分信息集”与“最小必要执行档位”，"
     "由两名不参与研发的独立标注者独立标注，计算标注者一致率，不一致样本由第三方裁决并记录理由；"
     "一致率达到预设水平后该指标才对外报告。"),
]

for prefix, marker, text in WRITES:
    if find_p(marker):
        print("跳过（已存在）:", marker)
        continue
    el = tbl_el(prefix)
    para_after_el(list(el.iterchildren())[-1], text)

doc.save(DOC)
print("saved | paragraphs:", len(doc.paragraphs))
