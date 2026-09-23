# -*- coding: utf-8 -*-
"""Restore the 目录 heading, then rebuild the TOC from toc_rows.txt."""
import copy, glob, os, io, re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCS = os.path.join(r"T:\创新创业\OwO-master\docs")
SRC = [x for x in glob.glob(os.path.join(DOCS, "Cuttle*v8.docx")) if "预览" not in x][0]
doc = Document(SRC)

def ptxt(p):
    return "".join(t.text or "" for t in p._element.iter(qn('w:t')))

# 恢复“目录”标题样式
for p in doc.paragraphs:
    if ptxt(p).strip() == "目录" and p.style.name != "Heading 1":
        p.style = doc.styles["Heading 1"]
        p.paragraph_format.space_after = 0
        print("restored 目录 as Heading 1")
        break

rows = []
for line in io.open(os.path.join(DOCS, ".qa_v5", "toc_rows.txt"), encoding="utf-8"):
    parts = line.rstrip("\n").split("\t")
    if len(parts) == 3:
        rows.append(tuple(parts))
print("toc rows:", len(rows))

# 删除旧 TOC 段落
toc_paras = []
for p in list(doc.paragraphs):
    el = p._element
    lbl = ptxt(p).strip()
    if el.findall('.//' + qn('w:fldChar')) or el.findall('.//' + qn('w:instrText')):
        toc_paras.append(p)
    elif el.findall('.//' + qn('w:tab')) and lbl and not lbl.endswith(("。", "；")) \
            and not re.match(r"^\d+\.\d+ ", lbl) is None:
        toc_paras.append(p)
    elif el.findall('.//' + qn('w:tab')) and lbl and not lbl.endswith(("。", "；")):
        toc_paras.append(p)
print("old toc paragraphs removed:", len(toc_paras))
anchor = toc_paras[0]._element.getprevious() if toc_paras else None
for p in toc_paras:
    p._element.getparent().remove(p._element)

model = None
for p in doc.paragraphs:
    if p.style.name == "Normal" and len(p.text) > 60:
        model = p; break

def make_entry(text, page, level):
    el = copy.deepcopy(model._element)
    for child in list(el):
        if child.tag != qn('w:pPr'):
            el.remove(child)
    pPr = el.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr'); el.insert(0, pPr)
    for tag in ('w:ind', 'w:tabs', 'w:spacing'):
        old = pPr.find(qn(tag))
        if old is not None:
            pPr.remove(old)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:firstLine'), '0')
    if level:
        ind.set(qn('w:left'), '340')
    pPr.append(ind)
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:after'), '0')
    sp.set(qn('w:line'), '260' if level == 0 else '248')
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)
    tabs = OxmlElement('w:tabs')
    tb = OxmlElement('w:tab')
    tb.set(qn('w:val'), 'right'); tb.set(qn('w:leader'), 'dot'); tb.set(qn('w:pos'), '9638')
    tabs.append(tb); pPr.append(tabs)
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rf.set(qn(a), 'Microsoft YaHei')
    rPr.append(rf)
    if level == 0:
        rPr.append(OxmlElement('w:b'))
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '21' if level == 0 else '20'); rPr.append(sz)
    r1 = OxmlElement('w:r'); r1.append(copy.deepcopy(rPr))
    t1 = OxmlElement('w:t'); t1.set(qn('xml:space'), 'preserve'); t1.text = text; r1.append(t1)
    r2 = OxmlElement('w:r'); r2.append(copy.deepcopy(rPr)); r2.append(OxmlElement('w:tab'))
    r3 = OxmlElement('w:r'); r3.append(copy.deepcopy(rPr))
    t3 = OxmlElement('w:t'); t3.text = page; r3.append(t3)
    for r in (r1, r2, r3):
        el.append(r)
    return el

def field_run(kind=None, instr=None):
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rf.set(qn(a), 'Microsoft YaHei')
    rPr.append(rf); r.append(rPr)
    if instr is not None:
        t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = instr; r.append(t)
    else:
        fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); r.append(fc)
    return r

prev = anchor
def add_field_para(runs):
    global prev
    el = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    ind = OxmlElement('w:ind'); ind.set(qn('w:firstLine'), '0'); pPr.append(ind)
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '0'); sp.set(qn('w:after'), '0')
    sp.set(qn('w:line'), '20'); sp.set(qn('w:lineRule'), 'exact')
    pPr.append(sp)
    rp = OxmlElement('w:rPr'); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '2'); rp.append(sz)
    pPr.append(rp)
    el.append(pPr)
    for r in runs:
        el.append(r)
    prev.addnext(el); prev = el

add_field_para([field_run('begin')])
add_field_para([field_run(instr=' TOC \\o "1-2" \\h \\z \\u '), field_run('separate')])
for style, text, page in rows:
    level = 0 if style == "Heading 1" else 1
    el = make_entry(text, page, level)
    prev.addnext(el); prev = el
add_field_para([field_run('end')])

doc.save(SRC)
print("TOC rebuilt with", len(rows), "entries")
