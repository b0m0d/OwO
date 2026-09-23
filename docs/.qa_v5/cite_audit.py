# -*- coding: utf-8 -*-
"""Reference citation audit: which of the 30 entries are actually cited in the body?"""
import re
from docx import Document
from docx.oxml.ns import qn

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)
paras = [p.text for p in doc.paragraphs]
cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]

# reference list entries
refs = {}
for p in paras:
    m = re.match(r"^\[(\d+)\]\s*(.+)$", p.strip())
    if m:
        refs[int(m.group(1))] = m.group(2).strip()

body = []
for p in paras:
    if re.match(r"^\[\d+\]", p.strip()):
        continue
    body.append(p)
body_text = "\n".join(body) + "\n" + "\n".join(cells)

cited = set()
for m in re.finditer(r"\[(\d+)\]", body_text):
    cited.add(int(m.group(1)))

print("reference entries:", len(refs))
print("cited in body:", sorted(cited))
print("never cited:", sorted(set(refs) - cited))
print()
for i in sorted(set(refs) - cited):
    print(f"  [{i}] {refs[i][:96]}")