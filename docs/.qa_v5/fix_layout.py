# -*- coding: utf-8 -*-
"""Close the two layout gaps: shrink 表 1 to fit one page, allow 图 4 on the chapter-3 opening page."""
import copy
from docx import Document
from docx.oxml.ns import qn

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""; r._element.getparent().remove(r._element)

def set_cell(cell, text):
    ps = cell.paragraphs
    set_text(ps[0], text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

def first_cells(t):
    return [c.text.strip() for c in t.rows[0].cells]

def get_table(sig, ncols=None):
    for t in doc.tables:
        fc = first_cells(t)
        if fc[:len(sig)] == sig and (ncols is None or len(fc) == ncols):
            return t
    raise KeyError(str(sig))

# 1) 表 1 缩为 4 行，保证完整落在摘要页
t = get_table(["项目要素", "核心内容"], 3)
while len(t.rows) > 4:
    tr = t.rows[len(t.rows) - 1]._element
    tr.getparent().remove(tr)
for ri, vals in enumerate([
 ["项目要素", "核心内容", "验证依据"],
 ["真实问题", "AI 入口、情境、行动与结果彼此割裂", "访谈、任务日记、跨应用流程观察"],
 ["项目方案", "在输入时刻理解情境并升级为受控行动", "交互原型、演示闭环、任务日志"],
 ["技术创新与产业路径", "意图胶囊 IC 与受控自治升级闭环；知识工作场景切入、团队许可与私有部署",
  "三组入口对照与消融实验；试点、留存、付费验证与合作材料"],
]):
    for ci, v in enumerate(vals):
        set_cell(t.rows[ri].cells[ci], v)

# 2) 图 4 与 3.1 正文同页：取消图片段落的 page-break-before
fig4 = None
for p in doc.paragraphs:
    if p.text.strip().startswith("图 4 Cuttle 一个入口与四个核心系统"):
        fig4 = p
        break
# 找到图 4 上方的图片段落（空段落）
if fig4 is not None:
    prev = fig4._element.getprevious()
    from docx.text.paragraph import Paragraph
    steps = 0
    while prev is not None and steps < 4:
        if prev.tag.endswith('}p'):
            pp = Paragraph(prev, fig4._parent)
            pf = pp.paragraph_format
            pf.page_break_before = False
            pf.keep_with_next = True
            if pp.text.strip() == "":
                break
        prev = prev.getprevious()
        steps += 1

doc.save(SRC)
print("layout gaps handled")