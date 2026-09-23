# -*- coding: utf-8 -*-
"""Fix 表 2's new 分组 header cell, and enforce uniform header centering + vertical centering
across every body table."""
import glob, os, copy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCS = os.path.join(r"T:\创新创业\OwO-master\docs")
SRC = [x for x in glob.glob(os.path.join(DOCS, "Cuttle*v8.docx")) if "预览" not in x][0]
doc = Document(SRC)

def clean_set(cell, text, bold=None):
    """Replace a cell's content with a single centered paragraph."""
    tc = cell._tc
    tcPr = tc.find(qn('w:tcPr'))
    # 保留 tcPr，清掉段落
    for p in tc.findall(qn('w:p')):
        tc.remove(p)
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); pPr.append(jc)
    ind = OxmlElement('w:ind'); ind.set(qn('w:firstLine'), '0'); pPr.append(ind)
    sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), '0'); pPr.append(sp)
    p.append(pPr)
    if text:
        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rf = OxmlElement('w:rFonts')
        for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
            rf.set(qn(a), 'Microsoft YaHei')
        rPr.append(rf)
        if bold:
            rPr.append(OxmlElement('w:b'))
            col = OxmlElement('w:color'); col.set(qn('w:val'), 'FFFFFF'); rPr.append(col)
        r.append(rPr)
        t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
        r.append(t)
        p.append(r)
    tc.append(p)

# ---------- 1) 表 2 分组表头 ----------
for t in doc.tables:
    fc = [c.text.strip() for c in t.rows[0].cells]
    if len(fc) == 5 and fc[0] == "" and fc[1].startswith("验证方向"):
        clean_set(t.rows[0].cells[0], "分组", bold=True)
        print("  表 2 分组表头已补")
        # 分组列正文也居中
        for r in t.rows[1:]:
            txt = r.cells[0].text.strip()
            clean_set(r.cells[0], txt)
        break

# ---------- 2) 所有表：表头水平居中 + 垂直居中 ----------
fixed_h = 0
for ti, t in enumerate(doc.tables):
    if ti <= 1:      # 封面表、目录不在处理范围
        continue
    for c in t.rows[0].cells:
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = 0
            p.paragraph_format.space_after = 0
        tcPr = c._tc.get_or_add_tcPr()
        va = tcPr.find(qn('w:vAlign'))
        if va is None:
            va = OxmlElement('w:vAlign'); tcPr.append(va)
        va.set(qn('w:val'), 'center')
        fixed_h += 1
print("header cells normalized:", fixed_h)

# ---------- 3) 正文单元格垂直居中、首行缩进归零 ----------
fixed_b = 0
for ti, t in enumerate(doc.tables):
    if ti <= 1:
        continue
    for r in t.rows[1:]:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent = 0
            tcPr = c._tc.get_or_add_tcPr()
            va = tcPr.find(qn('w:vAlign'))
            if va is None:
                va = OxmlElement('w:vAlign'); tcPr.append(va)
            va.set(qn('w:val'), 'center')
            fixed_b += 1
print("body cells normalized:", fixed_b)

doc.save(SRC)
print("saved")
