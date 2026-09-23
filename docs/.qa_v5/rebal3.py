# -*- coding: utf-8 -*-
"""Rebalance batch 3: chapter 6 restructure, chapter 7 growth, chapter 10/11 additions."""
import copy
from docx import Document
from docx.text.paragraph import Paragraph

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""; r._element.getparent().remove(r._element)

def find(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)

def E(prefix, text):
    set_text(find(prefix), text)

def H(prefix, text):
    p = find(prefix); set_text(p, text); return p

def model_p():
    for p in doc.paragraphs:
        if p.style.name == "Normal" and len(p.text) > 60:
            return p

def insert_after(anchor, texts):
    prev = anchor
    for t in texts:
        el = copy.deepcopy(model_p()._element)
        prev._element.addnext(el)
        np = Paragraph(el, prev._parent)
        set_text(np, t); prev = np
    return prev

def new_heading(ref, text):
    src = find(ref)
    el = copy.deepcopy(src._element)
    return el, src

def clone_after(anchor, ref_heading, text):
    src = find(ref_heading)
    el = copy.deepcopy(src._element)
    anchor._element.addnext(el)
    np = Paragraph(el, anchor._parent)
    set_text(np, text)
    return np

# ---- 第六章：重排为 6.1 / 6.2 / 6.3 / 6.4 ----
H("6.1 市场环境", "6.1 行业机会与目标市场")
p = find("中国互联网络信息中心第 57 次报告显示")
insert_after(p, [
 "需要 Cuttle 的岗位具有三个共同特征：信息来源跨越多个应用、结论需要标注出处、成果需要反复交付。"
 "运营、项目、咨询、研究、行政与产品岗位的日常工作正属于这一类，它们是目前最可能率先付费的个人用户；"
 "需要统一权限、审计与流程复用的中小团队与机构，则是团队许可的首批组织客户。"
 "这一需求随 AI 使用普及而扩大：模型能力越强，用户越需要把分散的模型调用收敛到一条可控、可追溯的链路中。",
])
H("6.2 竞争格局", "6.3 竞争格局")
H("6.3 核心竞争力", "6.4 核心竞争力")
anchor = find("Cuttle 采用自下而上的市场进入测算")
clone_after(anchor, "6.1 行业机会与目标市场", "6.2 市场进入规模")

# ---- 第七章：获客与留存 ----
H("7.3 增长与留存", "7.3 获客、转化与留存")
p = find("留存诊断：区分新鲜感流失")
insert_after(p, [
 "获客路径分为个人与组织两条。个人用户来自开发者社群与专业社区的内容传播、面向高频知识工作者的定向试点、"
 "以及可复用工作流模板的自然扩散，投入以内容与模板为主。组织客户来自个人用户的内部扩散："
 "一名成员在日常工作中稳定使用后，把流程带入所在团队，先形成小范围试点，再进入部门采购流程。",
 "续费动机来自使用过程中积累的资产，而不是模型能力本身：已验证的工作流、个人技能配置、"
 "来源与审计记录、跨应用适配习惯都会随使用加深而沉淀，迁移到其他工具意味着重新建立这些资产。"
 "这也是项目把来源可追溯与审批记录写入产品结构、而不只作为合规要求的原因。",
])
clone_after(find("续费动机来自使用过程中积累的资产"), "7.3 获客、转化与留存", "7.4 从个人版到团队版的转化")
E("7.4 产品版本与单位经济", "7.5 产品版本与单位经济")
insert_after(find("7.4 从个人版到团队版的转化"), [
 "个人专业版是团队许可的入口。成员在个人版中积累的工作流可以直接迁移为团队共享技能，"
 "团队版在此基础上增加成员管理、统一策略、审计记录与预算控制，解决个人工具无法覆盖的治理问题。"
 "这一路径决定了销售方式：前期不设独立销售团队，由产品与运营成员直接对接试点团队，"
 "把服务过程中的共性问题沉淀为标准化交付流程，再逐步扩展到渠道合作。",
])

# ---- 第十章：增长来源更明确 ----
E("收入增长来自三条可核验的路径",
  "收入增长来自三条可核验的路径，而不是场景数量的简单扩张。自然新增来自内容与模板传播带来的个人用户，"
  "以及试点场景向同类组织复制；团队导流来自个人用户在团队内部扩散后形成的组织采购；"
 "开发者渠道来自技能模板与开源组件的复用者在真实工作中产生的订阅。"
 "第二种增长主要来自前两条路径，第三年的增长主要来自第三条路径与团队客户的续订扩容。")

# ---- 第十一章：风险优先级与治理机制 ----
p = find("项目在每一阶段设置明确的停止或转向条件")
insert_after(p, [
 "项目按影响程度对风险分级治理。输入法信任、隐私泄露与执行误操作列为一级风险，由发布门槛控制，"
 "不达标即冻结发布；输入延迟、平台厂商集成同类能力与用户留存列为二级风险，"
 "以性能降级、差异化能力建设与首用体验改善应对。风险状态每两周随产品评审更新，"
 "达到预警条件时触发暂停、降级或转向决策，并在下一版本以回归测试固化整改结果。",
])

doc.save(SRC)
print("batch 3 done")