# -*- coding: utf-8 -*-
"""Detect pPr child-order violations (Word ignores out-of-order w:jc; LO tolerates)."""
import glob, os
from collections import Counter
from docx import Document
from docx.oxml.ns import qn

ORDER = ["pStyle","keepNext","keepLines","pageBreakBefore","framePr","widowControl",
    "numPr","suppressLineNumbers","pBdr","shd","tabs","suppressAutoHyphens","kinsoku",
    "wordWrap","overflowPunct","topLinePunct","autoSpaceDE","autoSpaceDN","bidi",
    "adjustRightInd","snapToGrid","spacing","ind","contextualSpacing","mirrorIndents",
    "suppressOverlap","jc","textDirection","textAlignment","textboxTightWrap",
    "outlineLvl","divId","cnfStyle","rPr","sectPr","pPrChange"]
IDX = {n: i for i, n in enumerate(ORDER)}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

DOCS = os.path.join(r"T:\创新创业\OwO-master\docs")
SRC = [x for x in glob.glob(os.path.join(DOCS, "Cuttle*v8.docx")) if "预览" not in x][0]
doc = Document(SRC)

bad = 0
mult = 0
checked = 0
for ti, t in enumerate(doc.tables):
    for ri, row in enumerate(t.rows):
        for ci, c in enumerate(row.cells):
            pprs = c._tc.findall(".//" + W + "p/" + W + "pPr")
            for pPr in pprs:
                checked += 1
                names = [ch.tag.replace(W, "") for ch in pPr]
                if len(set(names)) != len(names):
                    mult += 1
                seq = [IDX.get(n, -1) for n in names if n in IDX]
                if seq != sorted(seq):
                    bad += 1
                    if bad <= 5:
                        print("  OUT-OF-ORDER T%02d r%d c%d: %s" % (ti, ri, ci, names))
print("checked pPr:", checked, "| out-of-order:", bad, "| duplicate children:", mult)

# also confirm which T01 paragraphs lack jc
t = doc.tables[1]
for ri, row in enumerate(t.rows):
    for ci, c in enumerate(row.cells):
        for p in c.paragraphs:
            pPr = p._p.find(qn('w:pPr'))
            j = pPr.find(qn('w:jc')) if pPr is not None else None
            if j is None:
                print("  T01 r%d c%d NO_JC  text=%r" % (ri, ci, p.text[:30]))
