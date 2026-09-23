# -*- coding: utf-8 -*-
"""v8 stage 5 : clean references and rebuild [13]-[30] correctly."""
import copy
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def ptext(p):
    return "".join(t.text or "" for t in p._element.iter(qn('w:t')))

def set_para(p, text):
    """Replace all text content of a paragraph, keeping pPr and the first run's rPr."""
    el = p._element
    rpr = None
    for r in el.findall(qn('w:r')):
        rPr = r.find(qn('w:rPr'))
        if rPr is not None:
            rpr = copy.deepcopy(rPr)
        break
    for child in list(el):
        if child.tag != qn('w:pPr'):
            el.remove(child)
    from docx.oxml import OxmlElement
    r = OxmlElement('w:r')
    if rpr is not None:
        r.append(rpr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    el.append(r)
    return p

refs = [p for p in doc.paragraphs if ptext(p).strip().startswith("[")]
print("refs before:", len(refs))
for p in refs[12:]:
    p._element.getparent().remove(p._element)
doc.save(SRC)

doc = Document(SRC)
refs = [p for p in doc.paragraphs if ptext(p).strip().startswith("[")]
print("refs after cleanup:", len(refs), "last:", [ptext(p)[:40] for p in refs[-3:]])

REFS = [
 "[13] Yao S, Zhao J, Yu D, et al. ReAct: Synergizing Reasoning and Acting in Language Models. ICLR, 2023.",
 "[14] Schick T, Dwivedi-Yu J, Dessi R, et al. Toolformer: Language Models Can Teach Themselves to Use Tools. NeurIPS, 2023.",
 "[15] Shinn N, Cassano F, Gopinath A, et al. Reflexion: Language Agents with Verbal Reinforcement Learning. NeurIPS, 2023.",
 "[16] Wu Q, Bansal G, Zhang J, et al. AutoGen: Enabling Next-Gen LLM Applications via Multi-Agent Conversation. arXiv:2308.08155, 2023.",
 "[17] Hong S, Zhuge M, Chen J, et al. MetaGPT: Meta Programming for A Multi-Agent Collaborative Framework. ICLR, 2024.",
 "[18] Jimenez C E, Yang J, Wettig A, et al. SWE-bench: Can Language Models Resolve Real-World GitHub Issues? ICLR, 2024.",
 "[19] Zhou S, Xu F F, Zhu H, et al. WebArena: A Realistic Web Environment for Building Autonomous Agents. ICLR, 2024.",
 "[20] Xie T, Zhang D, Chen J, et al. OSWorld: Benchmarking Multimodal Agents for Open-Ended Tasks in Real Computer Environments. NeurIPS, 2024.",
 "[21] Dourish P. What We Talk About When We Talk About Context. Personal and Ubiquitous Computing, 2004, 8(1): 19-30.",
 "[22] Dey A K. Understanding and Using Context. Personal and Ubiquitous Computing, 2001, 5(1): 4-7.",
 "[23] Horvitz E. Principles of Mixed-Initiative User Interfaces. Proceedings of CHI, 1999: 159-166.",
 "[24] Amershi S, Weld D, Vorvoreanu M, et al. Guidelines for Human-AI Interaction. Proceedings of CHI, 2019: 1-13.",
 "[25] Lee J D, See K A. Trust in Automation: Designing for Appropriate Reliance. Human Factors, 2004, 46(1): 50-80.",
 "[26] Parasuraman R, Riley V. Humans and Automation: Use, Misuse, Disuse, Abuse. Human Factors, 1997, 39(2): 230-253.",
 "[27] Microsoft. Text Services Framework (TSF) 与 UI Automation 官方架构文档. Microsoft Learn, 2025.",
 "[28] NIST. Artificial Intelligence Risk Management Framework (AI RMF 1.0). NIST AI 100-1, 2023.",
 "[29] ISO/IEC 25010:2011. Systems and Software Quality Requirements and Evaluation (SQuaRE) - System and Software Quality Models.",
 "[30] Shneiderman B. Human-Centered AI. Oxford University Press, 2022.",
]

model = refs[2]           # [3] 中国互联网络信息中心 (has clean rPr)
anchor = refs[-1]
prev = anchor
from docx.oxml import OxmlElement
for txt in REFS:
    el = copy.deepcopy(model._element)
    # strip everything but pPr
    for child in list(el):
        if child.tag != qn('w:pPr'):
            el.remove(child)
    r = OxmlElement('w:r')
    rPr = el.find(qn('w:pPr'))  # no rPr from pPr
    src_r = model._element.find(qn('w:r'))
    if src_r is None:
        # model text sits inside hyperlink
        hl = model._element.find(qn('w:hyperlink'))
        src_r = hl.find(qn('w:r')) if hl is not None else None
    if src_r is not None and src_r.find(qn('w:rPr')) is not None:
        r.append(copy.deepcopy(src_r.find(qn('w:rPr'))))
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = txt
    r.append(t)
    el.append(r)
    prev._element.addnext(el)
    prev = Paragraph(el, prev._parent)

doc.save(SRC)
d2 = Document(SRC)
final = [ptext(p).strip() for p in d2.paragraphs if ptext(p).strip().startswith("[")]
print("refs final:", len(final))
for f in final:
    print("   ", f[:95])
print("paras", len(d2.paragraphs), "tables", len(d2.tables))