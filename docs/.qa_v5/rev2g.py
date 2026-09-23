# -*- coding: utf-8 -*-
from docx import Document
SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def set_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""; r._element.getparent().remove(r._element)

def set_cell(cell, text):
    ps = cell.paragraphs
    runs = ps[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""; r._element.getparent().remove(r._element)
    else:
        ps[0].add_run(text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

p = None
for x in doc.paragraphs:
    if x.text.strip().startswith("项目资源需求分三层表述"):
        p = x
set_text(p, "项目资源需求分三层表述。已有资源：学校实验室与实验工位、团队成员个人设备、指导教师的方法与安全指导、"
            "开源模型与本地推理能力，这些不计入现金需求，因此资金来源中团队自筹与学校支持的占比高于对外融资。"
            "新增现金需求分两批：第一年 50 万元，用于把原型推进到可验证的试点闭环并完成对照、消融与红队测试；"
            "第二年追加 110 万元，用于场景复制、团队版交付与兼容矩阵扩展。首轮对外表述的资源配置总额为 160 万元，"
            "按决策门分两批释放，任一批未达标即暂停后续投入；保守情景下前两年累计亏损约 59 万元，"
            "基准情景约 97 万元，该额度可覆盖基准情景所需，并保留约四成安全边际。")

t = doc.tables[21]
set_cell(t.rows[2].cells[0], "新增现金需求（第一年）")
set_cell(t.rows[2].cells[2], "50 万元")
set_cell(t.rows[3].cells[1], "场景复制、团队版交付、兼容矩阵扩展、支持与运维")
set_cell(t.rows[3].cells[2], "110 万元")
set_cell(t.rows[3].cells[3], "第二批试点、团队许可交付能力与合作拓展")
doc.save(SRC)

d = Document(SRC)
for r in d.tables[21].rows:
    print(" | ".join(c.text.strip()[:60] for c in r.cells))
print("---")
for x in d.paragraphs:
    if x.text.strip().startswith("项目资源需求分三层表述"):
        print(x.text.strip()[:320])