import docx, re
from docx.oxml.ns import qn
p=r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v6.docx"
d=docx.Document(p)
body=d.element.body
lines=[];pi=0;ti=0
for child in body.iterchildren():
    tag=child.tag.split('}')[1]
    if tag=='p':
        txt=''.join(n.text or '' for n in child.iter(qn('w:t')))
        if txt.strip(): lines.append(f'P{pi} :: {txt}')
        pi+=1
    elif tag=='tbl':
        lines.append(f'@@@ TABLE{ti}')
        ti+=1
open(r"T:\创新创业\OwO-master\docs\.qa_v5\v6_final.txt","w",encoding="utf-8").write("\n".join(lines))
# 术语与一致性检查
txts=[p.text for p in d.paragraphs]
alltext="\n".join(txts)
for t in d.tables:
    for r in t.rows:
        for c in r.cells:
            alltext+="\n"+c.text
checks = {
 "情境胶囊(旧术语)": alltext.count("情境胶囊"),
 "意图胶囊": alltext.count("意图胶囊"),
 "创新五": alltext.count("创新五"),
 "创新四": alltext.count("创新四"),
 "情境过度读取率": alltext.count("情境过度读取率"),
 "自治过度升级率": alltext.count("自治过度升级率"),
 "Context Overreach": alltext.count("Context Overreach"),
 "Autonomy Overreach": alltext.count("Autonomy Overreach"),
 "47.4": alltext.count("47.4"),
 "1740": alltext.count("1740"),
 "120 万元": alltext.count("120 万元"),
 "40 万元": alltext.count("40 万元"),
 "80 万元": alltext.count("80 万元"),
 "8% 至 12%": alltext.count("8% 至 12%"),
 "技能市场": alltext.count("技能市场"),
 "交易分成 15%": alltext.count("交易分成 15%"),
 "设备网格": alltext.count("设备网格"),
 "跨设备能力节点": alltext.count("跨设备能力节点"),
 "F1 不低于 0.90": alltext.count("F1 不低于 0.90"),
 "Macro-F1": alltext.count("Macro-F1"),
 "必须为 0": alltext.count("必须为 0"),
 "核心竞争力": alltext.count("核心竞争力"),
}
for k,v in checks.items(): print(f"{k}: {v}")
print("段落数:", len(d.paragraphs), "表格数:", len(d.tables))
print("总字数(段落):", sum(len(t) for t in txts))
