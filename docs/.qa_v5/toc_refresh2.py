# -*- coding: utf-8 -*-
"""Refresh TOC cached page numbers, addressing paragraphs by their w:tab structure."""
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
PAGES = {"项目摘要": "3", "第一章": "4", "第二章": "7", "第三章": "10", "第四章": "13",
         "第五章": "18", "第六章": "20", "第七章": "22", "第八章": "25", "第九章": "27",
         "第十章": "29", "第十一章": "32", "第十二章": "33", "参考资料": "34"}

doc = Document(SRC)

def visible(p):
    """Concatenate w:t text only, tabs excluded."""
    return "".join(t.text or "" for t in p._element.iter(qn('w:t')))

updated = 0
for p in doc.paragraphs:
    el = p._element
    if not el.findall('.//' + qn('w:tab')):
        continue
    label = visible(p).strip()
    # 去掉尾部可能残留的数字
    clean = label.rstrip("0123456789").strip()
    page = None
    for k, v in PAGES.items():
        if clean.startswith(k):
            page = v
            break
    if page is None:
        continue
    # 最后一段 run 是页码：清空非首 run 的文本，并把首 run 设为 标题+Tab 结构
    runs = p.runs
    if not runs:
        continue
    runs[0].text = clean
    # 若首 run 内没有 tab 元素，则补一个
    r0 = runs[0]._element
    if not r0.findall(qn('w:tab')):
        tab = OxmlElement('w:tab')
        r0.append(tab)
    # 找到承载页码的 run，或新建
    page_run = None
    for r in runs[1:]:
        if (r.text or "").strip().isdigit():
            page_run = r
            break
    if page_run is None:
        page_run = p.add_run("")
    page_run.text = page
    for r in runs[1:]:
        if r is not page_run:
            r.text = ""
    print(f"  {clean[:22]:26s} -> {page}")
    updated += 1

doc.save(SRC)
print("updated entries:", updated)
