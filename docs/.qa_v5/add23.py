# -*- coding: utf-8 -*-
import copy
from docx import Document
from docx.text.paragraph import Paragraph
SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)
def set_text(p, t):
    runs = p.runs
    runs[0].text = t
    for r in runs[1:]:
        r.text = ""; r._element.getparent().remove(r._element)
anchor = None
for p in doc.paragraphs:
    if p.text.strip() == "2.3 核心验证指标与决策门":
        anchor = p; break
if anchor is not None:
    model = None
    for p in doc.paragraphs:
        if p.style.name == "Normal" and len(p.text) > 60:
            model = p; break
    el = copy.deepcopy(model._element)
    anchor._element.addnext(el)
    np = Paragraph(el, anchor._parent)
    set_text(np, "以下七项为进入正式实验前预设的验证方向、判定标准与决策门。"
                 "判定数据来源包括任务日志、情境标注集、产品分析记录与试点合作材料，"
                 "样本量在实验前依据预实验效应量做统计功效分析后确定。")
    print("added 2.3 lead-in")
doc.save(SRC)