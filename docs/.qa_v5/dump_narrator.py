# -*- coding: utf-8 -*-
"""Locate the remaining 'above-mentioned' narrator phrases and the growth bullets."""
import re
from docx import Document
from docx.oxml.ns import qn

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def ptxt(p):
    return "".join(t.text or "" for t in p._element.iter(qn('w:t')))

print("=== 含“上述 / 这一X / 本章用于 / 该原则用于 / 这也是…的原因 / 项目要验证” ===")
for i, p in enumerate(doc.paragraphs):
    t = ptxt(p).strip()
    if re.search(r"上述|这一(思路|路径|原则|设计|判断|差别|命题|方式)|本章用于|该原则用于|这也是.{0,14}的原因|项目要验证|项目把这条命题", t):
        print(f"P{i:03d} [{p.style.name}] {t[:210]}")
        print()

print("=== 7.3 增长与留存 全部条目 ===")
hit = False
for i, p in enumerate(doc.paragraphs):
    t = ptxt(p).strip()
    if t.startswith("7.3"):
        hit = True
    if hit:
        print(f"P{i:03d} [{p.style.name}] {t[:150]}")
        if t.startswith("7.4"):
            break
