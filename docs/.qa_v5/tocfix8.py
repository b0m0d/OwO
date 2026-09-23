import docx
p=r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v6.docx"
vals=["3","5","9","11","13","18","20","23","25","27","29","31","33","34"]
d=docx.Document(p); toc=d.tables[1]
assert len(toc.rows)==len(vals)
for i,row in enumerate(toc.rows):
    par=row.cells[1].paragraphs[0]
    if par.runs:
        par.runs[0].text=vals[i]
        for r in par.runs[1:]: r.text=""
    else: par.add_run(vals[i])
d.save(p)
d2=docx.Document(p)
for row in d2.tables[1].rows: print(row.cells[0].text,"->",row.cells[1].text)
