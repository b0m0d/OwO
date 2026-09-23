# -*- coding: utf-8 -*-
"""Consistency sweep 1: market capacity logic, table note, two task chains, revenue basis,
usage-based revenue scope, timeline, naming."""
import copy
from docx import Document
from docx.text.paragraph import Paragraph

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""; r._element.getparent().remove(r._element)

def find(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)

def E(prefix, text):
    set_text(find(prefix), text)

def drop(prefix):
    p = find(prefix); p._element.getparent().remove(p._element)
    print("  dropped:", prefix[:40])

def model_p():
    for p in doc.paragraphs:
        if p.style.name == "Normal" and len(p.text) > 60:
            return p

def insert_after(anchor, texts):
    prev = anchor
    for t in texts:
        el = copy.deepcopy(model_p()._element)
        prev._element.addnext(el)
        np = Paragraph(el, prev._parent)
        set_text(np, t); prev = np

# ---- 1) 6.1 首年规模口径与 6.2 统一 ----
E("Cuttle 采用自下而上的市场进入测算",
  "首年不以宏观市场份额推算用户规模，而依据团队的产品验证与服务承载能力配置资源，"
  "将稳定服务规模控制在 200 人以内，并以真实试点数据逐步校准转化率、留存率与组织复制效率。"
  "中长期市场空间参考我国生成式人工智能用户规模（6.02 亿）、软件业务收入（15.48 万亿元）等公开数据，"
  "用于判断产品的长期扩展潜力。")
# ---- 2) 删除表 6 的旧比例注释 ----
drop("注：首年容量按服务能力配置，不采用样本池乘以转化率的推算法")

# ---- 3) 两条核心任务链统一 ----
for p in doc.paragraphs:
    t = p.text
    if "第一阶段只在团队能够直接进入的场景中形成样板" in t:
        set_text(p, "第一阶段围绕两条核心任务链形成试点样板：跨来源资料取数、口径核验到报告交付；"
                    "代码问题到修改、测试与可审阅差异。每个试点围绕任务完成率、时间节省、"
                    "安全理解与连续使用记录证据，并在不少于三个试点单位中验证复制条件。"
                    "第二阶段以开源 SDK、技能模板与开发者案例扩大高价值用户。"
                    "第三阶段依据真实采购需求提供团队许可与私有部署，合同范围绑定可交付能力、数据边界与支持成本。")
        print("  7.2 任务链统一")
for p in doc.paragraphs:
    t = p.text
    if "第七至十二个月完成兼容矩阵与 Origin、Return 适配器标准化" in t:
        set_text(p, "第七至十二个月完成兼容矩阵与 Origin、Return 适配器标准化，完成三组对照实验与消融实验并标定路由阈值，"
                    "在两条核心任务链和不少于三个试点单位中验证连续四周留存与安全门槛。")
        print("  8.3 任务链统一")

# ---- 4) 10.1 收入确认口径 ----
E("本章用于检验商业模型能否形成可持续经营",
  "本章用于检验商业模型能否形成可持续经营，不代表项目已经取得收入、融资或用户规模。"
  "收入按以下方式确认：个人订阅收入按年度平均在付用户数乘年费 240 元测算，"
  "表中年末在付用户数用于描述用户规模；团队许可按 4 万元每团队每年计价，自第二年起计入；"
  "私有部署按 18 万元每项目、以验收确认；每年在付用户由上年在付规模的 60% 续订加当年新增构成。"
  "本轮预测仅计入基础订阅、团队许可与私有部署收入，超额智能体调用收入及其对应的推理成本暂不计入，"
  "以保持测算审慎。三种情景分别设定用户与客户数量，按同一单价假设计算收入；"
  "第一年的规模由第六章的验证容量约束，个人付费用户不超过 200 人，"
  "三种情景分别按 80、120 与 180 人测算。全部单价与转化率均属待验证假设，验证方式与门槛见表 7。")

# ---- 5) 7.1 时间线 ----
E("Cuttle 的前两年聚焦个人专业版订阅与团队许可两种模式",
  "Cuttle 前 18 个月聚焦个人专业版订阅与团队许可两种模式。免费版承担体验与传播，提供低风险表达与少量任务；"
  "个人专业版面向高频知识工作者，是收入的主要来源；团队版面向需要统一策略、审计与共享能力的组织。"
  "按实施路线，第一年确认个人订阅收入，团队许可收入自第二年起计入；"
  "第 19 至 24 个月在组织采购需求得到验证后启动首个私有部署项目。"
  "技能生态属于远期方向，前两年不设为独立收入来源。")

# ---- 6) 命名：三维任务决策 / 交互与执行决策 ----
for p in doc.paragraphs:
    t = p.text.strip()
    if t == "4.2 核心创新二 三维自治决策与受控升级":
        set_text(p, "4.2 核心创新二 三维任务决策与受控自治升级")
        print("  4.2 重命名")
    if t == "3.3 交互与自治等级":
        set_text(p, "3.3 交互与执行决策")
        print("  3.3 重命名")
for p in doc.paragraphs:
    if "三维自治决策" in p.text:
        set_text(p, p.text.replace("三维自治决策", "三维任务决策"))
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            if "三维自治决策" in c.text:
                for pp in c.paragraphs:
                    set_text(pp, pp.text.replace("三维自治决策", "三维任务决策"))
            if "三维决策 Macro-F1" in c.text:
                for pp in c.paragraphs:
                    set_text(pp, pp.text.replace("三维决策 Macro-F1", "三维任务决策 Macro-F1"))

# ---- 7) 跨设备标注为扩展 ----
E("Return 是产品完成闭环的关键" if False else "Return to Origin 是产品完成闭环的关键",
  "Return to Origin 是产品完成闭环的关键。Word 中产生的任务返回为可审阅段落、批注或文档副本；"
  "聊天中产生的任务返回为候选回复，发送权仍由用户掌握；IDE 中产生的任务返回为差异、测试结果和回滚点。"
  "扩展场景下，跨设备任务返回为带状态与证据的任务卡。结果格式由原始应用、任务契约与风险等级共同决定。")

doc.save(SRC)
print("sweep 1 done")