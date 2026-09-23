# -*- coding: utf-8 -*-
"""Layout polish: compact 表 7 and 表 10 so the following sections flow without near-empty pages."""
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def first_cells(t):
    return [c.text.strip() for c in t.rows[0].cells]

def compact(t, top=40, side=70, line=0.94):
    for r in t.rows:
        for c in r.cells:
            tcPr = c._tc.get_or_add_tcPr()
            old = tcPr.find(qn('w:tcMar'))
            if old is not None:
                tcPr.remove(old)
            mar = OxmlElement('w:tcMar')
            for side_name, val in (('top', top), ('start', side), ('bottom', top), ('end', side)):
                e = OxmlElement('w:' + side_name)
                e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa')
                mar.append(e)
            tcPr.append(mar)
            for p in c.paragraphs:
                p.paragraph_format.space_before = 0
                p.paragraph_format.space_after = 0
                p.paragraph_format.line_spacing = line

for t in doc.tables:
    fc = first_cells(t)
    if fc[:3] == ["产品版本", "目标用户", "规划定价"]:
        compact(t); print("  compacted 表 7")
    if fc[:2] == ["风险", "概率"]:
        compact(t); print("  compacted 表 10")

doc.save(SRC)
print("layout polish done")