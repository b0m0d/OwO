# -*- coding: utf-8 -*-
from docx import Document
SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)
PAGES = {"项目摘要":"3","第一章":"5","第二章":"9","第三章":"13","第四章":"15","第五章":"21",
         "第六章":"23","第七章":"26","第八章":"28","第九章":"31","第十章":"32",
         "第十一章":"35","第十二章":"38","参考资料":"39"}
def set_cell(cell, text):
    ps = cell.paragraphs
    runs = ps[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""; r._element.getparent().remove(r._element)
    else:
        ps[0].add_run(text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)
t = doc.tables[1]
for row in t.rows:
    left = row.cells[0].text.strip()
    for k, v in PAGES.items():
        if left.startswith(k):
            set_cell(row.cells[1], v); break
doc.save(SRC)
d = Document(SRC)
for row in d.tables[1].rows:
    print("  ", row.cells[0].text.strip()[:26], "->", row.cells[1].text.strip())