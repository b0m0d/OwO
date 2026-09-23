# -*- coding: utf-8 -*-
"""Center EVERY cell paragraph in EVERY table (no index skipping) + vAlign center."""
import glob, os, shutil
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCS = os.path.join(r"T:\创新创业\OwO-master\docs")
SRC = [x for x in glob.glob(os.path.join(DOCS, "Cuttle*v8.docx")) if "预览" not in x][0]
BK = os.path.join(DOCS, ".qa_v5", "v8_before_centerall2.docx")
shutil.copy2(SRC, BK)
print("backup ->", os.path.basename(BK))

doc = Document(SRC)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

def in_cell(el):
    p = el.getparent()
    while p is not None:
        if p.tag == qn('w:tc'):
            return True
        if p.tag == qn('w:body'):
            return False
        p = p.getparent()
    return False

def pPr_of(p):
    pPr = p.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        p.insert(0, pPr)
    return pPr

def get_or_add(pPr, tag):
    el = pPr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        # keep schema order: insert before rPr / sectPr / pPrChange if present
        ref = None
        for after in ('w:rPr', 'w:sectPr', 'w:pPrChange'):
            cand = pPr.find(qn(after))
            if cand is not None:
                ref = cand
                break
        if ref is not None:
            ref.addprevious(el)
        else:
            pPr.append(el)
    return el

npar = 0
for p in doc.element.body.iter(qn('w:p')):
    if not in_cell(p):
        continue
    pPr = pPr_of(p)
    jc = get_or_add(pPr, 'w:jc')
    jc.set(qn('w:val'), 'center')
    # zero indents
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = get_or_add(pPr, 'w:ind')
    ind.set(qn('w:firstLine'), '0')
    ind.set(qn('w:left'), '0')
    ind.set(qn('w:right'), '0')
    sp = pPr.find(qn('w:spacing'))
    if sp is not None:
        sp.set(qn('w:before'), '0')
        sp.set(qn('w:after'), '0')
        if sp.get(qn('w:line')) is None:
            pass
    npar += 1

ntc = 0
for tc in doc.element.body.iter(qn('w:tc')):
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr'); tc.insert(0, tcPr)
    va = tcPr.find(qn('w:vAlign'))
    if va is None:
        va = OxmlElement('w:vAlign'); tcPr.append(va)
    va.set(qn('w:val'), 'center')
    ntc += 1

doc.save(SRC)
print("cell paragraphs centered:", npar, "| cells vAlign=center:", ntc)

# verify
doc2 = Document(SRC)
tot = cen = 0
for t in doc2.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                tot += 1
                pPr = p._p.find(qn('w:pPr'))
                j = pPr.find(qn('w:jc')) if pPr is not None else None
                if j is not None and j.get(qn('w:val')) == 'center':
                    cen += 1
print("VERIFY: cell paragraphs=%d  explicit-center=%d  missing=%d" % (tot, cen, tot - cen))
