# -*- coding: utf-8 -*-
from docx import Document
SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def set_cell(cell, text):
    ps = cell.paragraphs
    runs = ps[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
            r._element.getparent().remove(r._element)
    else:
        ps[0].add_run(text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

t = doc.tables[20]
for r in t.rows:
    lab = r.cells[0].text.strip()
    if lab.startswith("经营成本"):
        set_cell(r.cells[1], "40 / 50 / 65 万元")
        set_cell(r.cells[2], "55 / 105 / 190 万元")
        set_cell(r.cells[3], "80 / 190 / 260 万元")
    elif lab.startswith("经营结果 保守"):
        set_cell(r.cells[1], "负 32.8 万元"); set_cell(r.cells[2], "负 26.2 万元"); set_cell(r.cells[3], "负 16.8 万元")
        set_cell(r.cells[4], "三年均未转正；该情景用于观察下限")
    elif lab.startswith("经营结果 基准"):
        set_cell(r.cells[1], "负 40.4 万元"); set_cell(r.cells[2], "负 56.2 万元"); set_cell(r.cells[3], "正 9.2 万元")
        set_cell(r.cells[4], "第三年进入盈亏平衡上方")
    elif lab.startswith("经营结果 进取"):
        set_cell(r.cells[1], "负 50.6 万元"); set_cell(r.cells[2], "负 84.0 万元"); set_cell(r.cells[3], "正 92.0 万元")
        set_cell(r.cells[4], "第二年投入最大，第三年转正")
doc.save(SRC)
print("--- final finance table ---")
d = Document(SRC)
for ri, r in enumerate(d.tables[20].rows):
    print(f"R{ri}: " + " | ".join(c.text.strip()[:60] for c in r.cells))