# -*- coding: utf-8 -*-
"""Batch C: market table credibility, remove redundant competitor table, page break for ch.11."""
import copy
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)

def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""; r._element.getparent().remove(r._element)

def set_cell(cell, text):
    ps = cell.paragraphs
    set_text(ps[0], text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

def find(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)

def first_cells(t):
    return [c.text.strip() for c in t.rows[0].cells]

def get_table(sig, ncols=None):
    for t in doc.tables:
        fc = first_cells(t)
        if fc[:len(sig)] == sig and (ncols is None or len(fc) == ncols):
            return t
    raise KeyError(str(sig))

# ---- 表 6 首年容量：去掉伪精确链条 ----
t = get_table(["口径", "测算对象", "测算依据"], 4)
NEW6 = [
 ["口径", "测算对象", "测算方式", "结果与用途"],
 ["首年可交付容量", "团队可稳定服务并完整观察的用户数",
  "按产品验证与服务承载能力配置，不按市场规模推算；首批试点以团队可直接触达的组织、合作单位与知识工作场景为来源",
  "首年按不超过 200 名稳定用户配置服务能力，并以此约束付费用户规模"],
 ["验证与实验容量", "支撑对照实验与纵向观察的样本规模",
  "纵向核心组 15 至 25 人全程跟踪；独立验证组另招 20 至 30 人",
  "50 至 80 人形成深度数据，其中 35 至 55 人进入正式实验"],
 ["中长期市场空间", "产品扩展潜力的参考量级",
  "引用我国生成式人工智能用户规模、软件业务收入等公开数据作为行业背景",
  "用于判断长期扩展潜力，随试点转化率、留存率与复制效率持续校准"],
]
for ri, vals in enumerate(NEW6):
    for ci, v in enumerate(vals):
        set_cell(t.rows[ri].cells[ci], v)

# 表 6 表注
model = None
for p in doc.paragraphs:
    if p.style.name == "Normal" and len(p.text) > 60:
        model = p; break
grid_note = copy.deepcopy(model._element)
t._element.addnext(grid_note)
note = Paragraph(grid_note, t._parent)
set_text(note, "注：首年容量按服务能力配置，不采用样本池乘以转化率的推算法；"
               "上述比例区间来自前期访谈与任务日记，正式试点开始后由真实转化数据替换。")

# ---- 删除表 7（与图 8 重复），正文补一段差异总结 ----
for cand in list(doc.tables):
    if first_cells(cand)[:2] == ["类别", "代表产品"]:
        cand._element.getparent().remove(cand._element)
        print("  removed 表 7 产品竞争定位")
        break
for p in list(doc.paragraphs):
    if p.text.strip() == "表 7 产品竞争定位与任务生命周期覆盖":
        p._element.getparent().remove(p._element)
        print("  removed its caption")

anchor = find("输入法产品正在增加续写")
el = copy.deepcopy(model._element)
anchor._element.addnext(el)
np = Paragraph(el, anchor._parent)
set_text(np, "按任务生命周期的六个环节（意图捕获、情境构建、受控执行、结果验证、返回原处、治理审计）比较，"
             "输入法产品在意图捕获与表达效率上领先，通用助手与工作台在受控执行与验证上更完整，"
             "而把六个环节串成同一条链路、并让成果回到需求产生位置的方案尚未形成。"
             "Cuttle 的设计目标正是补齐这条链路，并以跨应用适配与可靠性积累作为差异化基础。")

# ---- 第十一章从新页开始 ----
h11 = find("第十一章 风险管理")
h11.paragraph_format.page_break_before = True

doc.save(SRC)
print("batch C done. tables:", len(doc.tables))