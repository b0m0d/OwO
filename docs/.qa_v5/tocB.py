# -*- coding: utf-8 -*-
from docx import Document
SRC = r"T:\创新创业\OwO-master\docs\Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
doc = Document(SRC)
PAGES = {"项目摘要":"3","第一章":"4","第二章":"7","第三章":"10","第四章":"13","第五章":"18",
         "第六章":"19","第七章":"22","第八章":"24","第九章":"27","第十章":"29",
         "第十一章":"32","第十二章":"34","参考资料":"35"}
def set_cell(cell, text):
    ps = cell.paragraphs
    runs = ps[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""; r._element.getparent().remove(r._element)
    else:
        ps[0].add_run(text)
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)
for row in doc.tables[1].rows:
    left = row.cells[0].text.strip()
    for k, v in PAGES.items():
        if left.startswith(k):
            set_cell(row.cells[1], v); break
doc.save(SRC)
print("TOC updated")