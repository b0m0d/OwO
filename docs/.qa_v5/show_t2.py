import glob, os
from docx import Document

DOCS = os.path.join(r"T:\创新创业\OwO-master\docs")
SRC = [x for x in glob.glob(os.path.join(DOCS, "Cuttle*v8.docx")) if "预览" not in x][0]
print("file:", os.path.basename(SRC))
d = Document(SRC)
for ti, t in enumerate(d.tables):
    fc = [c.text.strip() for c in t.rows[0].cells]
    if "验证方向" in " ".join(fc) or "分组" == fc[0]:
        print("table index", ti, "cols", len(t.columns), "rows", len(t.rows))
        print("header:", fc)
        for r in t.rows[1:]:
            print("  ", " | ".join(c.text.strip()[:30] for c in r.cells))
        break
else:
    print("表 2 not located; dumping all table headers:")
    for ti, t in enumerate(d.tables):
        print(ti, [c.text.strip()[:12] for c in t.rows[0].cells])
