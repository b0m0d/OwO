from pathlib import Path
import sys
from copy import deepcopy

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUTPUT = ROOT / "builGoal" / "Cuttle——新一代智能输入法-参赛商业计划书-图表完善版-v4.docx"
SOURCE_LOGO = Path(r"C:\Users\23843\AppData\Local\Temp\codex-clipboard-ffc7606e-74c1-4a51-92b2-2a357f7ca9b3.png")
ASSET_DIR = HERE / "cuttle-assets"
LOGO = ASSET_DIR / "cuttle-logo.png"
COVER = ASSET_DIR / "cuttle-cover-layout.png"
PAGE_BG = ASSET_DIR / "cuttle-page-background-color-v6-no-mascot.png"
FIGURE_DIR = ASSET_DIR / "figures"
FIGURE_FILES = {
    "图1-1": "fig1-1-background-evolution.png",
    "图1-2": "fig1-2-value-loop.png",
    "图2-1": "fig2-1-industry-trend.png",
    "图2-2": "fig2-2-pain-points.png",
    "图2-3": "fig2-3-pest.png",
    "图2-4": "fig2-4-swot.png",
    "图2-5": "fig2-5-market-growth.png",
    "图3-1": "fig3-1-product-layers.png",
    "图3-2": "fig3-2-user-journey.png",
    "图3-3": "fig3-3-tech-architecture.png",
    "图4-1": "fig4-1-business-model.png",
    "图5-1": "fig5-1-growth-loop.png",
    "图6-1": "fig6-1-fund-allocation.png",
    "图6-2": "fig6-2-financial-forecast.png",
    "图7-1": "fig7-1-organization.png",
    "图8-1": "fig8-1-risk-matrix.png",
    "图9-1": "fig9-1-roadmap.png",
    "图10-1": "fig10-1-social-impact.png",
}

sys.path.insert(0, str(HERE))
import build_owo_plan as base

BLACK = "000000"
MID = "666666"
LIGHT = "E6E6E6"
PALE = "F7F7F7"
PAGE_DXA = 9360


def clear_paragraph_content(paragraph):
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    p_pr = p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_full_page_background(header, image_path, page_width, page_height, alt_text, label=None):
    """Add one high-resolution page-sized image behind header/body content."""
    paragraph = header.paragraphs[0]
    clear_paragraph_content(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(9)

    shape = paragraph.add_run().add_picture(
        str(image_path), width=page_width, height=page_height
    )
    base.add_alt_text(shape, alt_text)
    inline = shape._inline

    anchor = OxmlElement("wp:anchor")
    for key, value in {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": "0",
        "behindDoc": "1",
        "locked": "0",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }.items():
        anchor.set(key, value)

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)

    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "page")
    pos_h = OxmlElement("wp:posOffset")
    pos_h.text = "0"
    position_h.append(pos_h)
    anchor.append(position_h)

    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "page")
    pos_v = OxmlElement("wp:posOffset")
    pos_v.text = "0"
    position_v.append(pos_v)
    anchor.append(position_v)

    extent = OxmlElement("wp:extent")
    extent.set("cx", str(int(page_width)))
    extent.set("cy", str(int(page_height)))
    anchor.append(extent)

    effect = OxmlElement("wp:effectExtent")
    for edge in ("l", "t", "r", "b"):
        effect.set(edge, "0")
    anchor.append(effect)
    anchor.append(OxmlElement("wp:wrapNone"))
    anchor.append(inline.docPr)
    c_nv = inline.find(qn("wp:cNvGraphicFramePr"))
    if c_nv is not None:
        anchor.append(c_nv)
    anchor.append(inline.graphic)
    inline.getparent().replace(inline, anchor)

    if label:
        run = paragraph.add_run(label)
        set_font(run, size=8.5, bold=True, east="微软雅黑")


def font(path, size):
    p = Path(path)
    if p.exists():
        return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()


def make_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    if not SOURCE_LOGO.exists():
        raise FileNotFoundError(SOURCE_LOGO)
    Image.open(SOURCE_LOGO).convert("RGBA").save(LOGO)

    w, h = 2480, 3508
    canvas = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(canvas)
    for x in range(-500, w + 500, 180):
        d.line([(x, h * 0.62), (x + 900, h)], fill=(238, 238, 238), width=4)
    for y in range(int(h * 0.66), h, 120):
        d.arc([80, y - 420, w - 80, y + 320], 190, 350, fill=(232, 232, 232), width=3)
    d.rectangle([0, 0, 42, h], fill=(20, 20, 20))
    d.text((180, 170), "中国国际大学生创新大赛（2026）", font=font(r"C:\Windows\Fonts\msyhbd.ttc", 58), fill="black")
    d.line([(180, 270), (2300, 270)], fill=(0, 0, 0), width=5)

    logo = Image.open(LOGO).convert("RGBA")
    logo = ImageOps.contain(logo, (920, 920))
    canvas.paste(logo, ((w - logo.width) // 2, 480), logo)

    title_font = font(r"C:\Windows\Fonts\msyhbd.ttc", 160)
    sub_font = font(r"C:\Windows\Fonts\msyhbd.ttc", 82)
    small_font = font(r"C:\Windows\Fonts\msyh.ttc", 46)
    title = "Cuttle"
    box = d.textbbox((0, 0), title, font=title_font)
    d.text(((w - (box[2] - box[0])) // 2, 1450), title, font=title_font, fill="black")
    subtitle = "新一代智能输入法"
    box = d.textbbox((0, 0), subtitle, font=sub_font)
    d.text(((w - (box[2] - box[0])) // 2, 1660), subtitle, font=sub_font, fill="black")
    d.line([(570, 1800), (1910, 1800)], fill=(0, 0, 0), width=4)
    tagline = "让输入理解情境，让表达自然抵达，让任务高效落地"
    box = d.textbbox((0, 0), tagline, font=small_font)
    d.text(((w - (box[2] - box[0])) // 2, 1870), tagline, font=small_font, fill=(45, 45, 45))

    meta_font = font(r"C:\Windows\Fonts\msyh.ttc", 44)
    meta_bold = font(r"C:\Windows\Fonts\msyhbd.ttc", 44)
    rows = [
        ("项目组别", "本科生组"),
        ("所属领域", "信息技术服务业"),
        ("参赛学校", "西南大学"),
        ("项目负责人", "杨俊熙"),
    ]
    y = 2680
    for label, value in rows:
        d.text((420, y), label + "：", font=meta_bold, fill="black")
        d.text((760, y), value, font=meta_font, fill="black")
        d.line([(420, y + 70), (2050, y + 70)], fill=(210, 210, 210), width=2)
        y += 150
    canvas.save(COVER, quality=95, dpi=(300, 300))

def set_font(run, size=10.5, bold=False, italic=False, color=BLACK, latin="Times New Roman", east="宋体"):
    base.set_run_font(run, latin=latin, east_asia=east, size=size, color=color, bold=bold, italic=italic)


def set_keep_with_next(paragraph, value=True):
    ppr = paragraph._p.get_or_add_pPr()
    tag = ppr.find(qn("w:keepNext"))
    if value and tag is None:
        ppr.append(OxmlElement("w:keepNext"))
    elif not value and tag is not None:
        ppr.remove(tag)


def add_page_field(paragraph):
    base.add_field(paragraph, "PAGE")


def set_cell_text(cell, text, size=9.2, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text))
    set_font(r, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    base.set_cell_margins(cell, top=90, bottom=90, start=110, end=110)


def build():
    make_assets()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.25)
    section.left_margin = Cm(2.45)
    section.right_margin = Cm(2.25)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.3)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.45

    for name, size, before, after, align in [
        ("Heading 1", 18, 20, 12, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 14, 14, 7, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 11.5, 10, 5, WD_ALIGN_PARAGRAPH.LEFT),
    ]:
        st = styles[name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.alignment = align
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.15

    num_ids = base.create_numbering(doc)
    current_decimal = num_ids["decimal"]

    if not PAGE_BG.exists():
        raise FileNotFoundError(PAGE_BG)
    add_full_page_background(
        section.header,
        PAGE_BG,
        section.page_width,
        section.page_height,
        "Cuttle商业计划书彩色页眉页脚统一背景",
        label="Cuttle——新一代智能输入法",
    )
    # The cover is a standalone full-page visual; it must not inherit the body background.
    clear_paragraph_content(section.first_page_header.paragraphs[0])

    footer = section.footer
    fp = footer.paragraphs[0]
    clear_paragraph_content(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    r = fp.add_run("第 ")
    set_font(r, size=8.5)
    add_page_field(fp)
    r = fp.add_run(" 页")
    set_font(r, size=8.5)

    def para(text="", bold_prefix=None, size=10.5, italic=False, first_indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.first_line_indent = Cm(0.74) if first_indent else Cm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.45
        if bold_prefix and text.startswith(bold_prefix):
            r = p.add_run(bold_prefix)
            set_font(r, size=size, bold=True)
            r = p.add_run(text[len(bold_prefix):])
            set_font(r, size=size, italic=italic)
        else:
            r = p.add_run(text)
            set_font(r, size=size, italic=italic)
        return p

    def lead(label, text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.right_indent = Cm(0.2)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.35
        base.set_paragraph_border(p, color=BLACK, size=12, space=6)
        r = p.add_run(label + "  ")
        set_font(r, size=10.5, bold=True, east="黑体")
        r = p.add_run(text)
        set_font(r, size=10.5)
        return p

    def bullet(text, level=0):
        p = doc.add_paragraph()
        base.apply_numbering(p, num_ids["bullet"])
        if level:
            p.paragraph_format.left_indent = Cm(1.0 + level * 0.5)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(text)
        set_font(r, size=10.3)
        return p

    def number(text, restart=False):
        nonlocal current_decimal
        if restart:
            current_decimal = base.clone_num_id(doc, num_ids["decimal"])
        p = doc.add_paragraph()
        base.apply_numbering(p, current_decimal)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(text)
        set_font(r, size=10.3)
        return p

    def h1(text, new_page=True):
        if new_page and len(doc.paragraphs) > 0:
            doc.add_page_break()
        return doc.add_heading(text, level=1)

    def h2(text):
        return doc.add_heading(text, level=2)

    def h3(text):
        return doc.add_heading(text, level=3)

    def caption(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(7)
        r = p.add_run(text)
        set_font(r, size=9, bold=True, east="宋体")
        set_keep_with_next(p, False)
        return p

    def figure_box(label, height_cm=5.1):
        key = label.split()[0]
        image_path = FIGURE_DIR / FIGURE_FILES[key]
        if not image_path.exists():
            raise FileNotFoundError(image_path)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_together = True
        set_keep_with_next(p, True)
        shape = p.add_run().add_picture(str(image_path), width=Cm(16.0))
        base.add_alt_text(shape, label)
        caption(label)
        return p

    def table(headers, rows, widths, aligns=None, font_size=9.1):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        base.set_table_geometry(t, widths, indent_dxa=110)
        base.set_table_borders(t, color=BLACK, size=8)
        base.set_repeat_table_header(t.rows[0])
        base.set_row_cant_split(t.rows[0])
        for i, head in enumerate(headers):
            set_cell_text(t.rows[0].cells[i], head, size=9.2, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        for row_data in rows:
            row = t.add_row()
            base.set_row_cant_split(row)
            for i, value in enumerate(row_data):
                align = aligns[i] if aligns else WD_ALIGN_PARAGRAPH.CENTER
                set_cell_text(row.cells[i], value, size=font_size, align=align)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        return t

    # Cover is deliberately a single full-page visual.
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(0)
    cover_shape = cp.add_run().add_picture(str(COVER), width=Cm(17.9), height=Cm(24.6))
    base.add_alt_text(cover_shape, "Cuttle新一代智能输入法参赛商业计划书封面")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("目录")
    set_font(r, size=20, bold=True, east="黑体")

    toc = [
        ("第一章  项目概述", "4"),
        ("1.1 项目名称与所属领域", "4"), ("1.2 项目背景", "4"), ("1.3 项目介绍", "5"), ("1.4 核心价值", "5"), ("1.5 当前基础与真实性边界", "6"),
        ("第二章  市场分析", "8"),
        ("2.1 行业与政策环境", "8"), ("2.2 市场痛点识别", "8"), ("2.3 目标用户与需求层级", "9"), ("2.4 竞争格局", "10"), ("2.5 PEST分析", "11"), ("2.6 SWOT分析", "11"), ("2.7 市场空间测算", "12"),
        ("第三章  产品体系与项目创新", "14"),
        ("3.1 产品总体形态", "14"), ("3.2 核心功能体系", "14"), ("3.3 典型场景", "16"), ("3.4 用户旅程", "16"), ("3.5 技术路线", "17"), ("3.6 项目创新点", "18"), ("3.7 安全与隐私", "19"),
        ("第四章  商业模式", "20"),
        ("4.1 盈利模式", "20"), ("4.2 产品版本与定价", "20"), ("4.3 收入来源", "21"), ("4.4 成本结构", "21"), ("4.5 合作伙伴策略", "21"),
        ("第五章  市场营销策略", "23"),
        ("5.1 市场定位", "23"), ("5.2 产品与渠道策略", "23"), ("5.3 宣传与用户增长", "24"), ("5.4 留存与口碑机制", "26"),
        ("第六章  财务分析", "27"),
        ("6.1 测算口径", "27"), ("6.2 融资计划", "27"), ("6.3 三年财务预测", "28"), ("6.4 敏感性与盈亏平衡", "30"),
        ("第七章  组织与管理", "31"),
        ("7.1 核心团队成员", "31"), ("7.2 指导教师", "31"), ("7.3 组织架构", "33"), ("7.4 协作与质量管理", "33"),
        ("第八章  风险防控", "35"), ("8.1 风险总述", "35"), ("8.2 风险分析与应对", "35"),
        ("第九章  战略规划", "37"), ("9.1 战略目标", "37"), ("9.2 实施路径", "37"), ("9.3 知识产权与生态", "38"),
        ("第十章  社会效益", "39"), ("10.1 生产力与教育价值", "39"), ("10.2 可信人工智能价值", "39"), ("10.3 就业与产业带动", "40"),
    ]
    for title, page in toc:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2.2)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run(title + "\t" + page)
        set_font(r, size=9.7, bold=("章" in title and "." not in title))

    h1("第一章  项目概述")
    h2("1.1 项目名称与所属领域")
    table(["项目要素", "具体内容"], [
        ("项目名称", "Cuttle——新一代智能输入法"),
        ("所属领域", "信息技术服务业"),
        ("项目定位", "面向个人电脑场景的情境感知智能输入与任务协作系统"),
        ("目标用户", "高校学生、教师与研究人员、内容与办公人员、软件开发者及高频电脑用户"),
        ("核心主张", "在用户授权范围内理解当前情境，辅助自然表达，并把复杂请求衔接到可审阅、可控制的任务执行流程"),
    ], [1900, 7460])
    lead("一句话定义", "Cuttle不是在传统输入法上增加一个聊天按钮，而是把输入入口、情境理解、桌面状态呈现和多步骤任务执行组织成一条连续的人机协作链路。")

    h2("1.2 项目背景")
    para("生成式人工智能正在从单次问答向办公、学习、内容创作和任务执行渗透。中国互联网络信息中心发布的第57次中国互联网络发展状况统计报告显示，截至2025年12月，我国生成式人工智能用户规模已达到6.02亿人，普及率为42.8%。用户基础快速扩大之后，竞争重点开始从“能否回答”转向“能否进入真实工作流、持续理解任务并稳定交付结果”。")
    para("现有通用人工智能助手通常以独立对话窗口为中心。用户在浏览器、文档、即时通信、代码编辑器等应用之间切换时，需要不断复制粘贴内容、重新描述背景和调整表达风格；当请求涉及文件处理、命令执行或跨应用操作时，又必须手工串联多个步骤。传统输入法虽然拥有最自然、最高频的文字入口，但大多仍停留在词频预测、语音转写和局部改写层面，难以理解“用户正在什么应用中、处理什么任务、下一步准备做什么”。")
    para("Cuttle选择从输入这一高频入口切入，将低摩擦表达与本地优先的任务智能体结合。简单请求在候选区或轻量面板中完成，复杂请求升级到可视化任务工作台；系统对文件、命令、网络和桌面动作执行最小权限控制、独立审批、过程审计和成果回退，从而形成“感知—表达—行动—复核”的完整闭环。")
    figure_box("图1-1  项目形成背景与用户需求演进图", 4.7)

    h2("1.3 项目介绍")
    para("Cuttle由四个彼此连接的产品层组成。第一层是智能输入入口，根据应用类型、任务主题和用户偏好提供补全、改写、翻译、快捷回复和专业词汇；第二层是情境感知引擎，在授权范围内读取应用、文件、界面语义和任务历史，形成可纠正的当前情境；第三层是桌面伙伴与工作台，把系统判断、建议、进度、审批和失败原因可视化；第四层是任务智能体与技能生态，将研究、文档、代码、表格、网页和桌面操作拆解为受控步骤并交付版本化成果。")
    para("产品采用分级交互。即时辅助只生成候选，不直接对外发送；主动建议以可忽略的轻量提示出现；长任务必须进入工作台，用户能够查看计划、进度、工具范围、审批节点和最终成果。分级设计既避免了为简单改写打开复杂系统，也防止低风险入口承担高风险自动化。")
    table(["产品层", "主要能力", "用户价值", "当前状态"], [
        ("智能输入入口", "场景化补全、改写、回复、术语与快捷指令", "减少重复表达与应用切换", "产品方案与历史输入法工程基础；统一版本待立项"),
        ("情境感知引擎", "应用、对象、任务主题、文件与界面语义", "建议与当前工作匹配", "模块化技术基础；需完成跨应用兼容矩阵"),
        ("桌面伙伴与工作台", "状态、建议、进度、审批、差异与恢复", "让AI可见、可纠正、可控制", "工作台处于产品化阶段，桌宠为后续交互形态"),
        ("任务智能体与技能", "任务拆解、工具调用、多智能体协同、成果交付", "把回答升级为任务闭环", "当前主要可验证工程主线"),
    ], [1600, 3000, 2500, 2260], font_size=8.7)

    h2("1.4 核心价值")
    h3("1.4.1 对用户的直接价值")
    bullet("减少上下文搬运：同一任务在资料阅读、文档写作、沟通和执行之间连续衔接。")
    bullet("降低表达成本：根据学习、办公、聊天、编程等场景给出相匹配的补全与语气。")
    bullet("提升复杂任务可控性：把目标、步骤、进度、审批、失败和成果放在统一界面。")
    bullet("保护个人数据与操作边界：优先本地处理，默认拒绝越权动作，敏感场景自动停止。")
    para("上述价值需要通过可观察指标持续验证。对于表达类任务，重点记录从产生意图到形成可发送文本所需的时间、修改次数和采纳率；对于复杂任务，重点记录任务成功率、人工接管次数、成果审阅时间和失败恢复效果；对于隐私与安全，重点检查越权拦截、敏感场景停止和审计记录完整性。只有当效率提升与风险控制同时成立时，Cuttle才真正形成区别于普通输入工具的长期价值。")
    h3("1.4.2 对机构与生态的价值")
    para("对于高校、实验室和小型团队，Cuttle可以把可复用任务流程沉淀为技能和模板，形成可审计的使用记录；对于开发者，开放技能接口使OCR、文档、代码、图表和业务系统能力可以按需接入；对于模型服务商，产品提供稳定、可替换的本地或兼容云端模型入口，不把用户锁定在单一模型厂商。")
    para("机构侧价值还体现在标准化与经验传承。课程团队可以把资料整理、报告检查和代码验收要求固化为模板，实验室可以将文献检索、证据抽取和研究记录形成规范流程，小型企业则可以把重复办公任务沉淀为受权限控制的技能。流程被保存后，新成员不必从零学习提示方式，管理者也能通过任务记录了解产出来源和异常环节。")
    figure_box("图1-2  Cuttle感知—表达—行动价值闭环图", 5.2)

    h2("1.5 当前基础与真实性边界")
    para("本计划书以当前仓库源码、测试报告和原项目申报材料为基础。现阶段最成熟、可量化的部分是本地优先的Agent工作台和任务执行底座，已经具备会话、工具、权限、审批、审计、任务编排、成果差异与恢复等技术能力。2026年8月31日的单智能体产品评测中，固定套件完成200次运行，成功195次，总成功率97.5%，覆盖代码、文档和研究类任务。该数据证明核心Agent执行面已形成较稳定的工程基础，但不能自动等同于完整桌面产品已通过发布验收。")
    para("智能输入法、统一桌面伙伴和输入法到工作台的连续交互属于本项目下一阶段的产品化重点。历史输入法工程与原申报书提供了产品方向和技术积累，但当前不把尚未完成的统一安装包、全应用适配、真实用户规模、商业收入和知识产权成果写成既成事实。比赛演示应使用经过复验的固定任务、固定安装包和可追溯数据，清楚区分“已实现”“可演示”“规划中”三类内容。")
    table(["证据层级", "可以陈述的内容", "提交前必须补强"], [
        ("已实现", "本地Agent循环、工具权限、审批审计、会话与成果管理、多智能体任务协同基础", "形成可安装版本、版本哈希、操作录像与复现实验"),
        ("可演示", "资料到报告、代码任务、文档整理等受控工作流", "在目标电脑完成全流程复测，固定演示数据"),
        ("规划中", "输入法统一入口、桌宠状态表达、跨应用情境连续、技能市场", "完成用户研究、交互原型和阶段验收指标"),
        ("情景测算", "市场规模、用户增长、定价、融资和三年财务", "由团队依据真实调研、合同和运营数据复核签字"),
    ], [1500, 3600, 4260])

    h1("第二章  市场分析")
    h2("2.1 行业与政策环境")
    para("国务院关于深入实施“人工智能+”行动的意见提出，到2027年新一代智能终端、智能体等应用普及率超过70%，到2030年超过90%，并鼓励发展提效型、陪伴型智能产品和智能助手新入口。输入法天然处于人机交互高频入口，本地情境感知与智能体任务执行又契合智能终端从工具向伙伴演进的方向。")
    para("行业供给大体分为四类：一是科大讯飞、百度、腾讯、字节等厂商的智能输入法，将大模型用于语音、改写和AI帮写；二是系统级桌面助手，拥有截图、文件搜索和系统设置等能力；三是Typeless、Wispr Flow等面向专业人群的语音与表达工具；四是开源Agent框架和自动化平台，强调工具调用与可扩展性。四类产品分别具备入口、生态、表达或执行优势，但“输入入口+情境理解+可控任务执行”的完整链路仍存在差异化空间。")
    para("行业正在从功能叠加转向系统协同。一方面，输入产品不断加入生成、翻译和语音能力；另一方面，智能体产品开始进入文件、浏览器和桌面工作流。二者的交汇点并不是把更多按钮放进输入框，而是建立能够跨应用保持任务连续、区分风险等级并交付可验证成果的底层机制。Cuttle将这一交汇点作为产品切入方向，优先解决个人电脑高频任务中的上下文断裂和执行不透明问题。")
    figure_box("图2-1  智能输入与桌面智能体行业演进趋势图", 5.2)

    h2("2.2 市场痛点识别")
    table(["核心痛点", "典型表现", "现有方式不足", "Cuttle应对"], [
        ("上下文反复丢失", "切换应用后重复解释背景、复制文件和粘贴内容", "通用对话以单次会话为中心", "在授权范围内保存任务主题、对象和成果关系"),
        ("表达与场景不匹配", "同一句话在论文、邮件、聊天中需反复改写", "传统输入法主要依赖词频与局部文本", "结合应用、关系和任务切换语气与术语"),
        ("复杂任务过程不可见", "用户不知道系统做了什么、为何失败", "自动化工具重执行、轻解释与恢复", "展示计划、进度、审批、证据、差异和回退"),
        ("自动化信任不足", "担心误删文件、误发消息和隐私泄露", "权限粒度不一，审批与执行耦合", "默认拒绝、独立审批、敏感熔断与审计"),
    ], [1600, 2600, 2500, 2660], font_size=8.6)
    para("四类痛点并非彼此独立。上下文丢失会增加表达成本，表达与场景不匹配又会迫使用户反复修改；当任务进一步升级为文件处理或跨应用操作时，过程不可见和权限不清将直接影响信任。项目因此不把某个单点功能作为唯一卖点，而是以连续工作流为单位设计产品和评估价值。")
    figure_box("图2-2  目标用户四类核心痛点结构图", 5.0)

    h2("2.3 目标用户与需求层级")
    para("项目采用“高频电脑任务优先”的用户选择原则，不追求在早期覆盖所有生成式人工智能用户。首批种子用户应拥有明确、重复、可验证的任务，例如课程资料到报告、科研文献整理、跨文档改写、代码修复和项目材料编制。这类用户能够提供任务日志、时间成本和成果质量对比，便于团队形成可复现的产品证据。")
    para("需求层级可分为基础表达、情境连续、任务执行和组织协作四层。基础表达解决补全、改写和翻译；情境连续减少跨应用重复说明；任务执行要求计划、工具和成果可控；组织协作进一步关注权限策略、模板复用和审计。项目早期先验证前两层的高频使用和第三层的核心闭环，再逐步进入团队级场景，避免在用户价值尚未明确时过早增加管理复杂度。")
    table(["用户群体", "高频任务", "核心诉求", "首个验证场景"], [
        ("高校学生", "资料阅读、课程报告、沟通与项目材料", "少切换、能追溯、可形成成果", "课程资料到结构化报告"),
        ("教师与研究人员", "文献、证据、教学材料和数据整理", "来源清晰、隐私可控、格式稳定", "多来源研究与文档交付"),
        ("内容与办公人员", "改写、汇总、表格、汇报和文件整理", "效率、格式一致、过程透明", "跨文档整理与版本审阅"),
        ("软件开发者", "代码修复、审查、重构与契约变更", "可执行、可测试、可回退", "本地仓库任务与差异审阅"),
        ("轻度数字用户", "查找文件、解释界面、重复操作", "低学习成本、明确确认", "桌面伙伴引导与受控自动化"),
    ], [1500, 2600, 2500, 2760], font_size=8.7)
    para("用户研究采用访谈、任务观察、可用性测试和连续使用记录相结合的方式。访谈用于理解动机与顾虑，任务观察用于识别真实操作路径，可用性测试用于发现界面问题，连续使用记录用于判断功能是否形成习惯。团队不以“用户表示愿意使用”替代行为证据，而以实际任务次数、结果采纳率和四周留存作为核心判断依据。")

    h2("2.4 竞争格局")
    para("竞争分析不按单一品牌输赢，而按用户完成任务所需的能力链条展开。通用AI助手具备模型与生态优势，系统级助手具备操作入口，AI输入工具具备表达低摩擦，RPA具备固定流程稳定性，开源Agent框架具备开发自由度。Cuttle的差异化不在于每项能力都强于成熟厂商，而在于用本地工作区、模型可替换、成果版本化、权限审批和情境连续把这些能力组织为面向个人生产力的完整产品。")
    para("竞争策略上，项目不与大型平台正面争夺通用问答和基础模型能力，而是聚焦可被验证的桌面任务闭环。模型能力可以接入和替换，Cuttle需要长期积累的是跨应用适配、权限治理、任务恢复、成果协作和用户工作习惯。上述能力越贴近真实环境，越难仅依靠模型升级复制，也更有机会形成稳定的产品壁垒。")
    table(["类别", "代表能力", "主要优势", "观察到的空白", "Cuttle定位"], [
        ("通用AI助手", "对话、搜索、文件与应用连接", "模型能力和服务生态强", "本地工作区、模型替换和审计深度因产品而异", "聚焦个人电脑真实任务闭环"),
        ("系统级助手", "截图、文件搜索、系统设置和入口", "与操作系统结合紧密", "多Agent成果协作和版本交付不是统一主线", "补足受控执行与成果管理"),
        ("AI输入工具", "补全、改写、翻译、语音", "使用频率高、学习成本低", "通常不承担复杂任务编排与执行", "作为即时入口连接长任务工作台"),
        ("RPA与自动化", "规则流程和企业系统自动化", "确定性流程稳定", "开放桌面环境配置成本高、自然语言协作弱", "自然语言规划+确定性工具+用户审批"),
        ("开源Agent框架", "工具调用、多Agent与扩展", "开发自由度高", "普通用户产品化、安全治理与安装门槛高", "提供桌面产品壳、评测和恢复体系"),
    ], [1350, 1900, 1700, 2500, 1910], font_size=8.2)

    h2("2.5 PEST分析")
    figure_box("图2-3  Cuttle项目PEST分析图", 5.5)
    number("政治与政策因素：人工智能+、智能终端、数据要素和数字化转型政策为智能助手和个人生产力工具提供发展方向，同时个人信息保护、数据安全和生成式人工智能服务规范提高了合规门槛。", restart=True)
    number("经济因素：模型推理成本持续下降，本地轻量模型可降低部分边际成本；但跨应用兼容、桌面发布、安全测试和用户支持需要持续投入，早期商业化必须聚焦高频高价值任务。")
    number("社会因素：用户对人工智能接受度快速提升，同时对误操作、隐私、幻觉和黑箱自动化保持警惕。产品必须用可见状态、明确确认和成果审阅建立信任。")
    number("技术因素：多模态模型、OCR、Windows UI Automation、本地推理、Agent工具调用和工作流技术日趋成熟，为情境感知与任务执行提供基础；系统兼容性和可靠性仍是工程难点。")
    para("PEST因素共同决定项目应采取稳健推进方式：政策机会和技术成熟度为产品提供窗口，用户规模扩大带来需求基础，但隐私合规、系统兼容和信任成本要求团队先建立安全边界与可复现案例。项目的市场节奏应与工程成熟度同步，避免营销承诺先于实际交付能力。")

    h2("2.6 SWOT分析")
    figure_box("图2-4  Cuttle项目SWOT分析图", 5.8)
    table(["优势 Strengths", "劣势 Weaknesses", "机会 Opportunities", "威胁 Threats"], [
        ("本地优先与权限审计进入主架构；已有Agent核心、评测和多智能体协同基础；输入入口高频；模型可替换。", "统一产品形态尚未完全落地；输入法和桌宠需重新立项；兼容矩阵与真实用户数据不足；团队商业资源有限。", "智能体和AI终端加速普及；高校学习办公场景集中；用户对隐私和可控自动化需求上升；技能生态可拓展。", "大型平台快速整合系统入口；基础模型能力同质化；桌面自动化安全事件影响信任；推理成本和接口政策波动。"),
    ], [2340, 2340, 2340, 2340], font_size=8.7)
    para("基于SWOT分析，项目近期采用SO与WO并行策略：利用已有Agent工程基础和高校场景，快速形成高质量案例；同时通过小范围试点补足统一产品形态、兼容矩阵和用户数据。面对大型平台竞争和安全事件风险，项目坚持模型可替换、本地优先和权限可审计，不以高风险全自动化作为传播重点。")

    h2("2.7 市场空间测算")
    para("市场测算采用自上而下与自下而上结合的情景法。外部用户规模只用于说明行业基础，不直接等同于Cuttle可服务市场。内部测算先假设6.02亿生成式人工智能用户中，8%至12%属于高频电脑学习办公人群，再以校园、开发者社区和内容渠道逐步触达。三年经营目标为注册用户30万人、付费用户约3.6万人，该目标属于规划情景，必须由真实访谈、留存和付费测试逐步校准。")
    table(["口径", "测算方法", "规模表达", "说明"], [
        ("潜在用户池", "CNNIC 2025年末生成式AI用户", "6.02亿人", "外部公开统计，不等同于可服务市场"),
        ("高频PC学习办公人群", "情景假设为8%至12%", "约4816万至7224万人", "需用调研验证"),
        ("三年注册用户", "校园与专业社区逐步扩散", "30万人", "经营目标"),
        ("三年付费用户", "第三年12%付费转化情景", "约3.6万人", "与稳定性、留存和定价相关"),
    ], [1700, 3000, 1800, 2860])
    para("市场空间最终需要由自下而上的验证结果决定。团队将按“深度访谈—种子用户—校园团队—专业社区—机构试点”逐级扩大，每一级都设置激活、留存、付费意愿和支持成本门槛。若某一级未达到门槛，则优先调整产品或缩小目标场景，而不是继续扩大宣传范围。")
    figure_box("图2-5  市场规模与三年用户增长数据图", 5.0)

    h1("第三章  产品体系与项目创新")
    h2("3.1 产品总体形态")
    para("Cuttle以输入法为品牌入口，以情境感知为理解中枢，以桌面伙伴为状态窗口，以任务工作台为行动界面。用户面对的是一套连续产品，而不是多个互不相干的聊天机器人。输入法回答“现在怎么说”，桌面伙伴说明“系统认为现在在做什么”，任务工作台负责“接下来怎样做完”。")
    para("四层产品并不要求用户一次性学习全部能力。首次使用从低风险的表达建议开始，用户接受并纠正系统判断后，逐步启用文件处理、任务工作台和技能复用。每一层都可以独立关闭，且不因为开启更高级能力而默认扩大数据读取范围。这样的渐进式设计有利于降低学习成本，也为团队逐层验证产品价值和风险控制提供明确路径。")
    figure_box("图3-1  Cuttle四层产品体系结构图", 6.2)
    table(["交互层级", "适用任务", "主要界面", "控制要求"], [
        ("即时辅助", "补全、改写、翻译、解释和快捷回复", "候选区或快捷面板", "默认只生成建议，不直接外发"),
        ("主动建议", "可总结、可解释、可复用的当前情境", "桌面伙伴提示卡", "可忽略、可关闭、可纠正"),
        ("长任务执行", "研究、报告、代码、文件和跨应用任务", "任务工作台", "计划可见、写操作审批、结果可回退"),
    ], [1700, 3200, 2300, 2160])

    h2("3.2 核心功能体系")
    h3("3.2.1 情境感知与模式切换")
    para("系统在用户授权范围内识别当前应用、输入焦点、对象类型、任务主题和最近确认的成果。当用户从论文写作切换到师生沟通，系统不再沿用学术段落语气；当用户进入密码、支付或验证码等敏感窗口，情境引擎自动停止读取与建议。所有判断均应通过桌面伙伴或状态栏可见，用户能够一键纠正“学习、办公、聊天、编程、勿扰”等模式。")
    para("情境信息按照确定性由高到低分层获取：优先使用应用提供的结构化接口、当前文件和用户显式选择，其次使用可访问性树和界面语义，最后才将视觉识别作为候选证据。系统对每项情境结论保留来源和有效期，应用切换、文件关闭或用户纠正后及时失效，避免将过期背景错误带入新任务。")
    h3("3.2.2 智能输入与表达辅助")
    para("输入层提供句子补全、快捷回复、风格调整、翻译、摘要、专业术语和个人表达偏好。它不只预测“哪个词更常见”，还判断当前关系、文体和任务目标，例如将口语内容改为正式邮件、将简短答复改为礼貌沟通、在代码编辑器中优先提供技术术语。个性化词汇和风格可以查看、修改和清除。")
    para("表达建议遵循“低打扰、可比较、保留原意”的原则。候选内容尽量短而明确，涉及语气或立场变化时同时展示原文与改写差异；系统不自动替用户发送消息，也不把一次选择永久固化为偏好。用户可分别管理专业词库、常用表达和禁用场景，确保个性化能力始终可见、可撤销。")
    h3("3.2.3 文件、知识与成果处理")
    para("用户可以将PDF、Word、表格、图片、代码或资料文件交给系统。简单任务即时生成摘要或提纲，复杂任务进入工作台，由Agent读取文件、提取结构、形成证据、生成文档或代码差异。成果以版本化Artifact交付，记录来源、生成过程、验证结果和修改差异，避免只留下不可追溯的一段回答。")
    para("成果处理采用“原始资料—中间证据—最终产物”三级结构。原始资料保持只读或建立副本，中间证据记录提取位置、判断依据和未解决问题，最终产物明确版本、生成时间和验证状态。用户可从最终结论追溯到证据，也可以只重新执行失败环节，而不必让整个任务从头开始。")
    h3("3.2.4 多智能体协同")
    para("Coordinator负责理解目标、拆解任务、分配资源和收口结果；Worker只获得完成本任务所需的上下文、工具和预算；Human节点在信息不足、判断冲突或高风险动作前介入。任务采用有向无环图表达依赖关系，已完成成果不会因后续调整或重试而丢失。用户只查看一项完整任务和统一进度，不需要分别指挥多个Agent。")
    para("协同机制强调职责隔离和结果契约。每个Worker在开始前明确输入、输出、可用工具和完成条件，结束时必须返回结构化结果、证据和异常说明；Coordinator负责检查依赖是否满足并决定合并、重试或请求人工判断。对于共享文件和关键接口，系统设置唯一责任人，避免多个智能体同时修改造成冲突。")

    h2("3.3 典型应用场景")
    table(["场景", "用户起点", "Cuttle连续动作", "最终交付"], [
        ("学习与课程报告", "阅读课程资料和教师要求", "识别主题—提取证据—在Word中提供情境写作—检查结构", "报告初稿、引用证据与修改建议"),
        ("科研与文献整理", "多篇论文、实验记录和研究问题", "分类—对照—提炼争议—形成研究矩阵", "综述提纲、证据表和待验证问题"),
        ("办公与项目材料", "会议记录、表格、旧文档和模板", "汇总—改写—格式统一—数据核对", "可审阅的正式文档与变更记录"),
        ("软件开发", "本地仓库、错误信息和验收要求", "定位—修改—测试—差异审阅—回退", "代码补丁、测试证据和说明"),
        ("沟通与表达", "聊天、邮件、社交平台或客户回复", "识别关系与语气—给出候选—保留用户选择", "得体表达与个人风格复用"),
    ], [1500, 2100, 3600, 2160], font_size=8.7)
    para("典型场景的共同特征是任务拥有明确起点、可检查过程和可交付结果。项目演示不以随机开放问答作为主体，而是选择资料、约束和验收标准均可固定的任务，完整展示情境建立、计划生成、权限审批、执行验证和成果审阅。这样既能说明产品创新，也便于评委和试点用户复现。")

    h2("3.4 用户旅程")
    figure_box("图3-2  从输入到任务交付的用户旅程图", 6.2)
    number("情境建立：用户打开课程资料或项目文件，Cuttle识别当前应用、对象与任务主题。", restart=True)
    number("即时表达：输入入口提供与当前任务匹配的术语、句子和风格建议，用户可直接采纳或纠正。")
    number("任务升级：当请求涉及多文件、多步骤或持续执行时，输入入口将任务升级到工作台。")
    number("安全执行：系统展示计划、工具范围、风险和审批节点，用户确认后再执行写入、命令或外部操作。")
    number("成果交付：系统返回可审阅的文档、代码差异、证据表或其他成果，并保留版本、审计和回退信息。")
    number("经验复用：经过验证的流程可沉淀为个人技能或团队模板，下次任务减少重复提示。")
    para("用户旅程中的每个节点都设置退出和纠正通道。用户可以拒绝某条建议、修改任务目标、缩小工具范围、暂停执行或回退成果；系统则记录纠正原因，用于改进后续建议。项目将重点测量各节点的等待时间、人工接管率和失败恢复率，持续减少不必要的确认，同时保留高风险动作的明确审批。")

    h2("3.5 技术路线")
    figure_box("图3-3  Cuttle技术架构与安全边界图", 6.5)
    table(["架构层", "关键组件", "主要职责", "设计原则"], [
        ("交互层", "输入面板、桌面伙伴、任务工作台", "任务入口、状态、审批、差异与结果审阅", "复杂能力保持可理解"),
        ("情境层", "应用信息、结构化API、文件状态、UI语义、OCR", "建立当前对象、主题、位置与用户意图", "确定性证据优先，视觉只作候选"),
        ("协同层", "会话、Goal、Plan、WorkSwarm、Workflow", "目标拆解、成员调度、Human介入与成果接力", "单一状态源、明确终态"),
        ("执行层", "文件、Git、浏览器、API、命令与桌面操作", "在明确范围内执行并验证动作", "工具最小授权、动作可验证"),
        ("资产层", "Project Space、Artifact、ChangeSet、Trace、Memory", "保存成果、证据、差异和可复用经验", "可追溯、可恢复"),
        ("安全层", "Policy、Approval、Audit、Credential、Sandbox", "阻断越权、保护凭据、记录决策", "默认拒绝、审批与Agent分离"),
        ("模型层", "本地模型、兼容云端模型、OCR与视觉模型", "理解、规划、生成和感知辅助", "模型可替换、密钥环境注入"),
    ], [1350, 2300, 3300, 2410], font_size=8.4)
    para("技术实施按照“稳定底座优先、统一接口先行、交互形态渐进接入”的顺序推进。首先保证会话、权限、审计、任务和成果管理稳定，再通过标准契约接入情境模块和输入入口，最后完善桌面伙伴的主动交互。各层之间只交换必要的结构化数据，避免界面组件直接访问模型凭据或底层工具。")
    para("兼容性方面建立操作系统版本、应用类型、输入环境和显示设置四维矩阵。能够通过结构化API完成的动作优先使用API；必须依赖桌面操作时，先识别可访问性节点，再以视觉定位补充，并在操作后验证界面状态。高频场景形成自动回归任务，发布前在固定环境和真实设备上同时复验。")

    h2("3.6 项目创新点")
    h3("3.6.1 从词频预测升级为全域情境表达")
    para("传统输入法主要依据输入历史与局部文本排序候选。Cuttle进一步结合当前应用、任务对象、关系和已确认资料，使表达建议能够随学习、办公、沟通和编程情境变化，同时允许用户查看和纠正系统判断。")
    para("该创新的重点不是读取更多数据，而是建立透明、有限且可纠正的情境模型。系统只使用当前任务需要的信息，并通过来源标识和模式提示告诉用户“为什么得到这条建议”。当证据不足时，产品应降低建议强度或主动询问，而不是把推测包装成确定结论。")
    h3("3.6.2 从单一入口升级为三级交互")
    para("即时辅助、主动建议和长任务执行使用不同界面和控制级别。低风险请求保持轻量，高风险操作进入工作台并等待审批，既保持输入法的流畅，也避免把复杂自动化藏在候选区。")
    para("三级交互还对应不同的信息密度和责任边界：即时辅助只展示少量候选，主动建议补充判断依据和可关闭入口，长任务工作台则完整呈现计划、工具、审批和成果。用户对能力的理解随风险同步增加，从而减少“界面看似简单、背后却执行大量动作”的失控感。")
    h3("3.6.3 从零散回答升级为成果中心协同")
    para("多智能体围绕版本化Artifact、任务图和Handoff协作，成果具有来源、证据、差异、审阅和恢复能力。用户得到的是一项完整任务的最终状态，而不是多个Agent互相冲突的文本。")
    para("成果中心使任务协作从对话记录转向可管理资产。报告、代码、数据表和研究矩阵都拥有明确版本，相关的中间证据和验证结果与成果绑定。后续成员可以在已有成果上继续工作，用户也能比较版本、选择接受部分修改或恢复到先前状态。")
    h3("3.6.4 从黑箱自动化升级为可信执行")
    para("权限、审批、审计、敏感熔断和回退不是外围功能，而是任务主流程。系统能够解释为什么需要某项权限、将执行哪些动作、失败后如何恢复，帮助用户逐步建立对智能体的信任。")
    para("可信执行不仅关注动作发生之前的授权，也关注动作完成之后的验证。系统需要检查文件是否按预期生成、命令是否成功、外部页面是否处于目标状态，并把验证结果写入审计记录。无法确认结果时应标记为待复核，而不是直接宣告任务完成。")
    h3("3.6.5 从一次性提示升级为技能复用")
    para("经过真实任务验证的过程可以沉淀为个人技能、团队模板或开发者插件。技能必须声明输入、输出、工具和权限范围，并通过健康检查和版本管理，形成可扩展但受治理的生态。")
    para("技能复用的门槛不是“曾经运行成功”，而是拥有明确适用条件、测试样例、失败处理和维护责任。技能升级后需要重新验证依赖与权限，出现异常可以回退到稳定版本。团队将优先开放低风险、结果易验证的技能，再逐步扩展到需要外部系统和写操作的场景。")

    h2("3.7 安全与隐私设计")
    bullet("权限默认拒绝：文件写入、命令执行、网络访问、消息发送和桌面操作必须经过策略。")
    bullet("审批与主智能体分离：主智能体不能为自身的越权工具调用自动授权。")
    bullet("凭据隔离：模型密钥只经环境变量或操作系统凭据注入，不进入代码、日志或普通Worker环境。")
    bullet("敏感场景熔断：密码、支付、验证码、隐私窗口和不可逆操作触发停止、预览或二次确认。")
    bullet("最小数据原则：只读取完成任务所需信息，优先本地处理，支持查看、清除和关闭情境记忆。")
    bullet("全过程审计：记录主体、目标、权限决定、关联任务、结果摘要和恢复动作。")
    para("隐私设置采用分级授权和最小保留策略。用户可以分别控制应用状态、文件内容、历史任务和个性化记忆，默认不把本地资料用于与当前任务无关的训练或分析。删除记录时同步清理索引和派生摘要；需要云端模型处理的内容在发送前展示范围，并尽量进行脱敏和片段化。")
    lead("安全承诺", "Cuttle的产品价值不是替用户做更多未经确认的事情，而是在提升效率的同时保持用户对数据、动作和结果的最终控制权。")

    h1("第四章  商业模式")
    h2("4.1 盈利模式")
    para("项目采用“个人订阅+团队许可+私有部署+技能生态”的分层商业模式。早期以个人专业版验证高频任务的持续付费，以校园与小型团队版验证协作和管理价值；当产品稳定性、安全治理和交付能力成熟后，再进入私有部署与技能生态。免费版承担获客与口碑，不以出售个人隐私数据作为收入来源。")
    para("商业化顺序遵循产品成熟度。个人订阅要求产品能够稳定完成用户每周重复任务，团队许可要求成员管理、统一策略和支持流程可复制，私有部署则进一步要求内网安装、模型适配、升级维护和服务责任清晰。任何新收入形态都必须在交付成本和安全风险可控后开启，避免短期项目收入拖累核心产品。")
    figure_box("图4-1  Cuttle多层商业模式结构图", 5.8)

    h2("4.2 产品版本与定价")
    table(["产品形态", "目标客户", "核心权益", "建议定价", "收入方式"], [
        ("个人免费版", "学生与轻度用户", "基础输入辅助、少量本地任务、核心安全功能", "免费", "获客与口碑"),
        ("个人专业版", "高频知识工作者与开发者", "更高任务额度、多智能体、专业技能包、优先支持", "19至39元/月", "订阅收入"),
        ("校园与团队版", "实验室、课程团队、小型组织", "团队空间、管理策略、统一部署和使用报告", "按席位或项目报价", "许可与服务"),
        ("私有部署版", "有数据合规要求的机构", "本地模型、内网部署、定制集成和审计", "项目制报价", "实施与维护"),
        ("技能与模板生态", "开发者、教师、服务商", "技能包、工作流和团队模板分发", "交易分成", "平台服务"),
    ], [1500, 1800, 3000, 1500, 1560], font_size=8.4)
    para("定价采用价值验证与成本约束相结合的方式。个人版通过小规模付费测试比较不同权益组合的转化和留存，团队版根据席位、管理功能和支持强度报价，私有部署则明确实施、模型、硬件和维护边界。正式定价前需测算单任务模型成本、用户支持成本和退款风险，避免仅参考同类产品标价。")

    h2("4.3 收入来源")
    number("订阅收入：来自个人专业版月度或年度订阅，是可持续现金流的核心。", restart=True)
    number("许可与服务：来自校园、实验室和小型团队的席位许可、部署、培训和支持。")
    number("私有部署与集成：面向有数据合规、内网和专有工作流需求的机构。")
    number("技能生态分成：经过平台审核的技能包、模板和行业工作流交易收入。")
    number("联合解决方案：与模型厂商、硬件厂商和软件服务商共同交付，按项目结算。")
    para("收入结构应保持订阅与标准化许可为核心，定制项目用于验证行业需求而非无限扩张。团队每季度检查各类收入的毛利、回款周期和研发占用，对重复出现的定制需求抽象为通用功能或技能模板；无法复用且持续占用核心研发资源的项目原则上不承接。")

    h2("4.4 成本结构")
    table(["成本类别", "主要构成", "控制方法"], [
        ("研发成本", "输入法、情境引擎、工作台、兼容性、安全、测试与发布工程", "聚焦核心场景，建立自动评测和兼容矩阵"),
        ("模型与基础设施", "云端推理、本地模型下载、更新、监控与错误诊断", "模型分级、预算上限、本地推理和缓存"),
        ("市场与服务", "校园试点、开发者社区、内容运营、用户支持和机构交付", "先验证留存和付费，再扩大获客"),
        ("合规与知识产权", "隐私评审、软著、商标、第三方许可与安全测试", "建立清单和阶段审查"),
    ], [1700, 4300, 3360])
    para("成本管理以单任务和单客户为基本单位。模型调用记录输入输出量与预算，桌面和兼容性问题记录支持时长，机构项目记录实施、培训和维护投入。通过任务分级、本地模型、缓存、自动诊断和标准交付包降低边际成本，并在新增功能立项时同步评估后续测试和支持负担。")

    h2("4.5 合作伙伴策略")
    para("项目优先建立四类合作关系：与高校实验室和课程团队共建真实任务试点；与国产模型和本地推理厂商优化成本与隐私；与开发者社区共建技能和兼容测试；与办公、教育和行业软件服务商探索接口集成。所有合作应以数据授权、责任边界、接口稳定性、知识产权和退出机制为前置条件。")
    para("合作采用由浅入深的三级机制。第一阶段以联合测试和案例共建验证需求，第二阶段形成标准接口、服务范围和数据处理协议，第三阶段在指标达标后再讨论预装、联合销售或规模交付。合作中产生的代码、数据、模型适配和品牌材料均明确归属，避免因口头承诺形成后续争议。")
    table(["合作伙伴", "合作内容", "双方价值", "准入条件"], [
        ("高校与实验室", "试点、用户研究、课程与科研任务", "获得真实反馈与创新工具", "知情同意、数据最小化、任务可复现"),
        ("模型与算力厂商", "模型接入、本地推理和成本优化", "获得高频终端应用场景", "接口稳定、价格透明、可替换"),
        ("开发者与服务商", "技能、模板、适配和行业方案", "获得分发和商业化渠道", "权限声明、健康检查、版本治理"),
        ("软件与硬件厂商", "接口集成、预装或联合方案", "完善智能终端体验", "安全评审、兼容性与用户控制"),
    ], [1700, 2600, 2700, 2360], font_size=8.7)

    h1("第五章  市场营销策略")
    h2("5.1 市场定位")
    para("Cuttle的市场定位不是“又一个大模型聊天软件”，而是“面向高频电脑任务的可信智能输入与执行入口”。传播内容应突出用户具体动作和结果：不再反复复制背景、表达会随场景变化、复杂任务过程可见、写操作需要确认、成果能够审阅和恢复。品牌语气保持克制、可信和可复现，避免以炫技动画替代真实任务闭环。")
    para("市场切入遵循“人群足够集中、任务足够高频、结果足够可验证”的原则。高校学生和科研团队具有资料整理、报告写作、代码实践和沟通表达等连续任务，且便于开展面对面访谈和长期观察；开发者与高频办公用户则更关注本地文件、版本差异和流程自动化，可用于检验产品在复杂任务中的稳定性。项目不在早期同时覆盖所有消费者场景，而是先在上述人群中建立可信案例。")
    para("品牌认知需要同时传达效率与边界。单纯强调“自动完成一切”会放大用户对误操作和隐私的担忧，因此宣传中应完整展示建议如何产生、权限如何确认、任务如何验证以及失败如何恢复。Cuttle希望形成的品牌印象不是激进替代用户，而是理解当前工作、降低重复劳动并始终尊重最终决定权。")
    lead("品牌主张", "让输入理解情境，让表达自然抵达，让任务高效落地。")

    h2("5.2 产品与渠道策略")
    para("产品策略以真实任务包而不是孤立功能作为推广单位。每个任务包包含适用人群、输入资料、操作步骤、权限范围、预期成果和评价指标，例如“课程资料生成结构化报告”“多篇文献形成证据矩阵”“本地仓库完成代码修复”。渠道传播、用户培训和产品验收围绕同一任务包展开，保证外部承诺与产品能力一致。")
    table(["阶段", "时间", "核心目标", "主要动作", "衡量指标"], [
        ("种子验证", "0至3个月", "找到3类高频任务", "校内访谈、内部强制使用、固定任务复现", "30名深度用户、100次真实任务、完成率"),
        ("校园试点", "4至8个月", "形成可展示案例", "课程资料、科研证据、代码任务工作坊", "3个试点团队、四周留存、可引用反馈"),
        ("专业社区增长", "9至15个月", "扩大高价值用户", "发布部分SDK、技能模板、技术内容和演示视频", "注册、激活、周任务数、推荐率"),
        ("机构合作", "16至24个月", "验证团队和私有部署", "与实验室、创新中心和软件团队共建", "付费试点、交付周期、续费意愿"),
    ], [1500, 1300, 2200, 2900, 1460], font_size=8.3)
    para("渠道选择按照反馈质量和扩散效率排序。校内实验室、课程团队和创新项目适合深度共创，能够提供连续任务和现场反馈；开发者社区适合验证技术可信度和技能扩展；短视频和图文平台用于传播具象案例，但不作为早期需求判断的唯一来源；机构合作则在产品稳定、交付流程和责任边界明确后逐步推进。")
    para("每个阶段设置退出条件。若种子用户无法连续完成目标任务，团队优先修正产品而非扩大安装；若校园试点四周留存不足，则重新聚焦任务价值和首次体验；若专业社区增长带来大量低质量用户，则调整内容与准入；机构试点若定制比例过高，则暂停扩张并提炼标准能力。")

    h2("5.3 宣传与用户增长")
    h3("5.3.1 赛事与校园传播")
    para("围绕“痛点—创新—证据—价值”制作路演材料，用同一条真实任务贯穿输入、情境、审批和成果交付。校园传播与课程项目、实验室和创新训练结合，招募愿意提供真实任务日志的种子用户，而不是只追求安装量。")
    para("赛事传播分为赛前准备、现场呈现和赛后沉淀三个阶段。赛前建立统一数据口径，固定演示设备、任务材料、版本哈希和异常预案；现场采用短链路说明核心价值，再通过完整案例展示计划、审批、执行和恢复；赛后将评委问题、失败环节和用户反馈转化为产品任务，避免参赛材料与实际研发脱节。")
    para("校园活动以小型工作坊和任务共创为主。团队可联合课程教师、实验室和学生组织，选择一类真实任务进行现场拆解，让参与者带着自己的资料完成试用。活动结束后收集任务完成时间、建议采纳率、操作困难和再次使用意愿，并邀请符合条件的用户进入四周观察组。相比一次性宣讲，这种方式更容易发现产品是否真正节省时间。")
    para("校园传播内容应尊重学术诚信和数据隐私。课程作业场景明确辅助边界，要求用户审阅并承担最终责任；科研资料和未公开项目默认在本地或授权环境处理；涉及他人信息时先完成脱敏。通过把安全说明纳入活动流程，项目可以把可信设计本身转化为差异化传播内容。")
    h3("5.3.2 开发者传播")
    para("公开接口契约、评测方法和示例技能，发布从目标到可审阅成果的过程型案例。技术内容重点说明本地工作区、模型可替换、权限策略、Artifact和恢复机制，形成可信、可复现的工程形象。")
    para("开发者内容采用“问题复现—设计选择—实现边界—验证结果”的结构。团队公开最小示例、接口定义、权限声明和测试方法，让开发者能够在自己的环境中复验，而不是只展示剪辑后的成功画面。对于仍处于规划阶段的输入法融合和桌面伙伴能力，明确标注原型、实验或待实现状态。")
    para("社区运营围绕贡献闭环展开：新开发者通过文档和示例完成首个技能，维护者对权限、依赖和测试进行审核，发布后持续观察健康状态和用户反馈。对兼容性问题、文档改进和安全缺陷设置不同贡献入口，并通过版本记录、贡献者名单和案例展示给予认可，逐步形成稳定的开发者关系。")
    para("技术传播的核心指标不只是浏览量，而是文档完成率、示例运行成功率、有效问题数量、外部贡献数和技能复用次数。若内容获得大量曝光却无法带来复现与贡献，团队应降低概念性表达，增加安装、调试、边界和失败处理细节。")
    h3("5.3.3 内容传播")
    para("制作“同一任务，传统方式与Cuttle方式”的对比内容，展示时间、步骤、错误和成果差异；避免纯概念宣传和无法复现的炫酷场景图。可围绕课程报告、科研综述、项目计划书和代码修复形成系列案例。")
    para("内容矩阵分为认知、理解、验证和转化四类。认知内容用简短场景说明用户痛点；理解内容解释情境感知、分级交互和可信执行；验证内容公开任务材料、过程记录和结果对比；转化内容提供明确的试用入口、适用条件和隐私提示。四类内容互相链接，使用户从看到概念逐步走向完成真实任务。")
    para("案例制作坚持数据可追溯。时间节省需要记录起止条件，质量提升需要说明评价方法，用户反馈需要获得授权并保留原始语境。失败案例同样具有价值，可以展示系统如何停止、请求人工判断和恢复成果。真实呈现限制能够降低过度承诺，也有助于建立长期信任。")
    para("传播节奏以稳定更新而非集中轰炸为主。每月围绕一个重点任务形成长文、短视频、操作清单和复盘，赛事节点再整合为专题内容。不同平台保持统一事实口径，但根据受众调整表达：校园用户突出学习和报告任务，开发者突出接口与验证，机构用户突出权限、部署和管理。")
    figure_box("图5-1  Cuttle用户增长与传播闭环图", 5.3)

    h2("5.4 留存与口碑机制")
    para("留存来自持续完成真实任务，而不是单次新鲜感。产品应围绕任务完成率、结果采纳率、四周留存、节省时间、错误与恢复率建立指标。用户在首次任务中快速看到可审阅成果，在后续任务中通过技能和记忆减少重复提示；当系统判断错误时，纠正成本必须低于重新描述成本。")
    para("首次体验采用渐进式引导：先让用户选择一个任务模板并解释最小权限，再在完成后展示成果、节省步骤和可复用技能。系统不要求用户一开始配置全部模型、工具和偏好，而是在出现真实需要时逐项说明。对于首次失败的用户，提供可理解的原因、人工替代路径和一键反馈，避免因黑箱错误直接流失。")
    para("留存运营按用户状态分层。新用户关注首次成功和安全理解，活跃用户关注任务效率与个性化，沉默用户需要识别是价值不足、兼容问题还是风险顾虑，团队用户则需要模板、权限和成员协作。触达消息仅在有明确价值时发送，不以频繁提醒制造虚假活跃。")
    table(["关键阶段", "用户问题", "产品机制", "核心指标"], [
        ("首次使用", "不知道能做什么", "场景化任务模板与权限说明", "首次任务成功率、完成时间"),
        ("形成习惯", "每次仍需重新提示", "情境连续、专业词汇和技能复用", "周任务数、四周留存"),
        ("建立信任", "担心误操作和泄露", "审批、审计、回退和敏感熔断", "越权拦截、误操作与恢复率"),
        ("主动推荐", "价值难以表达", "可分享的节省时间与成果案例", "推荐率、自然增长占比"),
    ], [1700, 2500, 3000, 2160])
    para("口碑案例只有在用户授权、结果可复验和边界描述完整时对外发布。团队将优先沉淀可公开的课程整理、开源代码和通用办公案例，对涉及个人、科研未公开数据或机构内部资料的任务只做匿名统计。推荐机制以分享任务模板、成果方法和节省时间为主，不设置诱导用户上传敏感内容的奖励。")

    h1("第六章  财务分析")
    h2("6.1 测算口径与关键假设")
    para("本章为参赛用途的中性情景测算，不代表项目已经形成收入、融资承诺或用户规模。测算以个人订阅、团队许可、私有部署和服务收入构成，成本包括研发人力、模型与基础设施、市场服务、合规和日常运营。正式提交前应由团队依据最新定价、用户转化、合同和成本凭证逐项复核。")
    para("测算采用三层口径。用户层关注注册、激活、留存和付费转化；收入层分别估算个人订阅、团队许可和项目服务；成本层记录固定研发投入与随任务增长的模型、支持和交付成本。各层之间通过明确公式关联，任何用户规模调整都同步反映到收入、模型消耗和服务成本，避免只修改收入而忽略相应支出。")
    para("财务数据每季度滚动更新。实际用户、付费、合同和成本凭证形成已发生口径，已签约未交付事项形成在手口径，其余仅保留为情景预测。对外材料优先展示区间和关键假设，不把远期预测包装为确定结果；当真实数据与预测偏差较大时，及时说明原因并调整经营计划。")
    table(["关键假设", "第一年", "第二年", "第三年", "说明"], [
        ("累计注册用户", "1万人", "8万人", "30万人", "校园与专业社区逐步扩散"),
        ("付费用户", "500人", "6400人", "3.6万人", "付费率5%、8%、12%情景"),
        ("团队与机构客户", "5个", "20个", "60个", "先试点后复制"),
        ("综合客单与服务收入", "低", "中", "中高", "与产品稳定性和交付深度相关"),
    ], [2200, 1400, 1400, 1400, 2960], font_size=8.7)

    h2("6.2 首轮资金需求情景")
    para("首轮资金需求按80万元情景设计，主要用于稳定发布、桌面端与输入入口开发、兼容性和安全测试、校园试点及知识产权。资金使用遵循研发优先、验证优先、现金流纪律和分阶段拨付原则，在留存与付费假设未被验证前不大规模投放。")
    para("资金按里程碑释放。第一阶段用于完成稳定安装包、核心任务闭环和安全基线；第二阶段在产品通过内部验收后投入校园试点和兼容测试；第三阶段仅在留存和用户价值达到门槛后增加市场与服务投入。每项支出对应负责人、预算、验收物和票据，重大调整经团队与指导教师复核。")
    table(["用途", "比例", "金额情景", "主要产出"], [
        ("产品与研发", "45%", "36万元", "稳定发布、输入入口、桌面端、技能与兼容测试"),
        ("市场与试点", "20%", "16万元", "用户研究、校园试点、内容与活动"),
        ("模型与基础设施", "15%", "12万元", "推理、更新、监控、测试资源"),
        ("合规与知识产权", "10%", "8万元", "软著、商标、隐私与安全评审"),
        ("运营与预备金", "10%", "8万元", "支持、设备、差旅和不确定性缓冲"),
    ], [2100, 1200, 1700, 4360])
    figure_box("图6-1  首轮资金使用结构图", 4.8)

    h2("6.3 三年财务预测")
    para("三年预测体现从研发验证到初步规模化的节奏。第一年以产品投入和试点为主，收入不足以覆盖研发成本；第二年依靠个人订阅和团队许可增长，目标在中后期接近盈亏平衡；第三年在留存、交付和支持流程稳定的前提下扩大用户与机构客户。若产品成熟度或付费验证未达到要求，团队规模和市场投入相应延后。")
    table(["指标", "第一年", "第二年", "第三年", "关键假设"], [
        ("营业收入", "24万元", "180万元", "650万元", "订阅、许可、私有部署与服务综合"),
        ("经营成本", "52万元", "150万元", "380万元", "研发前置投入，规模后支持成本增加"),
        ("经营结果", "-28万元", "30万元", "270万元", "第二年中后期达到盈亏平衡"),
        ("期末团队规模", "5至8人", "10至15人", "20至30人", "以产品和交付能力为扩张条件"),
    ], [1900, 1300, 1300, 1300, 3560])
    figure_box("图6-2  三年收入、成本与经营结果数据图", 5.5)
    h3("6.3.1 收入结构预测")
    para("个人订阅提供相对稳定的基础收入，重点取决于高频任务留存和权益差异；团队许可来自实验室、课程团队和小型组织，要求成员管理与统一策略可用；私有部署和服务收入金额较高但交付周期长，需要严格控制定制；技能和联合方案放在核心产品稳定之后，避免生态建设早于用户需求。")
    table(["收入来源", "第一年", "第二年", "第三年", "发展逻辑"], [
        ("个人订阅", "8万元", "70万元", "260万元", "由专业用户和开发者留存带动"),
        ("团队许可", "6万元", "45万元", "170万元", "校园、实验室和小团队复制"),
        ("私有部署与服务", "10万元", "55万元", "180万元", "按项目和交付深度计价"),
        ("技能与联合方案", "0万元", "10万元", "40万元", "核心产品稳定后逐步开放"),
    ], [2200, 1400, 1400, 1400, 2960])
    para("收入确认以实际交付为前提。订阅收入按服务周期确认，团队许可与部署项目按里程碑验收，联合方案明确分成口径和回款责任。免费用户、试用额度和意向合作不计入营业收入，赛事奖金和补贴单独列示，避免与可持续经营收入混淆。")
    h3("6.3.2 成本结构预测")
    para("成本控制不以压缩必要的安全和测试投入为代价。研发与测试覆盖核心功能、兼容性、发布和恢复；模型与基础设施按照任务类型设定预算；市场与服务随留存和付费验证增加；合规与运营保障知识产权、隐私审查和用户支持。各类成本均设置负责人和预警线。")
    table(["成本项目", "第一年", "第二年", "第三年", "控制重点"], [
        ("研发与测试", "28万元", "70万元", "150万元", "发布稳定性和兼容性优先"),
        ("模型与基础设施", "8万元", "25万元", "70万元", "本地模型、分级调用和预算上限"),
        ("市场与服务", "8万元", "30万元", "90万元", "以留存和付费验证决定投入"),
        ("合规与运营", "8万元", "25万元", "70万元", "安全、知识产权、支持和管理"),
    ], [2200, 1400, 1400, 1400, 2960])
    para("随着用户增长，团队重点关注边际成本变化。模型成本可通过本地推理、缓存和任务分级下降，用户支持成本则依赖产品稳定性、诊断工具和文档完善。若单个机构项目长期占用核心研发，需将其成本完整计入项目毛利，不能由通用产品预算隐性补贴。")

    h2("6.4 敏感性与盈亏平衡")
    para("财务结果对付费转化率、模型调用成本、团队交付效率和机构项目周期最敏感。若第三年付费率从12%下降至8%，付费用户将由约3.6万人降至约2.4万人，订阅收入和现金流显著下降；若模型调用成本上升，应通过本地模型、任务分级、缓存和预算上限控制单任务成本；若机构项目回款周期延长，应限制定制范围并设置里程碑付款。")
    para("盈亏平衡分析采用月度滚动方式，将固定研发与运营成本除以订阅和许可的平均贡献毛利，得到需要维持的付费用户与团队客户组合。由于产品早期数据有限，计划书不把单一平衡点作为承诺，而是设置基准、保守和压力三种情景，并为每种情景准备人员、模型和市场投入调整方案。")
    table(["变量", "基准情景", "不利情景", "应对措施"], [
        ("付费转化率", "5%/8%/12%", "第三年8%", "缩小获客范围，聚焦高频专业任务和留存"),
        ("模型单位成本", "逐年优化", "上涨30%", "本地模型、分级调用、缓存和任务预算"),
        ("机构回款周期", "按阶段回款", "延长至6个月", "里程碑付款、控制定制和现金流预警"),
        ("任务支持成本", "随规模缓慢增长", "高于订阅毛利", "增强自助诊断、技能健康检查和产品稳定性"),
    ], [1800, 2200, 2200, 3160])
    lead("财务纪律", "在真实留存和付费形成前，不以市场规模代替现金流；模型调用、用户支持和机构交付必须建立单任务成本与毛利监控。")

    h1("第七章  组织与管理")
    h2("7.1 核心团队成员")
    para("本章项目成员信息依据《西南大学大学生创新训练计划项目申报书——基于全域情境感知的智能输入系统》填写，并根据用户提供的最新指导教师资料将指导教师更新为赵恒军。能够确认的内容包括项目负责人、项目组成员、学院以及指导教师专业信息；未明确的成员专业分工和个人成果不作虚构，正式提交前由成员本人确认补充。")
    table(["姓名", "团队身份", "所在单位", "已核实信息", "本计划书职责"], [
        ("杨俊熙", "项目负责人", "西南大学计算机与信息科学学院 软件学院", "原申报书项目负责人；曾获2025年第18届中国大学生计算机设计大赛重庆市级赛三等奖", "统筹产品定位、研发进度、赛事材料与跨组协调"),
        ("张子豪", "项目组成员", "西南大学计算机与信息科学学院 软件学院", "原申报书项目组成员", "参与技术研发、测试与产品验证；具体模块由团队确认"),
        ("吴栩彪", "项目组成员", "西南大学计算机与信息科学学院 软件学院", "原申报书项目组成员", "参与技术研发、兼容性与演示验证；具体模块由团队确认"),
        ("赵恒军", "指导教师", "西南大学计算机与信息科学学院 软件学院", "博士、副教授；研究方向为信息物理系统、软件形式化方法、AI for Science", "技术路线、形式化验证、安全边界、科研规范与赛事指导"),
    ], [1200, 1400, 2500, 2400, 1860], font_size=8.2)
    para("团队分工采用“角色明确、模块可调整、贡献有证据”的原则。计划书中的职责用于说明当前组织方式，不代表永久固定岗位；随着产品阶段变化，成员可在研发、测试、用户研究和赛事材料之间调整。每项任务均保留负责人、协作者、交付物和验收记录，最终以实际版本、文档、测试和运营成果确认贡献。")

    h2("7.2 指导教师")
    table(["项目", "信息"], [
        ("姓名", "赵恒军"),
        ("性别", "男"),
        ("学历与职称", "博士、副教授"),
        ("研究方向", "信息物理系统、软件形式化方法、AI for Science"),
        ("邮箱", "zhaohj2016@swu.edu.cn"),
        ("所在单位", "西南大学计算机与信息科学学院 软件学院"),
    ], [1900, 7460])
    para("赵恒军老师自2016年6月至今任职于西南大学计算机与信息科学学院软件学院；2014年7月至2016年6月在中国科学院重庆绿色智能技术研究院信息所自动推理与认知中心担任助理研究员；2008年9月至2014年7月在中国科学院大学计算机软件与理论专业攻读博士学位；2010年10月至2011年5月赴美国新墨西哥大学计算机科学系访问学习；2004年9月至2008年7月在复旦大学数学与应用数学专业完成本科学习。其学习与工作经历横跨数学基础、逻辑推理、软件理论和智能系统，为项目开展跨学科技术研究提供了扎实支撑。")
    para("在教学方面，赵恒军老师主要讲授软件工程专业《操作系统原理》《类库与数据结构实验》等课程，并为研究生讲授《计算机科学中的逻辑学》。这些课程覆盖系统资源管理、软件基础结构、抽象建模和逻辑推理，与Cuttle涉及的桌面系统集成、任务状态管理、工具安全调用和可验证执行具有直接关联。指导过程中可将课程中的系统性思维转化为项目的工程规范、测试方法和风险审查机制。")
    para("在科研方面，赵恒军老师主要关注含智能组件的信息物理系统的验证与设计，研究涉及形式化方法、人工智能和控制等领域，相关方法亦可用于程序代码正确性分析与验证；其AI+Logic方向致力于推动逻辑推理自动化，并指导学生开展神经网络可解释性、AI医学图像、恶意软件识别、安全强化学习和自动驾驶等研究。上述积累能够帮助Cuttle把“模型能做什么”进一步落实为“系统在什么条件下可以安全地做、如何证明执行过程符合约束”。")
    para("赵恒军老师主持或负责国家自然科学基金项目，包括在研的“强化学习安全性的形式化方法研究及其在智能控制中的应用”和已结题的“采样控制系统的安全性验证方法及其应用”。代表性成果涉及强化学习智能控制在安全攸关信息物理系统中的应用。项目将重点吸收其中的安全约束、状态验证和风险前置思想，用于任务计划审查、权限策略、敏感动作熔断、执行后验证和异常恢复设计。")
    para("指导工作按照“需求与边界评审—技术方案评审—安全与形式化检查—阶段验收—赛事复盘”开展。赵恒军老师负责对系统边界、关键模型、风险假设和验证方法提出专业意见；项目负责人负责组织研发、用户验证和证据归档。涉及权限升级、不可逆操作、敏感数据、外部发布和重大架构变更的事项，均形成书面评审记录，确保指导意见能够落实到产品设计和工程验收中。")

    h2("7.3 组织架构设计")
    figure_box("图7-1  Cuttle项目组织架构图", 6.2)
    para("项目初期采用扁平化组织，避免在人员规模较小时虚设多个部门。设项目负责人、技术研发、产品与用户验证、赛事与运营四类职责，并由指导教师提供专业指导。随着试点和商业化推进，再将兼容性、安全、市场、客户成功和财务职能逐步专业化。")
    para("重大事项采用分级决策。日常研发和内容更新由对应负责人决定，跨模块接口和版本发布由团队评审，涉及安全边界、个人数据、外部合作、财务承诺和知识产权的事项由项目负责人组织并邀请指导教师审查。若成员意见无法统一，以用户安全、材料真实性和核心里程碑为优先原则。")
    table(["职责单元", "主要工作", "决策边界", "阶段产出"], [
        ("项目负责人", "定位、范围、进度、资源、对外沟通和材料真实性", "范围与优先级；重大事项报指导教师和团队评审", "计划书、路线图、决策记录和路演"),
        ("技术研发", "输入入口、Agent、工具、权限、桌面端、测试与发布", "技术方案与实现；核心契约变更需评审", "源码、测试、安装包和技术文档"),
        ("产品与用户验证", "需求访谈、原型、可用性、指标和试点", "用户需求优先级；不得以假设代替证据", "访谈、原型、试点数据和复盘"),
        ("赛事与运营", "申报、路演、财务假设、品牌、合作与资料归档", "对外数据需多方核验", "商业材料、演示和合作记录"),
        ("指导教师", "技术、科研、合规、资源和赛事指导", "提出专业意见，不替代团队实际贡献", "指导记录、评审意见和资源协调"),
    ], [1700, 3300, 2500, 1860], font_size=8.4)

    h2("7.4 协作与质量管理")
    para("项目管理将产品、技术和赛事材料纳入同一里程碑体系。每个阶段同时定义用户任务、工程产物、测试证据和对外表述，避免产品尚未完成而宣传材料先行。任务在开始前明确完成条件，结束后由非直接实施成员复核，发现问题进入缺陷与风险清单。")
    bullet("源码协作按文件认领，同一文件同一时间只允许一名成员修改；核心接口变更同步契约测试。")
    bullet("功能完成以源码、测试、构建产物和真实交互四类证据综合判断，不以文档声明替代验收。")
    bullet("每两周开展产品评审，检查用户价值、进度、风险、数据真实性和演示可复现性。")
    bullet("公开材料由项目负责人、技术负责人和指导教师三方复核，确保不夸大能力、不泄露凭据。")
    bullet("团队成员贡献以版本记录、任务单、设计文档、测试和运营证据归档，避免仅按名次列名。")
    table(["管理活动", "频率", "参与者", "主要输出"], [
        ("研发站会", "每周", "项目成员", "进度、阻塞和下周任务"),
        ("产品评审", "每两周", "团队与种子用户", "功能价值、问题和优先级"),
        ("安全与发布评审", "每个发布版本", "技术成员与指导教师", "权限、凭据、兼容性和回退清单"),
        ("经营与赛事复盘", "每月/赛前", "全体成员", "用户、财务、材料和演示真实性检查"),
    ], [1800, 1300, 2700, 3560])
    para("质量管理强调真实环境验证。源码测试通过后，还需检查安装包、输入入口、桌面端和相关依赖是否为最新构建，并在目标电脑完成关键任务。演示前固定版本和数据，记录设备环境与恢复方案；对外出现问题时保留事实边界，不以临时手工操作掩盖产品缺陷。")

    h1("第八章  风险防控")
    h2("8.1 风险总述")
    para("Cuttle同时涉及输入法、桌面交互、智能体执行和个人数据处理，风险不仅来自模型效果，还来自系统兼容、权限边界、数据合规、市场信任、团队持续性和现金流。项目建立“识别—评估—预警—应对—复盘”的闭环，风险责任落实到具体模块和里程碑。任何可能导致不可逆写入、外部发送、敏感数据泄露或大范围安装失败的风险，优先级高于功能扩张。")
    para("风险按发生概率、影响程度、可检测性和恢复难度综合分级。高影响风险即使概率较低也需设置预防控制，例如凭据泄露、误发消息和不可逆文件操作；高概率兼容问题则通过场景收敛、灰度发布和回退降低影响。风险清单在每次版本评审和重大演示前更新，负责人必须说明当前状态和剩余风险。")
    h2("8.2 具体风险分析与应对措施")
    table(["风险", "概率", "影响", "预警信号", "应对措施"], [
        ("技术可靠性", "中", "高", "任务失败、误操作、恢复失败", "固定评测、分层执行、断言、灰度和回退"),
        ("输入与桌面兼容", "高", "高", "特定应用、分辨率或升级后不可用", "建立应用和系统版本兼容矩阵，优先结构化接口"),
        ("隐私与安全", "中", "极高", "凭据泄露、越权、敏感数据外发", "默认拒绝、独立审批、凭据隔离、审计和安全测试"),
        ("模型幻觉", "中", "高", "证据不一致、结果不可验证", "证据绑定、工具校验、Human节点和结果审阅"),
        ("产品范围失控", "高", "高", "同时推进输入法、桌宠和Agent导致延期", "以Agent工作台为工程主线，其他形态按决策门立项"),
        ("用户留存不足", "中", "高", "体验新鲜但四周留存低", "聚焦高频任务，用时间和质量验证价值"),
        ("市场竞争", "高", "中至高", "系统厂商快速整合类似能力", "强调本地、可替换模型、开源扩展与成果协同"),
        ("财务与成本", "中", "中", "模型价格或支持成本上升", "本地模型、预算上限、单任务成本和现金流预警"),
        ("团队持续性", "中", "高", "核心成员投入下降", "职责备份、文档化、模块所有权和阶段激励"),
        ("知识产权与许可", "中", "高", "第三方组件或数据授权不清", "依赖清单、许可证审查、原创证据和申请计划"),
    ], [1500, 800, 900, 2600, 3560], font_size=8.1)
    figure_box("图8-1  项目风险矩阵与预警机制图", 5.5)
    h3("8.2.1 安全事件处置")
    number("立即停止相关工具和外部连接，冻结受影响版本与凭据。", restart=True)
    number("保留审计和复现证据，区分模型判断、工具实现、权限策略与用户操作。")
    number("评估影响范围并通知相关人员，必要时轮换凭据、撤回发布或关闭功能。")
    number("完成根因修复、回归测试和真实场景复验后再恢复，形成公开或内部复盘。")
    para("安全事件复盘区分直接原因、系统性原因和管理原因，避免只修复表面错误。直接原因可能是工具实现或模型判断，系统性原因包括权限策略、验证缺失和测试覆盖不足，管理原因则涉及发布流程和责任不清。整改措施需指定负责人、截止时间和验证证据，未完成前不得恢复高风险能力。")

    h1("第九章  战略规划")
    h2("9.1 短期与长期战略目标")
    para("战略目标以产品证据而不是时间自然到达为判断依据。每个阶段只有在上一阶段的安装、任务成功、安全、留存或付费门槛满足后才能进入；若外部环境或团队资源变化，则允许延长、缩小或终止阶段。该机制使愿景保持长期一致，同时避免为了赶进度牺牲稳定性。")
    table(["阶段", "时间", "产品目标", "市场目标", "验收门槛"], [
        ("参赛稳定版", "0至3个月", "收敛本地多Agent工作台和固定演示", "完成材料与种子访谈", "安装、核心任务、恢复和卸载可重复"),
        ("校园试用版", "4至8个月", "服务学习、研究、文档和代码任务", "3个试点团队", "30名深度用户，关键任务成功率达标"),
        ("产品融合版", "9至15个月", "加入桌面伙伴状态和轻量输入建议", "专业社区扩散", "不打断、可关闭、权限和隐私通过评审"),
        ("规模化版", "16至24个月", "团队许可、私有部署和技能分发", "形成付费试点", "交付流程可复制，毛利和支持成本可控"),
        ("后续研究", "24个月后", "评估系统级输入法入口和高级桌面自治", "拓展合作生态", "独立立项，不降低安全与稳定基线"),
    ], [1500, 1300, 2600, 2000, 1960], font_size=8.3)

    h2("9.2 战略实施路径与规划")
    figure_box("图9-1  Cuttle两年战略实施路线图", 6.2)
    para("战略实施遵循“先可重复、再可留存、后可付费、最后可扩展”的顺序。第一阶段收敛技术和演示证据；第二阶段验证真实任务和留存；第三阶段才把桌面伙伴与输入入口逐步接入统一情境；第四阶段验证团队许可、私有部署和技能生态。每个阶段设置继续、调整或停止的决策门，避免因为愿景完整而同时扩张所有子产品。")
    para("实施过程中保持一条工程主线和两条探索支线。工程主线持续完善Agent工作台、权限、审计和成果管理，确保已有能力稳定；输入入口和桌面伙伴作为探索支线，通过原型与小规模试用验证交互价值，达到决策门后再进入主线。这样的资源配置可以减少多端并行开发造成的质量分散。")
    h3("9.2.1 关键决策门")
    bullet("M1可参赛：计划书、路演、演示包、测试报告和真实性材料闭环。")
    bullet("M2可试用：安装、升级、异常恢复、数据迁移和隐私说明达到内测标准。")
    bullet("M3可留存：目标用户连续四周使用，任务成功、节省时间和满意度有数据。")
    bullet("M4可付费：至少一种权益被真实购买，支持成本和模型成本可控。")
    bullet("M5可扩展：团队版或技能生态不破坏权限、审计、兼容和恢复指标。")
    para("决策门由产品、技术、用户和经营四类证据共同组成。任何单项表现突出都不能替代整体判断，例如演示效果好但安装不稳定、用户增长快但留存低、机构意向多但交付成本过高，都不足以进入下一阶段。评审结论和未满足条件形成记录，供后续复盘。")

    h2("9.3 知识产权与生态规划")
    para("项目知识产权布局围绕情境感知、分级交互、受控Agent执行、成果协作和技能治理展开。近期优先完成软件著作权、商标检索与申请、第三方依赖许可证清单和原创证据归档；在形成稳定技术方案和可验证创新后，再评估发明专利申请。对开源部分明确代码许可证，对技能市场建立作者、依赖、数据和模型权利声明。")
    para("生态建设遵循开放与治理并重。公共接口尽量保持稳定，示例和测试帮助开发者降低接入成本；同时，技能必须声明数据来源、模型依赖、权限范围、外部连接和维护状态。涉及写操作、消息发送和敏感数据的技能采用更严格审核，并为用户提供清晰的风险提示和禁用选项。")
    table(["对象", "近期动作", "中期动作", "证据要求"], [
        ("软件著作权", "整理稳定版本、代码和设计说明", "按核心模块持续登记", "版本哈希、源码和开发记录"),
        ("商标与品牌", "检索Cuttle同类冲突与可注册性", "完成主要类别注册", "检索报告与申请材料"),
        ("专利", "识别可保护的技术方案", "验证后选择性申请", "现有技术检索、实验和发明人记录"),
        ("开源与第三方", "建立依赖与许可证清单", "自动化合规检查", "版本、许可证和修改说明"),
        ("技能生态", "定义技能契约与权限声明", "审核、签名、分发和收益规则", "作者、数据、模型和测试证明"),
    ], [1700, 2800, 2600, 2260], font_size=8.6)

    h1("第十章  社会效益")
    figure_box("图10-1  Cuttle项目社会效益结构图", 6.2)
    h2("10.1 生产力与教育价值")
    para("Cuttle把资料整理、表达修改和任务执行中的机械步骤交给可控工具，让学生、教师和知识工作者把更多时间用于理解、判断和创造。对于高校项目，系统能够支持资料阅读、证据对照、报告写作和代码实践，过程中的计划、差异、审批和成果记录也可成为人工智能素养与软件工程训练材料。")
    para("产品不把“更快生成”作为唯一目标，而把来源、结构、可验证和可修改作为成果质量的重要组成。用户能够看见系统如何从输入得到结果、哪些内容需要确认、失败后怎样恢复，这有助于培养对人工智能的批判性使用能力。")
    para("在教学应用中，教师可以把任务要求、评价标准和允许使用的工具写入模板，学生则提交过程记录和最终成果。系统保留资料来源、修改差异和人工确认节点，为讨论人工智能辅助边界、学术诚信和软件工程方法提供案例。项目不会以自动生成替代学习过程，而是帮助学生把时间投入到理解、论证和验证。")

    h2("10.2 可信人工智能与数字包容")
    para("本地优先、模型可替换和最小权限有助于保护学习、研究和工作资料；可视化状态和分级交互降低普通用户使用复杂Agent的门槛；桌面伙伴与轻量输入建议能够为不熟悉专业提示词的用户提供清晰入口。产品还可通过语音、翻译、界面解释和重复操作复用，帮助有不同数字能力的用户更平等地获得生产力工具。")
    para("数字包容要求产品不仅能被技术熟练用户使用，还应让普通用户理解系统正在做什么。界面采用清晰语言解释权限、风险和结果，复杂设置提供推荐值与撤销路径，关键动作避免只使用技术术语。后续用户研究将覆盖不同专业、设备条件和数字经验，观察首次成功率、求助次数和恢复能力。")
    table(["价值方向", "项目贡献", "建议评估指标"], [
        ("学习与工作效率", "减少资料搬运、重复改写和手工串联步骤", "任务节省时间、成果质量、用户满意度"),
        ("数字包容", "用可视化伙伴和分级交互降低复杂智能体门槛", "首次成功率、求助次数、不同用户群可用性"),
        ("可信人工智能", "让权限、审批、证据、差异和结果可见", "越权拦截、审计覆盖、误操作和恢复率"),
        ("创新就业", "培养AI产品、Agent工程、交互、安全和运营复合能力", "参与学生、实践岗位、技能与案例产出"),
    ], [1900, 4500, 2960])

    h2("10.3 就业与产业带动")
    para("项目发展将直接需要输入法与桌面开发、Agent工程、模型与数据、安全测试、产品设计、用户研究、客户成功和市场运营等岗位；技能生态还可为开发者、教师和行业服务商提供模板与插件分发机会。团队规模应随真实用户、收入和交付需求增长，不以宏大就业数字代替可执行计划。")
    para("从产业角度看，Cuttle探索的是模型能力进入个人电脑真实工作流的产品化路径。其本地工作区、模型可替换、权限审计和成果协作机制可以为国产模型、智能终端、教育数字化和中小软件团队提供可复用的集成范式，推动人工智能应用从演示走向稳定、负责和可持续的日常使用。")
    para("项目还可形成面向学生的实践训练链路：参与者从需求调研、交互设计和软件开发进入真实产品，再学习测试、安全、数据治理、内容传播和客户支持。能力评价以可验证的任务和成果为依据，既培养人工智能技术能力，也培养工程责任、协作和商业意识。")
    lead("项目愿景", "让每一次输入都更懂情境，让每一项自动化都尊重边界，让人工智能真正成为可理解、可协作、可控制的个人数字伙伴。")

    doc.core_properties.title = "Cuttle——新一代智能输入法 参赛商业计划书"
    doc.core_properties.subject = "内容与排版阶段稿"
    doc.core_properties.author = "Cuttle项目团队"
    doc.core_properties.keywords = "Cuttle, 智能输入法, 情境感知, Agent, 商业计划书"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
