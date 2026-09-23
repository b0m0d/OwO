from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"T:\创新创业\OwO-master")
ASSET_DIR = ROOT / "builGoal" / "项目计划书素材"
OUTPUT = ROOT / "builGoal" / "OwO情境感知智能交互系统-参赛项目计划书.docx"

NAVY = "12233F"
BLUE = "2E74B5"
CYAN = "2AA7C9"
VIOLET = "7157C8"
INK = "243447"
MUTED = "66758A"
LIGHT = "F4F6F9"
LIGHT_BLUE = "EAF3FA"
LIGHT_CYAN = "E9F8FA"
GOLD = "C89B3C"
GREEN = "2E7D5B"
RED = "A94442"
WHITE = "FFFFFF"
GRID = "CFD8E3"


def set_run_font(run, latin="Calibri", east_asia="Microsoft YaHei", size=None,
                 bold=None, color=None, italic=None):
    run.font.name = latin
    run._element.get_or_add_rPr()
    fonts = run._element.rPr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph_border(paragraph, color=BLUE, size=12, space=6):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    return run


def create_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abstract = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(existing_abstract, default=0) + 1
    num_ids = {}
    for kind, fmt, text, font in (
        ("bullet", "bullet", "•", "Arial"),
        ("decimal", "decimal", "%1.", "Calibri"),
    ):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "280")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "290")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), font)
        r_fonts.set(qn("w:hAnsi"), font)
        r_pr.append(r_fonts)
        lvl.extend([start, num_fmt, lvl_text, suff, p_pr, r_pr])
        abstract.append(lvl)
        numbering.append(abstract)

        num_id = max(existing_num + list(num_ids.values()), default=0) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)
        num_ids[kind] = num_id
        abstract_id += 1
    return num_ids


def clone_num_id(doc, base_num_id):
    numbering = doc.part.numbering_part.element
    base = None
    for num in numbering.findall(qn("w:num")):
        if int(num.get(qn("w:numId"))) == int(base_num_id):
            base = num
            break
    if base is None:
        raise ValueError(f"numbering id not found: {base_num_id}")
    abstract_id = base.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    new_id = max(existing, default=0) + 1
    new_num = OxmlElement("w:num")
    new_num.set(qn("w:numId"), str(new_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), abstract_id)
    new_num.append(abstract_ref)
    lvl_override = OxmlElement("w:lvlOverride")
    lvl_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    lvl_override.append(start_override)
    new_num.append(lvl_override)
    numbering.append(new_num)
    return new_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([r_fonts, color_el, underline])
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.extend([r_pr, text_el])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_alt_text(inline_shape, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description[:120])


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.36)
    section.footer_distance = Inches(0.36)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.28

    for name, size, color, before, after in (
        ("Title", 28, NAVY, 0, 8),
        ("Subtitle", 13.5, MUTED, 0, 8),
        ("Heading 1", 17, BLUE, 16, 8),
        ("Heading 2", 13.5, BLUE, 11, 6),
        ("Heading 3", 11.5, NAVY, 7, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False

    num_ids = create_numbering(doc)
    current_decimal_num_id = num_ids["decimal"]

    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("OwO 情境感知智能交互系统  |  参赛项目计划书")
    set_run_font(hr, size=8.5, color=MUTED, bold=True)
    set_paragraph_border(hp, color="DDE4EC", size=5, space=3)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("第 ")
    set_run_font(fr, size=8.5, color=MUTED)
    page_run = add_field(fp, " PAGE ")
    set_run_font(page_run, size=8.5, color=MUTED)
    fr2 = fp.add_run(" 页  共 ")
    set_run_font(fr2, size=8.5, color=MUTED)
    total_run = add_field(fp, " NUMPAGES ")
    set_run_font(total_run, size=8.5, color=MUTED)
    fr3 = fp.add_run(" 页")
    set_run_font(fr3, size=8.5, color=MUTED)

    def para(text="", bold_lead=None, align=None, size=None, color=None, after=None,
             italic=False, keep=False):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.keep_together = keep
        if bold_lead and text.startswith(bold_lead):
            r1 = p.add_run(bold_lead)
            set_run_font(r1, size=size or 10.5, color=color or INK, bold=True)
            r2 = p.add_run(text[len(bold_lead):])
            set_run_font(r2, size=size or 10.5, color=color or INK, italic=italic)
        else:
            r = p.add_run(text)
            set_run_font(r, size=size or 10.5, color=color or INK, italic=italic)
        return p

    def bullet(text, level=0):
        p = doc.add_paragraph()
        apply_numbering(p, num_ids["bullet"])
        if level:
            p.paragraph_format.left_indent = Inches(0.375 + 0.25 * level)
        r = p.add_run(text)
        set_run_font(r, size=10.3, color=INK)
        return p

    def number(text, restart=False):
        nonlocal current_decimal_num_id
        if restart:
            current_decimal_num_id = clone_num_id(doc, num_ids["decimal"])
        p = doc.add_paragraph()
        apply_numbering(p, current_decimal_num_id)
        r = p.add_run(text)
        set_run_font(r, size=10.3, color=INK)
        return p

    def heading(text, level=1):
        p = doc.add_heading(text, level=level)
        if level == 1:
            set_paragraph_border(p, color="D9E7F3", size=7, space=4)
        return p

    def callout(label, text, fill=LIGHT_BLUE, accent=BLUE):
        table = doc.add_table(rows=1, cols=1)
        set_row_cant_split(table.rows[0])
        set_table_geometry(table, [9360], indent_dxa=120)
        set_table_borders(table, color=fill, size=2)
        cell = table.cell(0, 0)
        set_cell_shading(cell, fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r1 = p.add_run(label + "  ")
        set_run_font(r1, size=10.5, bold=True, color=accent)
        r2 = p.add_run(text)
        set_run_font(r2, size=10.3, color=INK)
        doc.add_paragraph().paragraph_format.space_after = Pt(1)

    def table(headers, rows, widths, aligns=None, font_size=9.3, header_fill=LIGHT):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        set_table_geometry(t, widths, indent_dxa=120)
        set_table_borders(t)
        hdr = t.rows[0]
        set_repeat_table_header(hdr)
        set_row_cant_split(hdr)
        for idx, text in enumerate(headers):
            cell = hdr.cells[idx]
            set_cell_shading(cell, header_fill)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(text))
            set_run_font(r, size=9.2, bold=True, color=NAVY)
        for row_data in rows:
            row = t.add_row()
            set_row_cant_split(row)
            for idx, value in enumerate(row_data):
                cell = row.cells[idx]
                p = cell.paragraphs[0]
                p.alignment = (aligns[idx] if aligns else WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(str(value))
                set_run_font(r, size=font_size, color=INK)
        set_table_geometry(t, widths, indent_dxa=120)
        return t

    def figure(path, caption_text, width=6.5):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        shape = p.add_run().add_picture(str(path), width=Inches(width))
        add_alt_text(shape, caption_text)
        cap = doc.add_paragraph(caption_text, style="Caption")
        cap.paragraph_format.keep_with_next = False

    # Cover page: proposal_centerpiece pattern with a branded visual.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("大学生创新创业竞赛参赛材料")
    set_run_font(r, size=10.5, bold=True, color=CYAN)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("OwO")
    set_run_font(r, size=32, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("情境感知智能交互系统")
    set_run_font(r, size=25, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("让人工智能从被动问答走向可理解、可协作、可控制的桌面行动伙伴")
    set_run_font(r, size=12.5, color=MUTED)

    figure(ASSET_DIR / "owo-cover-visual.png", "封面视觉  情境感知、桌面伙伴与多智能体任务协作", width=6.25)

    meta = table(
        ["参赛信息", "待填写内容"],
        [
            ("参赛赛道", "高教主赛道／创意组／其他：待确认"),
            ("项目负责人", "待填写"),
            ("团队成员", "待填写"),
            ("学校与学院", "待填写"),
            ("指导教师", "待填写"),
            ("版本日期", "参赛稿 V1.0  2026年9月"),
        ],
        [2100, 7260],
        aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=9.3,
        header_fill="E9F1F8",
    )
    for row in meta.rows[1:]:
        set_cell_shading(row.cells[0], "F6F9FC")
    doc.add_page_break()

    # Authenticity and use note.
    heading("参赛说明与真实性边界", 1)
    callout(
        "材料定位",
        "本计划书以当前仓库源码、测试报告和产品设计文档为基础，面向创新创业赛事形成系统化参赛稿。为便于后续报名，学校、团队、知识产权、运营数据、用户调研、获奖证明等尚未提供的信息均保留为待填写或待佐证项。",
    )
    para("本项目的参赛叙事分为产品愿景、当前可验证能力和后续产品化计划三个层次。当前主交付形态是本地优先的多智能体工作台，已具备会话、工具、权限审批、审计、任务编排、产物交付等技术基础；智能输入法与桌面宠物是统一产品愿景中的交互表层，当前不作为已经完成的发布功能。")
    para("市场规模与财务部分采用情景测算，用于展示商业逻辑和资源配置方法，不代表已经形成真实收入、融资承诺或用户规模。正式提交前应由项目负责人补充真实访谈、试用记录、知识产权、团队经历和财务凭证，并对全部数据签字确认。")
    table(
        ["信息类别", "本稿处理方式", "提交前动作"],
        [
            ("产品与技术", "按源码、契约测试和验收报告分层描述", "补充最新安装包与实机演示证据"),
            ("市场数据", "引用公开权威资料并标明时间", "核对赛事截止日前的最新公开数据"),
            ("团队信息", "保留待填写项", "填写真实姓名、分工、学院与贡献"),
            ("财务预测", "基于价格、转化率和成本假设测算", "按实际运营计划复核并签字"),
            ("知识产权", "不虚构专利、软著或商标", "有证书再附扫描件，无证书则写申请计划"),
        ],
        [1900, 3600, 3860],
        font_size=9.0,
    )
    doc.add_page_break()

    # Headless-stable static table of contents. Page numbers are verified against
    # the rendered competition edition and can be refreshed if content changes.
    heading("目录", 1)
    toc_rows = [
        ("项目概要", "4"),
        ("项目背景与市场机会", "5"),
        ("产品定位与目标用户", "6"),
        ("产品与服务方案", "8"),
        ("技术路线与系统架构", "9"),
        ("研发基础与成熟度", "11"),
        ("项目创新点与核心优势", "11"),
        ("市场分析", "12"),
        ("竞争分析", "13"),
        ("商业模式", "14"),
        ("市场进入与营销策略", "15"),
        ("研发计划与实施路线", "16"),
        ("组织管理与团队建设", "16"),
        ("财务测算与资金规划", "17"),
        ("风险分析与应对", "18"),
        ("教育实效与社会价值", "19"),
        ("参赛展示与答辩方案", "19"),
        ("结论", "20"),
        ("参考资料", "20"),
        ("附录  提交前补充清单", "22"),
    ]
    toc = table(["章节", "页码"], toc_rows, [8200, 1160],
                aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
                font_size=9.3, header_fill="E9F1F8")
    for row in toc.rows[1:]:
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para("目录页码按本参赛稿渲染版核对；如后续增删内容，请重新更新。", size=8.8, color=MUTED, italic=True)
    doc.add_page_break()

    heading("项目概要", 1)
    para("OwO 是一套面向个人电脑场景的情境感知智能交互系统。项目希望把分散在输入、对话、文件、浏览器和任务执行中的人工智能能力，整合为一个持续理解当前任务、透明展示状态、经过用户授权后完成行动的桌面伙伴。产品愿景由情境表达入口、可视化桌面伙伴和任务智能体三层协同组成，分别回答用户现在怎么说、现在在做什么和接下来怎样做完。")
    para("项目当前的工程主线是本地优先多智能体工作台。系统采用 Coordinator、Worker 与 Human 节点协作，围绕任务图、共享工作区和版本化成果进行分工、接力、评审和恢复；所有文件、命令、网络和桌面操作均进入权限策略，涉及写入、发送、发布或其他高风险动作时必须由用户或独立审批器确认。该方案不是简单叠加多个聊天机器人，而是把目标分解、证据采集、执行、复核和产物交付纳入统一状态机。")
    para("在 2026 年 8 月 31 日的单智能体产品评测中，固定套件完成 200 次运行，成功 195 次，总成功率为 97.5%，覆盖代码、文档和研究类任务。该结果证明核心 Agent 执行面已具备较高稳定性，但不等同于全部桌面产品已通过发布验收。WorkSwarm 核心与服务接口已经建立，桌面端关键页面已完成阶段性重构，仍需继续完成全分辨率、异常启动、全部可见按钮真实请求和完整多智能体真实模型流程验收。")
    callout("核心价值主张", "以本地优先、情境连续、多智能体协同、审批审计和结果可恢复为核心，降低用户在多个应用之间反复描述背景、复制粘贴内容和手工串联步骤的成本。", fill=LIGHT_CYAN, accent=CYAN)
    table(
        ["维度", "项目回答"],
        [
            ("服务对象", "学生、教师、内容创作者、办公人员、开发者及有复杂电脑任务的个人用户"),
            ("核心问题", "人工智能与真实工作环境割裂，聊天结果难直接转化为可控行动"),
            ("解决方案", "情境感知入口、可视化状态、多智能体任务执行、权限审批与版本化交付"),
            ("初始落地", "代码、研究、文档和结构化信息处理的本地多智能体工作台"),
            ("长期形态", "覆盖表达、理解、行动与学习复用的个人桌面智能交互系统"),
        ],
        [1900, 7460],
        font_size=9.4,
    )

    heading("项目背景与市场机会", 1)
    heading("行业发展与政策环境", 2)
    para("生成式人工智能正从通用问答向学习、办公、内容创作和任务执行渗透。中国互联网络信息中心发布的第57次中国互联网络发展状况统计报告显示，截至2025年12月，我国生成式人工智能用户规模达到6.02亿人，普及率为42.8%，较2024年底增长141.7%。这说明用户教育已经跨过早期试用阶段，下一轮竞争重点将从是否使用人工智能转向能否嵌入真实工作流、稳定交付结果并建立信任。")
    para("国务院关于深入实施人工智能加行动的意见提出，到2027年新一代智能终端、智能体等应用普及率超过70%，到2030年超过90%，并鼓励发展提效型、陪伴型智能原生应用和智能助理新入口。教育部关于中国国际大学生创新大赛的通知同时强调项目应紧密结合现实需求，促进人工智能、数字技术与教育、消费生活等领域深度融合，并要求科技成果、知识产权、财务状况和运营材料真实合规。OwO 的本地智能体、陪伴式交互和学习办公场景与上述方向具有较高契合度。")
    table(
        ["外部信号", "公开依据", "对 OwO 的启示"],
        [
            ("用户基础扩大", "2025年末生成式AI用户6.02亿，普及率42.8%", "无需从零教育用户，应集中解决持续使用与任务闭环"),
            ("智能体普及加速", "政策提出2027年应用普及率超70%", "需要尽早形成安全、可审计的桌面智能体产品形态"),
            ("终端成为入口", "政策支持智能电脑、智能助理和陪伴型应用", "本地工作台与桌面伙伴具有入口价值"),
            ("赛事重视真实性", "教育部要求成果、知识产权、财务和贡献可核验", "材料必须区分已实现、演示与规划能力"),
        ],
        [1900, 3300, 4160],
        font_size=8.9,
    )

    heading("用户痛点", 2)
    table(
        ["痛点", "典型表现", "现有方式的不足", "OwO 对应方案"],
        [
            ("上下文反复丢失", "切换文档、浏览器和聊天应用后需重复解释", "通用对话以会话为中心，难连续理解桌面任务", "在授权范围内保存任务主题、文件和成果关系"),
            ("回答难变成结果", "模型给出建议后仍需手工复制、改写和执行", "输出停留在文本，缺乏步骤、验证和交付", "将目标拆解为任务图并交付版本化成果"),
            ("复杂任务不可见", "用户不知道系统在做什么、为何失败", "长任务缺少进度、证据和人工介入点", "桌面状态、任务进度、审批和失败原因可视化"),
            ("自动化信任不足", "担心误删文件、误发消息或隐私泄露", "强自治工具可能缺少统一权限与审计", "默认拒绝、独立审批、操作预览、审计和回退"),
        ],
        [1500, 2500, 2500, 2860],
        font_size=8.5,
    )
    heading("项目机会", 2)
    para("OwO 不与基础模型厂商竞争通用知识，而是聚焦模型与用户真实电脑之间的最后一公里。项目将模型能力组织为可落地的情境理解、协同执行和可信交付体系，可在基础模型快速迭代时保持模型可替换、数据本地优先和产品体验连续，从而形成独立于单一模型的系统价值。")

    heading("产品定位与目标用户", 1)
    heading("产品定位", 2)
    callout("一句话定义", "OwO 是能够理解用户当前情境、帮助用户表达，并在授权后协同完成桌面任务的本地优先智能交互系统。")
    figure(ASSET_DIR / "owo-system-ecosystem.png", "图 1  OwO 产品协同生态  表达辅助、桌面伙伴与多智能体执行围绕同一用户任务连续协作", width=6.55)
    table(
        ["产品层", "承担角色", "用户获得的价值", "当前状态"],
        [
            ("情境表达入口", "根据应用、对象和任务给出补全、改写、翻译和回复建议", "减少重复提示，提升表达速度与得体度", "产品规划，输入法融合暂停实施"),
            ("可视化桌面伙伴", "展示系统判断、主动建议、审批、进度和失败原因", "让AI状态可见、可纠正、可控制", "交互愿景，待形成独立桌宠客户端"),
            ("任务智能体工作台", "拆解目标，组织Agent与Human，调用受控工具，交付成果", "把问答升级为完整任务闭环", "当前工程主线，核心能力已进入验证"),
        ],
        [1900, 3100, 2760, 1600],
        font_size=8.8,
    )

    heading("目标用户", 2)
    table(
        ["用户群体", "高频任务", "核心诉求", "首个可验证场景"],
        [
            ("高校学生", "资料阅读、报告写作、课程研究、沟通", "少切换、能追溯、可形成成果", "资料到报告的一体化任务"),
            ("教师与研究人员", "文献整理、证据对照、教学材料", "来源清晰、结构稳定、隐私可控", "多来源研究与文档交付"),
            ("内容与办公人员", "改写、汇总、表格、汇报、文件整理", "效率、格式一致、过程透明", "跨文档整理与版本审阅"),
            ("软件开发者", "代码修复、审查、重构、契约变更", "可执行、可测试、可回滚", "本地仓库任务与差异审阅"),
            ("轻度数字用户", "查找文件、解释界面、重复操作", "低学习成本、明确确认", "桌面伙伴引导与受控自动化"),
        ],
        [1600, 2650, 2550, 2560],
        font_size=8.6,
    )

    heading("典型用户旅程", 2)
    figure(ASSET_DIR / "owo-user-journey.png", "图 2  典型用户旅程  从资料学习到正式写作，再到多步骤成果检查", width=6.55)
    number("用户将课程资料或项目文件交给 OwO，系统识别任务主题并建立受控工作区。", restart=True)
    number("表达入口和桌面伙伴根据当前应用切换学习、写作或沟通状态，用户可随时纠正。")
    number("简单请求即时完成，复杂请求由 Coordinator 拆解，Worker 分别处理检索、撰写、校对或代码任务。")
    number("涉及文件写入、命令、网络提交或桌面操作时，系统展示目标、范围与风险并等待确认。")
    number("系统交付可审阅的文档、代码差异、证据表或其他 Artifact，记录来源、版本、审计和失败信息。")

    heading("产品与服务方案", 1)
    heading("核心功能体系", 2)
    table(
        ["能力模块", "功能说明", "面向用户的结果"],
        [
            ("情境感知", "结合前台应用、授权文件、UI语义、OCR与任务历史形成场景快照", "减少重复描述，让建议与当前工作匹配"),
            ("智能表达", "提供句子补全、风格调整、翻译、摘要和快捷回复", "更快完成学术、职场和日常沟通"),
            ("文件与知识处理", "读取文档、表格、PDF、代码与网页，提取和组织结构化信息", "从散乱材料形成提纲、证据表和正式成果"),
            ("WorkSwarm 协同", "Coordinator与最多三个Worker及Human节点围绕任务图协作", "复杂任务可并行、接力和复核，不必逐个指挥"),
            ("成果交付", "版本化Artifact、差异审阅、评审记录、恢复和回退", "用户掌握最终结果，而非只获得一段回答"),
            ("安全控制", "默认拒绝、按工具授权、独立审批、审计、敏感面熔断", "自动化能力不越过用户边界"),
            ("技能复用", "把验证过的流程沉淀为技能或工作流，并进行健康检查", "重复任务越用越省事，同时保持可审查"),
        ],
        [1900, 4400, 3060],
        font_size=8.8,
    )
    heading("产品交互分级", 2)
    para("OwO 将交互复杂度分为即时辅助、主动建议和长任务执行三个层级。输入补全、润色等低风险需求即时响应；桌面伙伴以不打断的方式提示翻译、总结或待办提取；需要跨文件、多步骤和持续运行的任务则升级到工作台，由用户查看计划、进度、审批和成果。分级设计避免所有需求都被迫进入复杂聊天窗口，也避免轻量入口承担不适合的高风险执行。")
    table(
        ["级别", "适用任务", "交互方式", "控制要求"],
        [
            ("即时辅助", "补全、改写、翻译、解释", "候选或快捷操作", "默认只生成建议，不直接外发"),
            ("主动建议", "检测到可总结、可解释、可复用情境", "桌面伙伴提示卡", "可忽略、静默或禁用"),
            ("长任务执行", "研究、报告、代码、文件与跨应用任务", "工作台任务图与成果区", "计划可见，写操作审批，结果可回退"),
        ],
        [1600, 3100, 2500, 2160],
        font_size=8.8,
    )

    heading("技术路线与系统架构", 1)
    heading("总体架构", 2)
    para("系统采用本地优先、模型可替换、工具受控的分层架构。最上层是桌面工作台及未来的输入法和桌宠入口；中间层由会话、Goal与Plan、WorkSwarm、工作流和Artifact管线构成；执行层统一管理文件、Git、命令、浏览器、API和桌面操作；安全层贯穿权限、审批、审计、凭据、限流和错误语义；底层通过模型网关连接本地或兼容的云端模型。")
    table(
        ["架构层", "关键组件", "主要职责", "设计原则"],
        [
            ("交互层", "桌面工作台、状态可视化、未来输入法与桌宠", "任务入口、进度、审批、差异与成果审阅", "复杂能力保持用户可理解"),
            ("协同层", "Session、Goal、Plan、WorkSwarm、Workflow", "任务分解、成员调度、接力、Human介入", "单一状态源，明确终态"),
            ("执行层", "文件、Git、浏览器、API、受控命令、Computer Use", "在明确范围内执行并验证动作", "确定性工具优先，视觉只作候选"),
            ("资产层", "Project Space、Artifact、ChangeSet、Trace、Memory", "保存成果版本、证据、差异和可复用经验", "成果可追溯、可恢复"),
            ("安全层", "Policy、Approval、Audit、Credential、Sandbox", "阻断越权、保护密钥、记录决策", "默认拒绝，审批与主Agent分离"),
            ("模型层", "本地模型、兼容云端模型、OCR与视觉", "理解目标、规划、生成与感知辅助", "模型可替换，密钥仅由环境注入"),
        ],
        [1450, 2500, 3260, 2150],
        font_size=8.5,
    )

    heading("多智能体协同机制", 2)
    para("Coordinator 负责目标理解、拆解、资源边界和最终交付；Worker 只接收完成其任务所需的上下文、工具和预算；Human 节点在信息缺失、判断冲突或高风险动作前介入。任务以有向无环图表达依赖关系，已完成节点的成果不会因后续 steer 或重试而丢失。系统通过写租约和明确单写者减少并行冲突，通过Artifact与Handoff传递成果，而不是让多个智能体共享无边界的全部会话。")
    heading("情境感知与可靠执行", 2)
    para("感知链路采用分层证据：优先使用应用信息、结构化API、文件状态、浏览器语义引用和Windows UI Automation；OCR和视觉模型用于补充定位或解释，不直接获得高风险执行权。执行动作携带目标、预期效果和验证条件，失败时进入重观察、询问用户、降级或终止等已定义状态，避免模型在不确定界面上连续盲点。")
    heading("安全与隐私设计", 2)
    bullet("权限默认拒绝。文件写入、命令执行、网络访问、文本注入和远程回传均须通过策略。")
    bullet("审批与主智能体分离。主智能体不能为自身的越权工具调用自动授权。")
    bullet("模型密钥仅经环境变量或操作系统凭据引用注入，不进入代码、仓库、日志或普通Worker环境。")
    bullet("敏感窗口、密码、支付、验证码等场景触发熔断；不可逆操作必须预览、确认或提供回退。")
    bullet("审计记录动作主体、目标、权限决定、关联任务和结果摘要，为复盘和责任界定提供依据。")

    heading("研发基础与成熟度", 1)
    heading("当前工程基础", 2)
    table(
        ["能力", "当前证据", "成熟度判断", "下一验收"],
        [
            ("Agent执行与会话", "核心循环、工具调度、会话、项目规则、审计及API契约", "已实现并持续测试", "发布包与真实任务回归"),
            ("单智能体产品评测", "200次固定运行，195次成功，总成功率97.5%", "有量化基线", "修复5次失败并建立新批次复测"),
            ("WorkSwarm", "TeamRun、任务图、Artifact、Handoff、Human、steer、恢复及HTTP资源面", "核心与服务面已实现", "完整真实模型流程和桌面端验收"),
            ("桌面工作台", "五个一级页面已阶段性点击验证，关键面板持续拆分", "可演示但仍在产品化", "全分辨率、异常启动和全部按钮请求"),
            ("感知与桌面操作", "UIA、OCR、场景图、动作程序、结构化断言和模拟环境", "技术模块具备基础", "更多真实应用兼容矩阵"),
            ("输入法与桌宠", "已有产品方案和历史输入法工程", "统一产品愿景，当前暂停", "重新立项前完成交互与高敏输入评审"),
        ],
        [1700, 3520, 1900, 2240],
        font_size=8.5,
    )
    callout("成熟度说明", "测试通过证明相应源码与契约在给定环境下工作，不自动等同于安装包、桌面壳、所有外部应用和真实用户体验已完成。参赛演示应使用经过复验的固定任务与固定发布包。", fill="FFF7E8", accent=GOLD)

    heading("项目创新点与核心优势", 1)
    heading("产品创新", 2)
    table(
        ["创新点", "传统方式", "OwO 方案", "可验证价值"],
        [
            ("感知、表达、行动闭环", "输入法、桌宠、聊天助手彼此割裂", "统一情境和任务主线，按复杂度升级", "减少反复描述与应用切换"),
            ("可见的AI状态", "系统判断隐藏在对话或后台", "用桌面伙伴和工作台展示状态、进度、审批和失败", "提高可解释性和纠错效率"),
            ("成果中心协同", "多个Agent输出零散文本", "围绕版本化Artifact、任务图和Handoff协作", "减少冲突，交付结果可追溯"),
            ("本地优先安全边界", "自动化工具权限粒度不一", "统一默认拒绝、独立审批、审计和回退", "在提升效率时保留用户控制权"),
            ("过程经验可复用", "重复任务仍需重新提示", "把验证过的流程沉淀为技能、工作流和团队模板", "形成可审计的使用飞轮"),
        ],
        [1780, 2400, 3220, 1960],
        font_size=8.4,
    )
    heading("技术与工程优势", 2)
    bullet("Rust 核心与契约化HTTP接口有利于构建稳定、可测试的本地执行底座。")
    bullet("模型网关支持本地或兼容云端模型，避免产品价值绑定单一模型厂商。")
    bullet("权限、审批、审计、diff与revert是主流程，不是发布后追加的外围功能。")
    bullet("WorkSwarm限制成员数量和单写边界，在获得协同收益的同时控制复杂度。")
    bullet("产品评测套件提供重复运行、成功率、调用量、耗时和失败分类，可持续量化迭代。")

    heading("市场分析", 1)
    heading("市场空间判断", 2)
    para("OwO 所处市场不是单一输入法或单一聊天软件，而是个人生产力智能体、AI桌面助手和安全自动化的交叉市场。公开数据表明生成式人工智能用户已形成大规模基础，政策正在推动智能终端和智能体加速普及。项目的首要任务不是追求覆盖全部6.02亿用户，而是在学习、研究、文档和开发等高频电脑任务中找到愿意持续使用并付费的早期用户。")
    table(
        ["口径", "测算方法", "规模表达", "说明"],
        [
            ("潜在用户池", "CNNIC 2025年末生成式AI用户", "6.02亿人", "外部公开统计，不等同于OwO可服务市场"),
            ("可服务市场", "假设其中8%至12%为高频PC学习办公用户", "约4816万至7224万人", "内部情景假设，需用调研校准"),
            ("三年可获得用户", "按校园、开发者社区和内容渠道逐步扩散", "注册用户30万人", "经营目标，不是现有用户规模"),
            ("三年付费用户", "按第三年12%付费转化情景", "约3.6万人", "与产品稳定性、渠道和定价高度相关"),
        ],
        [1700, 3100, 2100, 2460],
        font_size=8.7,
    )
    para("上述市场测算使用自上而下的情景法，仅用于说明增长路径。正式商业验证应增加至少三类证据：目标用户访谈与任务日志、可用版本的留存和付费意愿测试、校园或团队试点的实际采购反馈。", size=9.2, color=MUTED, italic=True)

    heading("客户需求优先级", 2)
    table(
        ["细分市场", "需求强度", "付费可能", "进入难度", "进入顺序"],
        [
            ("高校学生与科研团队", "高", "中", "低至中", "第一阶段"),
            ("独立开发者与小型技术团队", "高", "中至高", "中", "第一阶段"),
            ("内容与知识工作者", "中至高", "中", "中", "第二阶段"),
            ("学校与实验室私有部署", "高", "高", "高", "第二阶段"),
            ("普通大众桌面伙伴", "中", "低至中", "高", "第三阶段"),
        ],
        [2300, 1300, 1550, 1550, 2660],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
        font_size=8.8,
    )

    heading("竞争分析", 1)
    heading("竞争格局", 2)
    para("竞争对手应按产品类别而非单一品牌理解。通用AI助手在模型能力和生态方面强，操作系统助手拥有系统入口，AI输入法具有低摩擦表达优势，RPA在固定流程上稳定，开源Agent框架便于技术团队定制。OwO 的差异化不在于每一项能力都超过这些成熟产品，而在于将本地优先、多智能体成果协作、权限审批和情境可视化组合成一条适合个人生产力的完整链路。")
    table(
        ["类别", "代表能力", "优势", "项目组观察到的空白", "OwO应对"],
        [
            ("通用AI桌面助手", "ChatGPT桌面及应用连接", "模型与服务生态强", "本地工作区深度、可替换模型与审计要求因产品而异", "聚焦本地项目、受控工具与成果交付"),
            ("操作系统助手", "Microsoft Copilot Vision、文件搜索、截图与设置引导", "系统入口和应用生态强", "个人多Agent成果协作不是唯一主线", "以WorkSwarm和版本化Artifact形成差异"),
            ("AI输入与写作工具", "补全、改写、翻译、快捷回复", "使用频率高、学习成本低", "通常不承担复杂任务编排和执行", "作为未来交互入口连接长任务工作台"),
            ("RPA与自动化", "固定流程和企业系统自动化", "规则流程稳定、企业价值明确", "配置成本较高，对开放环境适应有限", "结合自然语言规划、确定性工具与用户审批"),
            ("开源Agent框架", "工具调用、多Agent编排、代码扩展", "开发自由度高", "普通用户产品化、安装与安全治理门槛高", "提供桌面产品壳、评测、权限与恢复体系"),
        ],
        [1500, 2150, 1800, 2250, 1660],
        font_size=8.0,
    )
    para("注：竞品信息依据各产品官方公开页面整理，空白与应对为项目组基于公开功能的比较判断，不代表对竞争产品的完整测评。", size=8.8, color=MUTED, italic=True)

    heading("SWOT 分析", 2)
    table(
        ["优势 Strengths", "劣势 Weaknesses", "机会 Opportunities", "威胁 Threats"],
        [
            ("本地优先与权限审计进入主架构；已有核心代码和量化评测；WorkSwarm围绕成果协作；模型可替换。",
             "统一产品形态尚未完全落地；输入法和桌宠仍属规划；桌面兼容矩阵与真实用户数据不足；团队和商业资源待补齐。",
             "生成式AI用户快速增长；政策推动智能体和AI终端；高校学习办公场景集中；用户对隐私与可控自动化需求上升。",
             "大型平台快速整合系统入口；基础模型能力同质化；隐私安全事故影响信任；推理成本和接口政策波动。"),
        ],
        [2340, 2340, 2340, 2340],
        font_size=8.6,
        header_fill="E9F1F8",
    )

    heading("商业模式", 1)
    heading("价值交换与收入来源", 2)
    table(
        ["产品形态", "目标客户", "核心权益", "建议定价", "收入方式"],
        [
            ("个人免费版", "学生与轻度用户", "基础对话、少量本地任务、核心安全功能", "免费", "获客与口碑"),
            ("个人专业版", "高频知识工作者与开发者", "更高任务额度、WorkSwarm、专业技能包、优先支持", "19至39元/月", "订阅收入"),
            ("校园与团队版", "实验室、课程团队、小型组织", "团队空间、管理员策略、统一部署和使用报告", "按席位或项目报价", "许可与服务"),
            ("私有部署版", "有数据合规要求的机构", "本地模型、内网部署、定制集成和审计", "项目制报价", "实施与维护"),
            ("技能与模板生态", "开发者、教师、服务商", "技能包、工作流和团队模板分发", "交易分成", "平台服务"),
        ],
        [1550, 2100, 3000, 1500, 1210],
        font_size=8.4,
    )
    heading("成本结构", 2)
    bullet("研发成本：核心引擎、桌面端、兼容性、安全、测试和发布工程。")
    bullet("模型与基础设施：云端推理、下载与更新服务、遥测和错误诊断；本地模型可降低部分边际成本。")
    bullet("市场与服务：校园试点、开发者社区、内容运营、用户支持和机构交付。")
    bullet("合规与知识产权：隐私评审、软件著作权、商标、第三方许可证与安全测试。")
    heading("商业验证顺序", 2)
    number("先验证任务价值：目标用户是否愿意持续把真实资料、代码或文档任务交给系统。", restart=True)
    number("再验证留存：首周、四周留存和每周完成任务数是否达到预设门槛。")
    number("再验证付费：专业额度、团队协作、私有部署和专用技能中哪一项产生真实付费。")
    number("最后扩展生态：只有在核心任务闭环稳定后，才开放技能市场和第三方开发者收入。")

    heading("市场进入与营销策略", 1)
    heading("进入路径", 2)
    table(
        ["阶段", "目标", "核心动作", "衡量指标"],
        [
            ("种子验证", "找到3个高频任务", "校内访谈、可用性测试、开发者内测、固定任务复现", "30名深度用户、100次真实任务、任务完成率"),
            ("校园试点", "形成可展示案例", "课程资料到报告、科研证据表、代码任务工作坊", "3个试点团队、4周留存、可引用反馈"),
            ("社区增长", "扩大专业用户", "开源部分SDK、发布技能模板、技术内容与演示视频", "注册、激活、周任务数、推荐率"),
            ("机构合作", "验证团队和私有部署", "与实验室、创新中心、软件团队共建", "试点转付费、交付周期、续费意向"),
        ],
        [1600, 2000, 3600, 2160],
        font_size=8.6,
    )
    heading("品牌与传播", 2)
    para("传播内容围绕一个统一故事展开：用户阅读资料，桌面伙伴识别并处理，打开文档后获得情境写作帮助，切换沟通场景时得到得体表达，最后由多智能体工作台完成研究、报告或代码任务。所有演示都应展示审批、进度和结果证据，避免只呈现炫酷动画而缺少真实闭环。")
    bullet("赛事传播：围绕痛点、创新、实机证据和社会价值制作路演PPT与3分钟演示。")
    bullet("校园传播：与课程项目、实验室和创新训练结合，收集真实任务与反馈。")
    bullet("开发者传播：公开接口契约、评测方法和示例技能，建立可信技术形象。")
    bullet("内容传播：制作从一个目标到可审阅成果的过程型案例，突出透明与可控。")

    heading("研发计划与实施路线", 1)
    heading("阶段目标", 2)
    table(
        ["阶段", "时间", "产品目标", "关键交付", "验收门槛"],
        [
            ("阶段一 参赛稳定版", "0至3个月", "收敛本地多Agent工作台", "安装包、固定演示、失败复测、帮助文档", "核心流程可重复，发布包与源码一致"),
            ("阶段二 校园试用版", "4至8个月", "服务学习、研究、文档和代码任务", "3类技能包、试点管理、反馈与留存分析", "30名深度用户，关键任务成功率达标"),
            ("阶段三 产品融合版", "9至15个月", "加入桌面伙伴状态和轻量建议", "桌宠交互、情境切换、可纠正偏好", "不打断、可关闭、权限和隐私通过评审"),
            ("阶段四 规模化版", "16至24个月", "团队许可与生态扩展", "管理策略、私有部署、技能分发", "形成付费试点和可复制交付流程"),
            ("后续研究", "24个月后", "评估输入法入口和高级桌面自治", "TSF方案、兼容矩阵、影子预演", "独立立项，不降低安全基线"),
        ],
        [1400, 1200, 2600, 2500, 1660],
        font_size=8.2,
    )
    heading("里程碑与决策门", 2)
    bullet("M1 可参赛：计划书、路演、演示包、测试报告和真实性材料闭环。")
    bullet("M2 可试用：安装、升级、错误恢复、数据迁移和隐私说明达到内测标准。")
    bullet("M3 可留存：目标用户连续四周使用，任务成功、时间节省和满意度有数据。")
    bullet("M4 可付费：至少一种权益被真实购买，支持成本和模型成本可控。")
    bullet("M5 可扩展：团队版或技能生态不破坏权限、审计、兼容和恢复指标。")

    heading("组织管理与团队建设", 1)
    heading("建议团队结构", 2)
    para("因当前未提供参赛成员信息，本节提供可直接替换的组织模板。最终提交必须按成员真实贡献填写，不得仅为满足人数而挂名。")
    table(
        ["角色", "建议职责", "当前负责人", "参赛佐证"],
        [
            ("项目负责人／产品", "产品定位、用户调研、进度、路演与跨组协调", "待填写", "需求文档、访谈、版本记录"),
            ("核心引擎负责人", "Agent循环、工具、权限、会话、审计和模型网关", "待填写", "代码提交、测试与设计文档"),
            ("多Agent负责人", "WorkSwarm、任务图、Artifact、恢复与评测", "待填写", "模块代码、验收报告"),
            ("桌面与交互负责人", "桌面工作台、桌宠交互、可用性和视觉系统", "待填写", "原型、前端代码、测试记录"),
            ("市场与运营负责人", "市场调研、校园试点、合作、内容和商业验证", "待填写", "调研表、试点协议、运营数据"),
            ("指导教师与顾问", "技术、商业、法律、赛事与行业指导", "待填写", "指导记录与真实简介"),
        ],
        [1800, 3900, 1500, 2160],
        font_size=8.5,
    )
    heading("协作与质量管理", 2)
    bullet("代码协作按文件认领，同一文件同一时间只允许一名成员修改，核心接口变更同步契约测试。")
    bullet("功能完成以源码、测试、构建产物和真实交互四类证据综合判断，不以文档声明替代验收。")
    bullet("每两周形成产品评审，检查用户价值、进度、风险、数据真实性和演示可复现性。")
    bullet("公开材料由负责人、技术负责人和指导教师三方复核，确保不夸大能力、不泄露凭据。")

    heading("财务测算与资金规划", 1)
    heading("三年经营情景", 2)
    para("以下为参赛用途的中性情景测算。收入由个人订阅、团队许可、私有部署与服务构成；成本包括研发人力、模型与基础设施、市场服务、合规和日常运营。所有数据均为规划假设，正式商业计划应由团队根据实际定价、用户转化和合同复核。")
    table(
        ["指标", "第一年", "第二年", "第三年", "关键假设"],
        [
            ("累计注册用户", "1万人", "8万人", "30万人", "校园与开发者社区逐步扩散"),
            ("付费用户数", "500人", "6400人", "3.6万人", "付费率5%、8%、12%"),
            ("团队与机构客户", "5个", "20个", "60个", "先试点后复制"),
            ("营业收入", "24万元", "180万元", "650万元", "订阅、许可和服务综合"),
            ("经营成本", "52万元", "150万元", "380万元", "研发投入前置，规模后成本增加"),
            ("经营结果", "负28万元", "30万元", "270万元", "第二年中后期达到盈亏平衡"),
        ],
        [1900, 1300, 1300, 1300, 3560],
        font_size=8.5,
    )
    heading("首轮资金需求情景", 2)
    table(
        ["用途", "比例", "金额情景", "主要产出"],
        [
            ("产品与研发", "45%", "36万元", "发布稳定性、桌面端、技能与兼容测试"),
            ("市场与试点", "20%", "16万元", "校园试点、用户研究、内容与活动"),
            ("模型与基础设施", "15%", "12万元", "推理、更新、监控和测试资源"),
            ("合规与知识产权", "10%", "8万元", "软著、商标、隐私与安全评审"),
            ("运营与预备金", "10%", "8万元", "支持、设备、差旅与不确定性缓冲"),
        ],
        [2500, 1200, 1700, 3960],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.8,
    )
    callout("财务纪律", "在真实收入形成前，不以市场规模代替现金流。模型调用、用户支持和机构交付必须建立单任务成本与毛利监控；若付费转化低于假设，应优先缩小范围而非盲目扩大获客。", fill="FFF7E8", accent=GOLD)

    heading("风险分析与应对", 1)
    table(
        ["风险", "概率", "影响", "预警信号", "应对措施"],
        [
            ("技术可靠性", "中", "高", "任务失败、误操作、恢复失败", "固定评测、分层执行、断言、灰度与回退"),
            ("隐私与安全", "中", "极高", "密钥泄露、越权、敏感数据外发", "默认拒绝、独立审批、密钥隔离、审计与安全测试"),
            ("产品范围失控", "高", "高", "同时推进输入法、桌宠、Agent导致交付延迟", "以工作台为主线，其他表层按决策门立项"),
            ("用户留存不足", "中", "高", "体验新鲜但四周留存低", "聚焦高频任务，以真实时间节省和成果质量验证"),
            ("平台竞争", "高", "中至高", "系统厂商快速集成功能", "强调本地、可替换模型、开源扩展和成果协同"),
            ("成本波动", "中", "中", "模型价格或调用量上升", "本地模型、模型分级、预算上限和成本可视化"),
            ("合规与知识产权", "中", "高", "数据授权不清、第三方许可冲突", "数据最小化、许可证清单、真实材料审查"),
            ("团队持续性", "中", "高", "核心成员投入下降", "职责备份、文档化、模块所有权和阶段激励"),
        ],
        [1750, 900, 1000, 2500, 3210],
        font_size=8.1,
    )

    heading("教育实效与社会价值", 1)
    heading("教育实效", 2)
    bullet("专创融合：把人工智能、软件工程、人机交互、隐私安全和商业验证组织为同一项目实践。")
    bullet("能力培养：团队成员经历需求、架构、开发、测试、路演、用户研究和项目管理的完整过程。")
    bullet("课程支持：为资料整理、证据对照、报告写作和代码实践提供可审阅的智能协作工具。")
    bullet("开放示范：通过评测套件、契约和技能模板沉淀可复用的创新训练资源。")
    heading("社会价值", 2)
    para("OwO 的社会价值不只在于提高输入或生成速度，更在于让普通用户理解人工智能正在做什么、能否执行以及如何纠正。透明的审批和审计可以降低自动化误操作风险；本地优先与模型可替换有助于保护学习、研究和工作资料；陪伴式入口可降低复杂智能体的使用门槛，使更多用户能够获得可控的数字生产力工具。")
    table(
        ["价值方向", "项目贡献", "建议评估指标"],
        [
            ("学习效率", "减少资料整理和格式处理，把时间转向理解与创造", "任务节省时间、成果质量、学习者满意度"),
            ("数字包容", "用可视化伙伴和分级交互降低操作门槛", "首次成功率、求助次数、不同用户群可用性"),
            ("可信人工智能", "让权限、审批、证据和结果可见", "越权拦截、审计覆盖、误操作与恢复率"),
            ("创新就业", "培养AI产品、Agent工程、交互和运营复合能力", "参与学生、实习岗位、技能与案例产出"),
        ],
        [1900, 4300, 3160],
        font_size=8.8,
    )

    heading("参赛展示与答辩方案", 1)
    heading("三分钟演示主线", 2)
    number("痛点开场：用户在资料、文档和沟通应用间切换，重复复制内容并重新解释任务。", restart=True)
    number("展示统一入口：导入资料，系统建立任务上下文，桌面状态可见且可纠正。")
    number("展示长任务：Coordinator拆分研究、撰写和校对任务，Worker并行处理并持续回传进度。")
    number("展示可信控制：文件写入或外部操作前出现审批，用户查看范围后确认。")
    number("展示成果：打开版本化Artifact、证据和差异，说明失败可追溯、成果可恢复。")
    heading("评委高频问题准备", 2)
    table(
        ["可能问题", "建议回答要点", "必须出示的证据"],
        [
            ("与通用AI助手有什么不同", "不竞争基础模型，聚焦本地工作区、多Agent成果协作、审批审计和恢复", "竞品矩阵、实机任务闭环"),
            ("输入法和桌宠是否已经完成", "它们是统一产品愿景；当前交付主线是多Agent工作台，明确说明阶段", "成熟度表、路线图"),
            ("多Agent真的更好吗", "只在预设任务满足质量、成功率或耗时收益时启用，不默认堆Agent", "单Agent基线、WorkSwarm评测计划"),
            ("如何避免误操作和泄密", "默认拒绝、独立审批、凭据隔离、敏感熔断、审计和回退", "审批流程与安全测试"),
            ("商业模式是否成立", "先以高频任务验证留存和付费，再扩展团队、私有部署与技能生态", "用户访谈、试点、付费测试"),
        ],
        [2000, 4700, 2660],
        font_size=8.4,
    )
    heading("提交材料清单", 2)
    bullet("项目计划书、路演PPT、三分钟演示视频或现场演示脚本。")
    bullet("最新可安装版本、版本哈希、复现实验说明和演示数据。")
    bullet("源码与功能证据清单、测试报告、失败记录与修复复测。")
    bullet("团队成员与实质贡献说明、指导教师信息和签字承诺。")
    bullet("软件著作权、专利、商标、获奖、用户反馈或试点协议等真实附件。")
    bullet("市场调研原始表、访谈纪要、财务假设表与引用来源。")

    heading("结论", 1)
    para("OwO 选择了一条兼顾用户体验与安全边界的桌面智能体路线。它以情境连续性连接理解、表达与行动，以多智能体协同处理复杂任务，以版本化成果替代零散回答，并把权限、审批、审计和恢复作为产品核心。当前项目已经形成可持续开发的本地Agent工程基础和量化评测方法，下一阶段应集中完成发布稳定性、真实用户验证和工作台产品化，再逐步扩展桌面伙伴、表达入口和商业生态。")
    para("参赛的关键不在于展示最多功能，而在于用一条真实、可复现的任务证明项目能够解决明确痛点，并诚实说明已完成、可演示和后续规划的边界。只要团队持续以真实证据驱动迭代，OwO 有机会成为面向学习、研究、文档与开发场景的可信个人智能工作台，并进一步演化为具有陪伴感和情境理解能力的桌面智能交互系统。")

    heading("参考资料", 1)
    references = [
        ("教育部关于举办中国国际大学生创新大赛（2026）的通知", "https://hudong.moe.gov.cn/srcsite/A08/s5672/202607/t20260731_1445670.html"),
        ("中国国际大学生创新大赛高教主赛道商业计划书参考模板", "https://cxcyxy.whmc.edu.cn/info/1091/2592.htm"),
        ("CNNIC 第57次中国互联网络发展状况统计报告", "https://www3.cnnic.cn/n4/2026/0304/c88-11549.html"),
        ("CNNIC 生成式人工智能应用发展报告（2025）发布信息", "https://www2.cnnic.cn/n4/2025/1020/c326-11388.html"),
        ("国务院关于深入实施人工智能加行动的意见", "https://www.gov.cn/gongbao/2025/issue_12266/material/gwygb202525.pdf"),
        ("Microsoft Support  Copilot on Windows 功能说明", "https://support.microsoft.com/en-us/microsoft-copilot/getting-started-with-copilot-on-windows"),
        ("OpenAI Help Center  Apps in ChatGPT", "https://help.openai.com/en/articles/11487775-apps-in-chatgpt"),
        ("OwO Agent SDK README 与当前源码", str(ROOT / "agent-sdk" / "README.md")),
        ("OwO V1-R1 产品评测验收报告", str(ROOT / "agent-sdk" / "evals" / "v1" / "results" / "accept-20260831-142050" / "acceptance-report.md")),
        ("OwO 主开发技术文档 Agent SDK v1", str(ROOT / "builGoal" / "主开发技术文档-Agent-SDK-v1.md")),
    ]
    reference_num_id = clone_num_id(doc, num_ids["decimal"])
    for idx, (title, url) in enumerate(references, start=1):
        p = doc.add_paragraph()
        apply_numbering(p, reference_num_id)
        r = p.add_run(title + "  ")
        set_run_font(r, size=9.3, color=INK)
        if url.startswith("http"):
            add_hyperlink(p, url, url)
        else:
            r2 = p.add_run(url)
            set_run_font(r2, size=8.5, color=MUTED)

    heading("附录  提交前补充清单", 1)
    table(
        ["事项", "当前状态", "负责人", "完成标准"],
        [
            ("比赛名称、赛道与格式", "待确认", "项目负责人", "对照当届通知逐项通过"),
            ("团队成员与真实贡献", "待填写", "项目负责人", "成员确认并有证据"),
            ("用户调研与试点", "待补充", "市场负责人", "原始问卷、访谈和反馈可追溯"),
            ("知识产权与许可证", "待核验", "技术负责人", "证书或申请材料、依赖清单完整"),
            ("最新发布包实机验收", "部分完成", "技术负责人", "安装、配对、核心流程、恢复和卸载通过"),
            ("多Agent收益证据", "待专项评测", "评测负责人", "与单Agent相比至少一项指标显著改善"),
            ("财务假设复核", "情景测算", "财务负责人", "价格、转化、成本有依据并签字"),
            ("PPT与演示视频", "待制作", "路演负责人", "3分钟主线清晰、无夸大、可复现"),
        ],
        [2600, 1500, 1800, 3460],
        font_size=8.5,
    )

    core = doc.core_properties
    core.title = "OwO 情境感知智能交互系统参赛项目计划书"
    core.subject = "大学生创新创业竞赛项目计划书"
    core.author = "OwO 项目团队"
    core.keywords = "OwO, 情境感知, 桌面智能体, 多智能体, WorkSwarm, 创新创业"
    core.comments = "基于当前源码、测试报告、产品文档与公开赛事资料生成的参赛初稿。"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
