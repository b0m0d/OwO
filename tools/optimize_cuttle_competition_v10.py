from __future__ import annotations

import argparse
import importlib.util
import io
import math
import re
import tempfile
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageEnhance, ImageFilter
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt


ROOT = Path(r"T:\创新创业\OwO-master")
SOURCE = ROOT / "docs" / "Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
OUTPUT = ROOT / "docs" / "Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉精修版-v10.docx"
ASSET_DIR = ROOT / "docs" / "cuttle_v10_assets"
GPT_CAPSULE_SOURCE = Path(
    r"C:\Users\23843\.codex\generated_images\01a094c3-6e31-7743-96af-a80f753acd12\exec-0aa36619-c51c-4a36-b625-9ef451ecbf83.png"
)

spec = importlib.util.spec_from_file_location("cuttle_v9", ROOT / "tools" / "optimize_cuttle_competition_v9.py")
v9 = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(v9)


def style_cover(doc: Document, _source: Path):
    paragraphs = doc.paragraphs
    p0, p1, p2, p3, p4 = paragraphs[:5]

    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(22)
    p0.paragraph_format.space_after = Pt(15)
    for run in p0.runs:
        v9.set_run(run, name=v9.FONT_HEAD, size=10.5, bold=True, color=v9.TEAL)

    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(18)
    if doc.inline_shapes:
        logo = doc.inline_shapes[0]
        logo.width = Inches(1.04)
        logo.height = Inches(1.04)

    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(4)
    p2.paragraph_format.space_after = Pt(10)
    for run in p2.runs:
        v9.set_run(run, name=v9.FONT_HEAD, size=28, bold=True, color=v9.GRAPHITE)

    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(18)
    for run in p3.runs:
        v9.set_run(run, name=v9.FONT_HEAD, size=14.5, bold=True, color=v9.TEAL)

    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.left_indent = Cm(1.15)
    p4.paragraph_format.right_indent = Cm(1.15)
    p4.paragraph_format.first_line_indent = Cm(0)
    p4.paragraph_format.line_spacing = 1.45
    p4.paragraph_format.space_after = Pt(28)
    for run in p4.runs:
        v9.set_run(run, size=10.8, color=v9.INK)

    cover_table = doc.tables[0]
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_table.autofit = False
    v9.remove_all_table_borders(cover_table)
    for row in cover_table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Cm(0.84)
        v9.prevent_row_split(row)
        for col_index, cell in enumerate(row.cells):
            v9.set_shading(cell, v9.PAPER)
            v9.set_cell_margins(cell, top=50, start=80, bottom=50, end=80)
            v9.set_cell_width(cell, 1.35 if col_index == 0 else 4.95)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    v9.set_run(
                        run,
                        name=v9.FONT_HEAD if col_index == 0 else v9.FONT_CN,
                        size=8.8 if col_index == 0 else 10.2,
                        bold=col_index == 1,
                        color=v9.MUTED if col_index == 0 else v9.GRAPHITE,
                    )


def inject_ic_explanation(doc: Document):
    addition = (
        "之所以称为“胶囊”，是因为它把一次任务所需的六类信息像载荷一样封装进同一对象，"
        "并设置明确的读取边界与失效条件；后续模块只解封当前步骤所需字段，避免整段上下文无差别流转。"
    )
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("意图胶囊（Intent Capsule, IC）是本项目最具定义性的技术对象"):
            if addition not in paragraph.text:
                paragraph.add_run(addition)
            break


def palette_only(image: Image.Image) -> Image.Image:
    image = v9.palette_enhance(image, crop_title=False)
    return image


def reframe_crop(image: Image.Image, box: tuple[int, int, int, int]) -> Image.Image:
    original_size = image.size
    image = image.crop(box)
    image = palette_only(image)
    canvas = Image.new("RGB", original_size, (252, 251, 248))
    margin_x, margin_y = 54, 26
    scale = min((original_size[0] - margin_x * 2) / image.width, (original_size[1] - margin_y * 2) / image.height)
    resized = image.resize((round(image.width * scale), round(image.height * scale)), Image.Resampling.LANCZOS)
    x = (original_size[0] - resized.width) // 2
    y = (original_size[1] - resized.height) // 2
    canvas.paste(resized, (x, y))
    return canvas


def draw_centered(draw: ImageDraw.ImageDraw, xy, text, font, fill):
    box = draw.textbbox((0, 0), text, font=font)
    draw.text((xy[0] - (box[2] - box[0]) / 2, xy[1]), text, font=font, fill=fill)


def rebuild_ic_capsule(width=1700, height=930) -> Image.Image:
    canvas = Image.new("RGB", (width, height), (252, 251, 248))
    draw = ImageDraw.Draw(canvas)
    labels = [
        ("I", "Intent", "任务意图", (190, 92, 62)),
        ("C", "Context", "最小情境", (8, 122, 125)),
        ("O", "Origin", "来源位置", (185, 138, 47)),
        ("P", "Provenance", "证据来源", (38, 124, 91)),
        ("F", "Freshness", "时效", (51, 83, 106)),
        ("S", "Scope", "授权范围", (141, 111, 75)),
    ]
    cx, cy = width // 2, height // 2
    ring = 245
    points = []
    for index in range(6):
        angle = -math.pi / 2 + index * (math.pi / 3)
        points.append((cx + int(ring * math.cos(angle)), cy + int(ring * math.sin(angle))))
    draw.ellipse((cx-322, cy-322, cx+322, cy+322), outline=(225, 220, 210), width=3)
    draw.ellipse((cx-270, cy-270, cx+270, cy+270), outline=(215, 211, 203), width=12)
    for index, ((letter, english, chinese, color), (x, y)) in enumerate(zip(labels, points)):
        draw.line((cx, cy, x, y), fill=color, width=7)
        draw.ellipse((x-72, y-72, x+72, y+72), fill=(247, 245, 240), outline=color, width=9)
        draw.ellipse((x-37, y-37, x+37, y+37), fill=color)
        draw_centered(draw, (x, y-22), letter, v9.font(42, True), (255, 255, 255))
        if y < cy:
            draw_centered(draw, (x, y-155), english, v9.font(28, True), color)
            draw_centered(draw, (x, y-120), chinese, v9.font(25), (102, 109, 115))
        else:
            draw_centered(draw, (x, y+115), english, v9.font(28, True), color)
            draw_centered(draw, (x, y+150), chinese, v9.font(25), (102, 109, 115))
    draw.ellipse((cx-138, cy-138, cx+138, cy+138), fill=(36, 44, 51))
    draw.ellipse((cx-112, cy-112, cx+112, cy+112), outline=(8, 122, 125), width=6)
    draw_centered(draw, (cx, cy-38), "IC", v9.font(58, True), (255, 255, 255))
    draw_centered(draw, (cx, cy+35), "可校验 · 可失效 · 可追溯", v9.font(24), (220, 220, 216))
    return canvas


def rebuild_market_funnel(width=1800, height=850) -> Image.Image:
    canvas = Image.new("RGB", (width, height), (252, 251, 248))
    draw = ImageDraw.Draw(canvas)
    stages = [
        ((190, 130, 1610, 300), (36, 44, 51), "交付容量", "≤ 200 人", "真实场景承载"),
        ((350, 330, 1450, 500), (8, 122, 125), "深度观察", "50–80 人", "纵向核心组 + 独立验证组"),
        ((520, 530, 1280, 700), (185, 138, 47), "正式实验", "35–55 人", "对照实验 + 消融实验"),
    ]
    for index, (rect, color, label, value, note) in enumerate(stages):
        x0, y0, x1, y1 = rect
        points = [(x0 + 50, y0), (x1 - 50, y0), (x1, y1), (x0, y1)]
        draw.polygon(points, fill=color)
        draw.text((x0 + 80, y0 + 42), f"0{index + 1}  {label}", font=v9.font(32, True), fill=(255, 255, 255))
        draw.text((x0 + 540, y0 + 31), value, font=v9.font(55, True), fill=(255, 255, 255))
        # The band itself carries the stage meaning; omit a secondary note so the visual stays clean.
        if index < 2:
            draw.polygon([(895, y1 + 14), (925, y1 + 14), (910, y1 + 38)], fill=(183, 178, 168))
    draw_centered(draw, (900, 758), "样本逐级收敛", v9.font(28, True), (102, 109, 115))
    return canvas


def rebuild_opportunity(width=1800, height=760) -> Image.Image:
    canvas = Image.new("RGB", (width, height), (252, 251, 248))
    draw = ImageDraw.Draw(canvas)
    colors = [(36, 81, 122), (182, 91, 62), (8, 122, 125)]
    stages = [
        (300, "01", "任务现场", "资料、对象与输入位置"),
        (900, "02", "用户承担", "手工搬运上下文与回忆"),
        (1500, "03", "可验证行动", "从输入焦点回到原位置"),
    ]
    draw.line((260, 385, 1540, 385), fill=(215, 211, 203), width=9)
    for i, (x, number, title, detail) in enumerate(stages):
        color = colors[i]
        r = 116 if i == 1 else 96
        draw.ellipse((x-r, 385-r, x+r, 385+r), fill=(246, 243, 236), outline=color, width=8)
        draw.ellipse((x-42, 343, x+42, 427), fill=color)
        draw_centered(draw, (x, 355), number, v9.font(38, True), (255, 255, 255))
        draw_centered(draw, (x, 548), title, v9.font(34, True), color)
        draw_centered(draw, (x, 594), detail, v9.font(25), (102, 109, 115))
        if i < 2:
            draw.polygon([(x+150, 385), (x+184, 368), (x+184, 402)], fill=color)
    draw_centered(draw, (900, 92), "输入位置、任务承接与验证结果应当在同一条链路上闭合", v9.font(30, True), (36, 44, 51))
    return canvas


def rebuild_loop(width=1800, height=700) -> Image.Image:
    canvas = Image.new("RGB", (width, height), (252, 251, 248))
    draw = ImageDraw.Draw(canvas)
    cx, cy, radius = 900, 360, 205
    draw.ellipse((cx-radius, cy-radius, cx+radius, cy+radius), outline=(215, 211, 203), width=28)
    colors = [(36, 81, 122), (8, 122, 125), (185, 138, 47), (182, 91, 62), (107, 82, 171), (36, 44, 51)]
    labels = [("输入时刻", "文字 / 选区 / 语音"), ("意图胶囊", "对象 · 来源 · 时效"), ("意图路由", "文本 / 工具 / Agent"), ("受控执行", "权限 · 预算 · 审批"), ("成果验证", "证据 · 差异 · 验收"), ("返回原处", "Word / 浏览器 / IDE")]
    for i, ((title, detail), color) in enumerate(zip(labels, colors)):
        angle = -math.pi / 2 + i * (2 * math.pi / 6)
        x = cx + int(radius * math.cos(angle))
        y = cy + int(radius * math.sin(angle))
        draw.ellipse((x-42, y-42, x+42, y+42), fill=color)
        draw_centered(draw, (x, y-17), f"0{i+1}", v9.font(26, True), (255, 255, 255))
        tx = x + (100 if x >= cx else -100)
        align = "left" if x >= cx else "right"
        box = draw.textbbox((0, 0), title, font=v9.font(27, True))
        draw.text((tx if align == "left" else tx-(box[2]-box[0]), y-32), title, font=v9.font(27, True), fill=color)
        box2 = draw.textbbox((0, 0), detail, font=v9.font(21))
        draw.text((tx if align == "left" else tx-(box2[2]-box2[0]), y+6), detail, font=v9.font(21), fill=(102, 109, 115))
    draw.ellipse((cx-112, cy-112, cx+112, cy+112), fill=(36, 44, 51))
    draw_centered(draw, (cx, cy-30), "CUTTLE", v9.font(38, True), (255, 255, 255))
    draw_centered(draw, (cx, cy+22), "Return to Origin", v9.font(24), (222, 220, 214))
    draw.arc((cx-radius-14, cy-radius-14, cx+radius+14, cy+radius+14), 10, 52, fill=(8, 122, 125), width=9)
    draw.arc((cx-radius-14, cy-radius-14, cx+radius+14, cy+radius+14), 70, 112, fill=(185, 138, 47), width=9)
    draw.arc((cx-radius-14, cy-radius-14, cx+radius+14, cy+radius+14), 130, 172, fill=(182, 91, 62), width=9)
    draw.arc((cx-radius-14, cy-radius-14, cx+radius+14, cy+radius+14), 190, 232, fill=(107, 82, 171), width=9)
    draw.arc((cx-radius-14, cy-radius-14, cx+radius+14, cy+radius+14), 250, 292, fill=(36, 81, 122), width=9)
    draw.arc((cx-radius-14, cy-radius-14, cx+radius+14, cy+radius+14), 310, 352, fill=(36, 44, 51), width=9)
    return canvas


def rebuild_architecture(width=1800, height=950) -> Image.Image:
    canvas = Image.new("RGB", (width, height), (252, 251, 248))
    draw = ImageDraw.Draw(canvas)
    draw.rounded_rectangle((300, 72, 1500, 178), radius=52, fill=(36, 44, 51))
    draw_centered(draw, (900, 104), "统一输入入口", v9.font(34, True), (255, 255, 255))
    draw_centered(draw, (900, 144), "键盘 · 拼音 · 语音 · 选区 · 快捷指令", v9.font(24), (220, 220, 216))
    draw.line((900, 178, 900, 770), fill=(215, 211, 203), width=12)
    modules = [
        ("Context Engine", "读取完成本任务的最小情境", (8, 122, 125), 220),
        ("Intent Router", "判断自治程度与执行位置", (185, 138, 47), 395),
        ("Agent Fabric", "按契约拆分、执行与回滚", (182, 91, 62), 570),
        ("Verification Return", "绑定证据并返回原位置", (107, 82, 171), 745),
    ]
    for title, detail, color, y in modules:
        draw.ellipse((856, y+30, 944, y+118), fill=color)
        draw.rounded_rectangle((360, y, 780, y+150), radius=34, fill=(247, 245, 240), outline=color, width=5)
        draw.text((405, y+33), title, font=v9.font(29, True), fill=color)
        draw.text((405, y+85), detail, font=v9.font(23), fill=(63, 69, 74))
        draw.rounded_rectangle((1020, y, 1440, y+150), radius=34, fill=(247, 245, 240), outline=color, width=5)
        metric = ["最小情境", "自治阈值", "最小权限", "证据闭环"][modules.index((title, detail, color, y))]
        draw.text((1070, y+52), metric, font=v9.font(31, True), fill=color)
    return canvas


def rebuild_validation(width=1800, height=850) -> Image.Image:
    canvas = Image.new("RGB", (width, height), (252, 251, 248))
    draw = ImageDraw.Draw(canvas)
    draw.line((180, 680, 1580, 160), fill=(215, 211, 203), width=16)
    milestones = [(330, 610, "发现问题", "半结构化访谈"), (900, 395, "验证机制", "Wizard of Oz + 对照实验"), (1450, 190, "验证价值", "核心组 + 独立验证组")]
    colors = [(36, 81, 122), (8, 122, 125), (182, 91, 62)]
    for i, ((x, y, title, detail), color) in enumerate(zip(milestones, colors)):
        draw.ellipse((x-42, y-42, x+42, y+42), fill=color)
        draw.ellipse((x-16, y-16, x+16, y+16), fill=(252, 251, 248))
        draw.text((x-125, y+62), title, font=v9.font(31, True), fill=color)
        draw.text((x-125, y+108), detail, font=v9.font(23), fill=(102, 109, 115))
        draw.text((x-12, y-95), f"0{i+1}", font=v9.font(28, True), fill=color)
    draw.text((165, 72), "从问题发现到价值验证，是一条逐级收敛的研究路径", font=v9.font(33, True), fill=(36, 44, 51))
    return canvas


def rebuild_escalation(width=1800, height=820) -> Image.Image:
    canvas = Image.new("RGB", (width, height), (252, 251, 248))
    draw = ImageDraw.Draw(canvas)
    cx, cy = 900, 350
    vertices = [(900, 120), (580, 585), (1220, 585)]
    colors = [(8, 122, 125), (185, 138, 47), (182, 91, 62)]
    titles = [("自治程度", "文本表达 → 智能体执行"), ("协作形态", "单执行者 → 多执行者"), ("执行位置", "本地 → 云端 → 跨设备")]
    for i, (vertex, color, (title, detail)) in enumerate(zip(vertices, colors, titles)):
        draw.line((cx, cy, vertex[0], vertex[1]), fill=(215, 211, 203), width=6)
        draw.ellipse((vertex[0]-48, vertex[1]-48, vertex[0]+48, vertex[1]+48), fill=color)
        draw_centered(draw, (vertex[0], vertex[1]-17), str(i+1), v9.font(30, True), (255, 255, 255))
        tx = vertex[0] - 150 if vertex[0] < cx else vertex[0] - 150
        ty = vertex[1] - 95 if vertex[1] < cy else vertex[1] + 70
        draw.text((tx, ty), title, font=v9.font(29, True), fill=color)
        draw.text((tx, ty+46), detail, font=v9.font(22), fill=(102, 109, 115))
    draw.ellipse((cx-130, cy-130, cx+130, cy+130), fill=(36, 44, 51))
    draw_centered(draw, (cx, cy-35), "受控自治", v9.font(37, True), (255, 255, 255))
    draw_centered(draw, (cx, cy+25), "三维决策", v9.font(27), (220, 220, 216))
    draw.text((340, 700), "高风险动作暂停确认 · 低风险动作自动执行 · 所有升级保留审计与回滚", font=v9.font(27, True), fill=(63, 69, 74))
    return canvas


def rebuild_go_to_market(width=1800, height=760) -> Image.Image:
    canvas = Image.new("RGB", (width, height), (252, 251, 248))
    draw = ImageDraw.Draw(canvas)
    draw.line((190, 380, 1610, 380), fill=(215, 211, 203), width=13)
    stages = [(380, "01", "高频 PC 场景", "知识工作者"), (900, "02", "团队许可", "开发者与研究者"), (1420, "03", "私有部署", "机构与行业客户")]
    colors = [(36, 81, 122), (8, 122, 125), (182, 91, 62)]
    for (x, number, title, detail), color in zip(stages, colors):
        draw.ellipse((x-96, 284, x+96, 476), fill=(247, 245, 240), outline=color, width=8)
        draw.ellipse((x-39, 341, x+39, 419), fill=color)
        draw_centered(draw, (x, 353), number, v9.font(31, True), (255, 255, 255))
        draw_centered(draw, (x, 545), title, v9.font(32, True), color)
        draw_centered(draw, (x, 590), detail, v9.font(24), (102, 109, 115))
    return canvas


def rebuild_roadmap(width=1800, height=740) -> Image.Image:
    canvas = Image.new("RGB", (width, height), (252, 251, 248))
    draw = ImageDraw.Draw(canvas)
    points = [(220, 440), (550, 440), (900, 440), (1240, 440), (1570, 440)]
    colors = [(36, 81, 122), (8, 122, 125), (185, 138, 47), (182, 91, 62), (107, 82, 171)]
    labels = [("0–3 个月", "问题验证"), ("4–6 个月", "团队原型"), ("7–12 个月", "兼容与标准化"), ("13–18 个月", "团队版"), ("19–24 个月", "规模复制")]
    draw.line((points[0][0], 440, points[-1][0], 440), fill=(36, 44, 51), width=8)
    for i, ((x, y), color, (period, title)) in enumerate(zip(points, colors, labels)):
        draw.ellipse((x-18, y-18, x+18, y+18), fill=color)
        box_y = 130 if i % 2 == 0 else 505
        draw.line((x, y-18 if i % 2 == 0 else y+18, x, box_y+120 if i % 2 == 0 else box_y), fill=color, width=4)
        draw.rounded_rectangle((x-125, box_y, x+125, box_y+120), radius=22, fill=(247, 245, 240), outline=color, width=4)
        draw_centered(draw, (x, box_y+22), period, v9.font(25, True), color)
        draw_centered(draw, (x, box_y+65), title, v9.font(30, True), (36, 44, 51))
    draw.text((150, 55), "二十四个月研发与市场路线", font=v9.font(34, True), fill=(36, 44, 51))
    return canvas


def rebuild_competition_map(width=1600, height=1000) -> Image.Image:
    canvas = Image.new("RGB", (width, height), (252, 251, 248))
    draw = ImageDraw.Draw(canvas)
    left, top, right, bottom = 245, 85, 1500, 820
    grid = (222, 218, 209)
    for fraction in (0.25, 0.5, 0.75):
        x = left + int((right - left) * fraction)
        y = bottom - int((bottom - top) * fraction)
        draw.line((x, top, x, bottom), fill=grid, width=2)
        draw.line((left, y, right, y), fill=grid, width=2)
    draw.line((left, bottom, right + 18, bottom), fill=(36, 44, 51), width=5)
    draw.polygon([(right + 18, bottom), (right - 2, bottom - 11), (right - 2, bottom + 11)], fill=(36, 44, 51))
    draw.line((left, bottom, left, top - 18), fill=(36, 44, 51), width=5)
    draw.polygon([(left, top - 18), (left - 11, top + 2), (left + 11, top + 2)], fill=(36, 44, 51))

    vertical = list("执行与治理深度")
    start_y = 290
    for idx, char in enumerate(vertical):
        draw_centered(draw, (88, start_y + idx * 49), char, v9.font(32, True), (36, 44, 51))
    draw_centered(draw, ((left + right) // 2, 892), "任务生命周期覆盖程度", v9.font(35, True), (36, 44, 51))

    points = [
        (0.10, 0.16, "搜狗输入法", (8, 122, 125)),
        (0.17, 0.28, "讯飞输入法", (8, 122, 125)),
        (0.10, 0.41, "Wispr Flow", (185, 138, 47)),
        (0.23, 0.52, "ChatGPT", (36, 44, 51)),
        (0.16, 0.65, "Codex", (36, 44, 51)),
        (0.26, 0.77, "MiMo Desktop", (107, 82, 171)),
        (0.36, 0.88, "Claude Code", (107, 82, 171)),
    ]
    for px, py, label, color in points:
        x = left + int((right - left) * px)
        y = bottom - int((bottom - top) * py)
        draw.ellipse((x - 13, y - 13, x + 13, y + 13), fill=color)
        draw.text((x + 23, y - 19), label, font=v9.font(26, True), fill=color)

    x = left + int((right - left) * 0.72)
    y = bottom - int((bottom - top) * 0.84)
    draw.ellipse((x - 47, y - 47, x + 47, y + 47), fill=(235, 215, 205))
    draw.ellipse((x - 27, y - 27, x + 27, y + 27), fill=(182, 91, 62))
    draw.text((x + 50, y - 38), "Cuttle", font=v9.font(42, True), fill=(182, 91, 62))
    draw.text((x + 52, y + 16), "全生命周期 · 深治理", font=v9.font(25), fill=(102, 109, 115))
    return canvas


def rebuild_finance(width=1800, height=920) -> Image.Image:
    canvas = Image.new("RGB", (width, height), (252, 251, 248))
    draw = ImageDraw.Draw(canvas)
    left, right, zero_y = 230, 1650, 445
    scale = 2.75
    for value in (-100, -50, 0, 50):
        y = zero_y - int(value * scale)
        draw.line((left, y, right, y), fill=(63, 69, 74) if value == 0 else (220, 216, 207), width=4 if value == 0 else 2)
        draw.text((160, y - 16), str(value), font=v9.font(22), fill=(102, 109, 115))
    vertical = list("经营结果")
    for idx, char in enumerate(vertical):
        draw_centered(draw, (72, 300 + idx * 48), char, v9.font(29, True), (36, 44, 51))
    draw_centered(draw, (72, 565), "（万元）", v9.font(23), (102, 109, 115))

    scenarios = ["保守", "基准", "进取"]
    colors = [(185, 138, 47), (8, 122, 125), (36, 44, 51)]
    results = [[-39.0, -48.6, -62.8], [-33.5, -64.5, -97.6], [-27.4, -13.6, 51.2]]
    years = ["第一年", "第二年", "第三年"]
    group_x = [460, 930, 1400]
    bar_w = 80
    for year_index, center in enumerate(group_x):
        draw_centered(draw, (center, 800), years[year_index], v9.font(30, True), (36, 44, 51))
        for scenario_index, value in enumerate(results[year_index]):
            x0 = center - 135 + scenario_index * 105
            x1 = x0 + bar_w
            y1 = zero_y - int(value * scale)
            top, bottom = min(zero_y, y1), max(zero_y, y1)
            draw.rounded_rectangle((x0, top, x1, bottom), radius=9, fill=colors[scenario_index])
            label_y = top - 36 if value >= 0 else bottom + 8
            draw.text((x0 - 4, label_y), f"{value:+.1f}", font=v9.font(21, True), fill=colors[scenario_index])
    for index, (name, color) in enumerate(zip(scenarios, colors)):
        x = 1110 + index * 180
        draw.rounded_rectangle((x, 55, x + 28, 83), radius=5, fill=color)
        draw.text((x + 40, 52), name, font=v9.font(24), fill=(63, 69, 74))
    return canvas


def enhance_embedded_images(docx_path: Path):
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(docx_path, "r") as archive:
        media_names = [name for name in archive.namelist() if name.startswith("word/media/") and name.lower().endswith(".png")]
        raw = {name: archive.read(name) for name in media_names}

    crop_bounds = {
        2: (0, 125, 1800, 570),
        3: (0, 125, 1800, 620),
        4: (0, 120, 1800, 735),
        5: (0, 115, 1800, 800),
        6: (0, 105, 1800, 735),
        10: (0, 115, 1800, 655),
        11: (0, 110, 1800, 680),
    }
    processed = {}
    for name, data in raw.items():
        match = re.search(r"image(\d+)\.png$", name)
        image_number = int(match.group(1)) if match else 0
        image = Image.open(io.BytesIO(data)).convert("RGB")
        if image_number == 2:
            image = rebuild_opportunity()
        elif image_number == 3:
            image = rebuild_loop()
        elif image_number == 4:
            image = rebuild_validation()
        elif image_number == 5:
            image = rebuild_architecture()
        elif image_number == 6:
            image = rebuild_escalation()
        elif image_number == 7:
            image = rebuild_ic_capsule()
        elif image_number == 8:
            image = rebuild_market_funnel()
        elif image_number == 9:
            image = rebuild_competition_map()
        elif image_number == 12:
            image = rebuild_finance()
        elif image_number == 10:
            image = rebuild_go_to_market()
        elif image_number == 11:
            image = rebuild_roadmap()
        elif image_number in crop_bounds:
            image = reframe_crop(image, crop_bounds[image_number])
        else:
            image = palette_only(image)
        asset_path = ASSET_DIR / f"enhanced-{Path(name).name}"
        image.save(asset_path, format="PNG", optimize=True, dpi=(240, 240))
        processed[name] = asset_path.read_bytes()

    temp_docx = docx_path.with_suffix(".media.tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as source_zip, zipfile.ZipFile(temp_docx, "w", zipfile.ZIP_DEFLATED) as target_zip:
        for item in source_zip.infolist():
            target_zip.writestr(item, processed.get(item.filename, source_zip.read(item.filename)))
    temp_docx.replace(docx_path)


def build_document(source: Path, output: Path):
    if not source.exists():
        raise FileNotFoundError(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(source)
    v9.set_document_defaults(doc)
    v9.style_header_footer(doc)
    style_cover(doc, source)
    v9.build_compact_toc(doc)
    inject_ic_explanation(doc)
    v9.style_body(doc)
    v9.style_tables(doc)
    v9.style_inline_images(doc)
    doc.save(output)
    enhance_embedded_images(output)
    return output


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--refresh-toc", type=Path)
    args = parser.parse_args()
    if args.refresh_toc:
        v9.refresh_toc_from_pdf(args.output, args.refresh_toc)
        print(args.output)
    else:
        print(build_document(args.source, args.output))


if __name__ == "__main__":
    main()
