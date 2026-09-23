from __future__ import annotations

import argparse
import math
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"T:\创新创业\OwO-master")
SOURCE = ROOT / "docs" / "Cuttle——情境原生智能输入系统-参赛商业计划书-国金重构版-v8.docx"
OUTPUT = ROOT / "docs" / "Cuttle——情境原生智能输入系统-参赛商业计划书-专业视觉优化版-v9.docx"
ASSET_DIR = ROOT / "docs" / "cuttle_v9_assets"

GRAPHITE = "242C33"
INK = "2C3339"
TEAL = "087A7D"
GOLD = "B98A2F"
TERRACOTTA = "B65B3E"
MUTED = "666D73"
PAPER = "FCFBF8"
SOFT = "F1EFEA"
GRID = "D7D3CB"
WHITE = "FFFFFF"

FONT_CN = "宋体"
FONT_HEAD = "微软雅黑"
FONT_LATIN = "Aptos"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def get_or_add(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def clear_paragraph(paragraph):
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_run(run, *, name=FONT_CN, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = rgb(color)
    if italic is not None:
        run.italic = italic


def set_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=70, start=105, bottom=70, end=105):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, spec in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in spec.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_table_borders(table, outer=GRAPHITE, inner=GRID):
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            spec = {
                "top": {"val": "single", "sz": 10 if row_index == 0 else 4, "color": outer if row_index == 0 else inner},
                "bottom": {"val": "single", "sz": 8 if row_index == len(table.rows) - 1 else 4, "color": outer if row_index == len(table.rows) - 1 else inner},
                "start": {"val": "single", "sz": 8 if col_index == 0 else 4, "color": outer if col_index == 0 else inner},
                "end": {"val": "single", "sz": 8 if col_index == len(row.cells) - 1 else 4, "color": outer if col_index == len(row.cells) - 1 else inner},
                "insideH": {"val": "single", "sz": 4, "color": inner},
                "insideV": {"val": "single", "sz": 4, "color": inner},
            }
            set_cell_border(cell, **spec)


def remove_all_table_borders(table):
    nil = {"val": "nil", "sz": 0, "color": WHITE}
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top=nil, bottom=nil, start=nil, end=nil, insideH=nil, insideV=nil)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_paragraph_border(paragraph, edge: str, color: str, size: int, space: int):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    node = p_bdr.find(qn(f"w:{edge}"))
    if node is None:
        node = OxmlElement(f"w:{edge}")
        p_bdr.append(node)
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), str(size))
    node.set(qn("w:space"), str(space))
    node.set(qn("w:color"), color)


def set_cell_width(cell, width_inches: float):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def insert_table_after(document, paragraph, rows: int, cols: int):
    table = document.add_table(rows=rows, cols=cols)
    paragraph._p.addnext(table._tbl)
    return table


def delete_paragraph(paragraph):
    node = paragraph._element
    node.getparent().remove(node)
    paragraph._p = paragraph._element = None


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "2"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_run(run, name=FONT_LATIN, size=8.5, color=MUTED)


def set_document_defaults(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.75)
    section.left_margin = Cm(2.05)
    section.right_margin = Cm(2.05)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.7)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.line_spacing = 1.42
    normal.paragraph_format.space_after = Pt(4)

    title = doc.styles["Title"]
    title.font.name = FONT_HEAD
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEAD)
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = rgb(GRAPHITE)

    h1 = doc.styles["Heading 1"]
    h1.font.name = FONT_HEAD
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEAD)
    h1.font.size = Pt(16.5)
    h1.font.bold = True
    h1.font.color.rgb = rgb(GRAPHITE)
    h1.paragraph_format.space_before = Pt(13)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = FONT_HEAD
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEAD)
    h2.font.size = Pt(12.5)
    h2.font.bold = True
    h2.font.color.rgb = rgb(GRAPHITE)
    h2.paragraph_format.space_before = Pt(9)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    h3.font.name = FONT_HEAD
    h3._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEAD)
    h3.font.size = Pt(11.2)
    h3.font.bold = True
    h3.font.color.rgb = rgb(TEAL)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def style_header_footer(doc: Document):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("CUTTLE")
    set_run(r1, name=FONT_LATIN, size=8.5, bold=True, color=GRAPHITE)
    r2 = p.add_run("   情境原生智能输入系统  ·  参赛商业计划书")
    set_run(r2, size=8, color=MUTED)
    set_paragraph_border(p, "bottom", GRID, 5, 3)

    footer = section.footer
    fp = footer.paragraphs[0]
    clear_paragraph(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    add_page_field(fp)

    first_header = section.first_page_header
    clear_paragraph(first_header.paragraphs[0])
    first_footer = section.first_page_footer
    clear_paragraph(first_footer.paragraphs[0])


def style_cover(doc: Document, source: Path):
    paragraphs = doc.paragraphs
    p0, p1, p2, p3, p4 = paragraphs[:5]

    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p0.paragraph_format.space_before = Pt(24)
    p0.paragraph_format.space_after = Pt(4)
    for run in p0.runs:
        set_run(run, name=FONT_HEAD, size=10.5, bold=True, color=TEAL)

    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.space_before = Pt(3)
    p1.paragraph_format.space_after = Pt(24)

    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_before = Pt(6)
    p2.paragraph_format.space_after = Pt(10)
    for run in p2.runs:
        set_run(run, name=FONT_HEAD, size=29, bold=True, color=GRAPHITE)

    p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p3.paragraph_format.space_after = Pt(22)
    for run in p3.runs:
        set_run(run, name=FONT_HEAD, size=15, bold=True, color=TEAL)

    p4.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p4.paragraph_format.left_indent = Cm(0.35)
    p4.paragraph_format.right_indent = Cm(0.9)
    p4.paragraph_format.first_line_indent = Cm(0)
    p4.paragraph_format.line_spacing = 1.5
    p4.paragraph_format.space_after = Pt(24)
    set_paragraph_border(p4, "left", GOLD, 18, 8)
    for run in p4.runs:
        set_run(run, size=11, color=INK)

    if doc.inline_shapes:
        logo = doc.inline_shapes[0]
        logo.width = Inches(0.78)
        logo.height = Inches(0.78)

    cover_table = doc.tables[0]
    cover_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cover_table.autofit = False
    remove_all_table_borders(cover_table)
    for row_index, row in enumerate(cover_table.rows):
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Cm(0.82)
        prevent_row_split(row)
        for col_index, cell in enumerate(row.cells):
            set_shading(cell, PAPER)
            set_cell_margins(cell, top=45, start=70, bottom=45, end=70)
            set_cell_width(cell, 1.25 if col_index == 0 else 5.1)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run(
                        run,
                        name=FONT_HEAD if col_index == 0 else FONT_CN,
                        size=8.8 if col_index == 0 else 10.2,
                        bold=col_index == 1,
                        color=MUTED if col_index == 0 else GRAPHITE,
                    )

    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source, "r") as archive:
        flow_image = Image.open(archive.open("word/media/image3.png")).convert("RGB")
    flow_image = palette_enhance(flow_image, crop_title=True)
    flow_image = flow_image.crop((0, 46, flow_image.width, flow_image.height))
    flow_path = ASSET_DIR / "cover-flow-strip.png"
    flow_image.save(flow_path, format="PNG", optimize=True, dpi=(240, 240))
    strip = doc.add_paragraph()
    strip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    strip.paragraph_format.space_before = Pt(18)
    strip.paragraph_format.space_after = Pt(0)
    strip.add_run().add_picture(str(flow_path), width=Inches(6.18))
    cover_table._tbl.addnext(strip._p)


def build_compact_toc(doc: Document):
    original = list(doc.paragraphs)
    toc_heading = next(p for p in original if p.text.strip() == "目录")
    toc_index = original.index(toc_heading)
    body_heading = next(
        p
        for p in original[toc_index + 1 :]
        if p.text.strip() == "项目摘要" and p.style and p.style.name == "Heading 1"
    )
    body_index = original.index(body_heading)
    between = original[toc_index + 1 : body_index]
    entries = []
    for p in between:
        text = p.text.strip()
        match = re.match(r"^(.*?)(\d+)$", text)
        if match:
            title = match.group(1).strip()
            page = max(1, int(match.group(2)) - 1)
            level = 0 if title in ("项目摘要", "参考资料") or title.startswith("第") else 1
            entries.append((title, page, level))

    toc_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    toc_heading.paragraph_format.space_before = Pt(10)
    toc_heading.paragraph_format.space_after = Pt(12)
    clear_paragraph(toc_heading)
    r = toc_heading.add_run("目录")
    set_run(r, name=FONT_HEAD, size=21, bold=True, color=GRAPHITE)
    r = toc_heading.add_run("  CONTENTS")
    set_run(r, name=FONT_LATIN, size=9, bold=True, color=GOLD)

    half = math.ceil(len(entries) / 2)
    row_count = half
    table = insert_table_after(doc, toc_heading, row_count, 4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [2.76, 0.3, 2.76, 0.3]
    for row_index, row in enumerate(table.rows):
        row.height = Cm(0.46)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        prevent_row_split(row)
        for col_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[col_index])
            set_shading(cell, PAPER)
            set_cell_margins(cell, top=15, start=30, bottom=15, end=30)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(
                cell,
                top={"val": "nil", "sz": 0, "color": PAPER},
                start={"val": "nil", "sz": 0, "color": PAPER},
                end={"val": "nil", "sz": 0, "color": PAPER},
                bottom={"val": "dotted", "sz": 3, "color": GRID},
            )

    for side in range(2):
        for row_index in range(row_count):
            entry_index = row_index + side * half
            title_cell = table.cell(row_index, side * 2)
            page_cell = table.cell(row_index, side * 2 + 1)
            if entry_index >= len(entries):
                continue
            title, page, level = entries[entry_index]
            tp = title_cell.paragraphs[0]
            pp = page_cell.paragraphs[0]
            tp.paragraph_format.first_line_indent = Cm(0)
            pp.paragraph_format.first_line_indent = Cm(0)
            tp.paragraph_format.space_before = Pt(0)
            tp.paragraph_format.space_after = Pt(0)
            tp.paragraph_format.line_spacing = 1.0
            pp.paragraph_format.space_before = Pt(0)
            pp.paragraph_format.space_after = Pt(0)
            pp.paragraph_format.line_spacing = 1.0
            tp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if level == 1:
                tp.paragraph_format.left_indent = Cm(0.22)
            tr = tp.add_run(title)
            pr = pp.add_run(str(page))
            set_run(
                tr,
                name=FONT_HEAD if level == 0 else FONT_CN,
                size=8.45 if level == 0 else 7.45,
                bold=level == 0,
                color=GRAPHITE if level == 0 else MUTED,
            )
            set_run(pr, name=FONT_LATIN, size=7.6, bold=level == 0, color=GOLD if level == 0 else MUTED)

    page_break = doc.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    table._tbl.addnext(page_break._p)
    for p in between:
        delete_paragraph(p)
    return entries


def style_body(doc: Document):
    body_started = False
    references_started = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "项目摘要":
            body_started = True
        if not body_started:
            continue

        if paragraph.style and paragraph.style.name == "Heading 1":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.page_break_before = False
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            set_paragraph_border(paragraph, "bottom", GOLD, 7, 4)
            for run in paragraph.runs:
                set_run(run, name=FONT_HEAD, size=16.5, bold=True, color=GRAPHITE)
            if text == "参考资料":
                references_started = True
            continue

        if paragraph.style and paragraph.style.name == "Heading 2":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Cm(0.28)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            set_paragraph_border(paragraph, "left", TEAL, 14, 7)
            for run in paragraph.runs:
                set_run(run, name=FONT_HEAD, size=12.5, bold=True, color=GRAPHITE)
            continue

        if paragraph.style and paragraph.style.name == "Heading 3":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            for run in paragraph.runs:
                set_run(run, name=FONT_HEAD, size=11.2, bold=True, color=TEAL)
            continue

        if text.startswith("图 ") or text.startswith("表 "):
            is_table = text.startswith("表 ")
            match = re.match(r"^([图表]\s*\d+)\s*(.*)$", text)
            clear_paragraph(paragraph)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.keep_with_next = is_table
            if match:
                prefix = paragraph.add_run(match.group(1).replace(" ", ""))
                set_run(prefix, name=FONT_HEAD, size=9, bold=True, color=GOLD if is_table else TEAL)
                title = paragraph.add_run("  " + match.group(2))
                set_run(title, size=9.2, color=INK)
            continue

        has_drawing = bool(paragraph._p.xpath(".//w:drawing"))
        if has_drawing:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(2)
            continue

        if references_started and text:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Cm(0.55)
            paragraph.paragraph_format.first_line_indent = Cm(-0.55)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0.8)
            for r_node in paragraph._p.xpath(".//w:r"):
                r_pr = r_node.get_or_add_rPr()
                color = get_or_add(r_pr, "w:color")
                color.set(qn("w:val"), INK)
                underline = get_or_add(r_pr, "w:u")
                underline.set(qn("w:val"), "none")
                fonts = get_or_add(r_pr, "w:rFonts")
                fonts.set(qn("w:eastAsia"), FONT_CN)
                size = get_or_add(r_pr, "w:sz")
                size.set(qn("w:val"), "16")
            continue

        if text:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Pt(21)
            paragraph.paragraph_format.line_spacing = 1.42
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.widow_control = True
            for run in paragraph.runs:
                set_run(run, size=10.5, color=INK)


def style_tables(doc: Document):
    for table_index, table in enumerate(doc.tables[2:], start=1):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_borders(table)
        set_repeat_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            prevent_row_split(row)
            for col_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell, top=55, start=78, bottom=55, end=78)
                if row_index == 0:
                    set_shading(cell, GRAPHITE)
                elif col_index == 0:
                    set_shading(cell, SOFT)
                else:
                    set_shading(cell, PAPER)

                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                    paragraph.paragraph_format.left_indent = Cm(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.12
                    content = paragraph.text.strip()
                    if row_index == 0:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        short = len(content) <= 13 and not re.search(r"[，；。：]", content)
                        numeric = bool(re.fullmatch(r"[\d\s./%—－–+负正万元人个至不超过]+", content))
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if short or numeric else WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        set_run(
                            run,
                            size=8.2 if table_index == 9 else 8.7,
                            bold=row_index == 0 or (col_index == 0 and len(content) <= 18),
                            color=WHITE if row_index == 0 else INK,
                        )


def style_inline_images(doc: Document):
    for index, shape in enumerate(doc.inline_shapes):
        if index == 0:
            continue
        ratio = shape.height / shape.width if shape.width else 0.55
        width = Inches(6.28)
        shape.width = width
        shape.height = int(width * ratio)
        doc_pr = shape._inline.docPr
        doc_pr.set("title", f"Cuttle 图示 {index}")
        if not doc_pr.get("descr"):
            doc_pr.set("descr", "Cuttle 商业计划书中的技术、验证或经营信息图")


def replace_color_near(arr, source, target, radius=18):
    src = np.array(source, dtype=np.int16)
    tgt = np.array(target, dtype=np.float32)
    work = arr.astype(np.int32)
    dist = np.sqrt(np.sum((work - src) ** 2, axis=2))
    mask = dist <= radius
    if np.any(mask):
        alpha = np.clip(1 - dist[mask] / max(radius, 1), 0.45, 1.0)[:, None]
        arr[mask] = (arr[mask].astype(np.float32) * (1 - alpha) + tgt * alpha).astype(np.uint8)


def palette_enhance(image: Image.Image, crop_title=True) -> Image.Image:
    image = image.convert("RGB")
    if crop_title and image.height >= 700:
        top = max(78, int(image.height * 0.115))
        image = image.crop((0, top, image.width, image.height))
    arr = np.array(image)
    mappings = [
        ((255, 255, 255), (252, 251, 248), 6),
        ((243, 246, 248), (241, 239, 234), 14),
        ((22, 59, 101), (36, 44, 51), 22),
        ((28, 154, 165), (8, 122, 125), 22),
        ((239, 106, 74), (182, 91, 62), 20),
        ((227, 166, 47), (185, 138, 47), 20),
        ((102, 113, 126), (102, 109, 115), 16),
        ((29, 39, 51), (36, 42, 47), 16),
    ]
    for source, target, radius in mappings:
        replace_color_near(arr, source, target, radius)
    image = Image.fromarray(arr)
    image = ImageEnhance.Contrast(image).enhance(1.035)
    image = image.filter(ImageFilter.UnsharpMask(radius=0.7, percent=115, threshold=2))
    return image


def font(size: int, bold=False):
    path = Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc")
    return ImageFont.truetype(str(path), size=size)


def rebuild_market_funnel(width=1800, height=700) -> Image.Image:
    canvas = Image.new("RGB", (width, height), (252, 251, 248))
    draw = ImageDraw.Draw(canvas)
    draw.line((220, 110, 220, 585), fill=(183, 178, 168), width=5)
    items = [
        (150, (36, 44, 51), "交付容量", "≤ 200 人", "按首年服务承载能力配置", 1.00),
        (330, (8, 122, 125), "深度观察", "50–80 人", "纵向核心组 15–25 人  ·  独立验证组 20–30 人", 0.62),
        (510, (185, 138, 47), "正式实验", "35–55 人", "进入对照实验与消融实验", 0.43),
    ]
    for y, color, label, value, note, scale in items:
        draw.ellipse((196, y - 24, 244, y + 24), fill=color)
        draw.text((295, y - 55), label, font=font(34, True), fill=color)
        draw.text((295, y - 4), value, font=font(54, True), fill=(36, 44, 51))
        draw.text((700, y + 8), note, font=font(28), fill=(102, 109, 115))
        draw.rounded_rectangle((700, y + 58, 700 + int(760 * scale), y + 72), radius=7, fill=color)
    draw.text((295, 625), "验证样本逐级收敛：先确认可交付，再验证机制，最后进入正式实验", font=font(25), fill=(102, 109, 115))
    return canvas


def rebuild_finance(width=1800, height=820) -> Image.Image:
    canvas = Image.new("RGB", (width, height), (252, 251, 248))
    draw = ImageDraw.Draw(canvas)
    plot = (170, 90, 1660, 610)
    zero_y = 380
    draw.line((plot[0], zero_y, plot[2], zero_y), fill=(63, 69, 74), width=4)
    draw.text((75, zero_y - 17), "0", font=font(23), fill=(102, 109, 115))
    for value in (-100, -50, 50):
        y = zero_y - int(value * 2.6)
        draw.line((plot[0], y, plot[2], y), fill=(220, 216, 207), width=2)
        draw.text((55, y - 16), f"{value}", font=font(22), fill=(102, 109, 115))
    scenarios = ["保守", "基准", "进取"]
    colors = [(185, 138, 47), (8, 122, 125), (36, 44, 51)]
    results = [[-39.0, -48.6, -62.8], [-33.5, -64.5, -97.6], [-27.4, -13.6, 51.2]]
    years = ["第一年", "第二年", "第三年"]
    group_x = [430, 910, 1390]
    bar_w = 82
    for year_index, center in enumerate(group_x):
        draw.text((center - 56, 660), years[year_index], font=font(31, True), fill=(36, 44, 51))
        for scenario_index, value in enumerate(results[year_index]):
            x0 = center - 150 + scenario_index * 110
            x1 = x0 + bar_w
            y1 = zero_y - int(value * 2.6)
            top, bottom = min(zero_y, y1), max(zero_y, y1)
            draw.rounded_rectangle((x0, top, x1, bottom), radius=10, fill=colors[scenario_index])
            label = f"{value:+.1f}"
            label_y = top - 38 if value >= 0 else bottom + 8
            draw.text((x0 - 5, label_y), label, font=font(22, True), fill=colors[scenario_index])
    draw.text((170, 735), "单位：万元  ·  图示为营业收入减经营成本后的年度经营结果", font=font(24), fill=(102, 109, 115))
    for index, (name, color) in enumerate(zip(scenarios, colors)):
        x = 1120 + index * 170
        draw.rounded_rectangle((x, 22, x + 28, 50), radius=5, fill=color)
        draw.text((x + 40, 19), name, font=font(23), fill=(63, 69, 74))
    return canvas


def enhance_embedded_images(docx_path: Path):
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(docx_path, "r") as archive:
        media_names = [name for name in archive.namelist() if name.startswith("word/media/") and name.lower().endswith(".png")]
        raw = {name: archive.read(name) for name in media_names}

    processed = {}
    for name, data in raw.items():
        image_number_match = re.search(r"image(\d+)\.png$", name)
        image_number = int(image_number_match.group(1)) if image_number_match else 0
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            tmp.write(data)
            tmp_path = Path(tmp.name)
        try:
            image = Image.open(tmp_path).convert("RGB")
            if image_number == 8:
                image = rebuild_market_funnel()
            elif image_number == 12:
                image = rebuild_finance()
            elif image_number == 1:
                image = palette_enhance(image, crop_title=False)
            elif image_number >= 13:
                image = palette_enhance(image, crop_title=False)
            else:
                crop_title = image_number != 9
                image = palette_enhance(image, crop_title=crop_title)
            asset_path = ASSET_DIR / f"enhanced-{Path(name).name}"
            image.save(asset_path, format="PNG", optimize=True, dpi=(240, 240))
            processed[name] = asset_path.read_bytes()
        finally:
            tmp_path.unlink(missing_ok=True)

    temp_docx = docx_path.with_suffix(".media.tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as source_zip, zipfile.ZipFile(temp_docx, "w", zipfile.ZIP_DEFLATED) as target_zip:
        for item in source_zip.infolist():
            payload = processed.get(item.filename, source_zip.read(item.filename))
            target_zip.writestr(item, payload)
    temp_docx.replace(docx_path)


def build_document(source: Path, output: Path):
    if not source.exists():
        raise FileNotFoundError(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(source)
    set_document_defaults(doc)
    style_header_footer(doc)
    style_cover(doc, source)
    build_compact_toc(doc)
    style_body(doc)
    style_tables(doc)
    style_inline_images(doc)
    doc.save(output)
    enhance_embedded_images(output)
    return output


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def refresh_toc_from_pdf(docx_path: Path, pdf_path: Path):
    from pypdf import PdfReader

    doc = Document(docx_path)
    toc_table = doc.tables[1]
    reader = PdfReader(str(pdf_path))
    pages = [normalize_text(page.extract_text() or "") for page in reader.pages]

    for row in toc_table.rows:
        for side in range(2):
            title_cell = row.cells[side * 2]
            page_cell = row.cells[side * 2 + 1]
            title = title_cell.text.strip()
            if not title:
                continue
            needle = normalize_text(title)
            page_number = None
            for index, page_text in enumerate(pages[2:], start=3):
                if needle in page_text:
                    page_number = index
                    break
            if page_number is None:
                continue
            p = page_cell.paragraphs[0]
            clear_paragraph(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            level = 0 if title in ("项目摘要", "参考资料") or title.startswith("第") else 1
            run = p.add_run(str(page_number))
            set_run(run, name=FONT_LATIN, size=7.6, bold=level == 0, color=GOLD if level == 0 else MUTED)
    doc.save(docx_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--refresh-toc", type=Path)
    args = parser.parse_args()
    if args.refresh_toc:
        refresh_toc_from_pdf(args.output, args.refresh_toc)
        print(args.output)
    else:
        print(build_document(args.source, args.output))


if __name__ == "__main__":
    main()
